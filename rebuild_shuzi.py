#!/usr/bin/env python3
"""Rebuild 数字.xlsx into 数字_已修复.xlsx: legal order, keep case semantics, MP→HJ."""
from __future__ import annotations

import json
import math
import re
import shutil
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

import openpyxl
from openpyxl.styles import numbers

SRC = Path("/Users/xixi/Documents/工作文件/复盘/e2e测试数据/数字.xlsx")
DST = Path("/Users/xixi/Documents/工作文件/复盘/e2e测试数据/数字_已修复.xlsx")
EPS = 1e-9
NAME_RE = re.compile(r"【(.+?)】(.+?)位置_(翻前|翻牌圈|转牌圈|河牌圈)(.+)")

# Clockwise from BTN. 6/7-max use HJ (same noun as 9-max), no MP.
CW_FROM_BTN = {
    3: ["BTN", "SB", "BB"],
    4: ["BTN", "SB", "BB", "UTG"],
    5: ["BTN", "SB", "BB", "UTG", "CO"],
    6: ["BTN", "SB", "BB", "UTG", "HJ", "CO"],
    7: ["BTN", "SB", "BB", "UTG", "UTG+1", "HJ", "CO"],
    8: ["BTN", "SB", "BB", "UTG", "UTG+1", "LJ", "HJ", "CO"],
    9: ["BTN", "SB", "BB", "UTG", "UTG+1", "UTG+2", "LJ", "HJ", "CO"],
}

PREFLOP_ORDER = {
    3: ["BTN", "SB", "BB"],
    4: ["UTG", "BTN", "SB", "BB"],
    5: ["UTG", "CO", "BTN", "SB", "BB"],
    6: ["UTG", "HJ", "CO", "BTN", "SB", "BB"],
    7: ["UTG", "UTG+1", "HJ", "CO", "BTN", "SB", "BB"],
    8: ["UTG", "UTG+1", "LJ", "HJ", "CO", "BTN", "SB", "BB"],
    9: ["UTG", "UTG+1", "UTG+2", "LJ", "HJ", "CO", "BTN", "SB", "BB"],
}


def postflop_order(n: int) -> list[str]:
    pf = PREFLOP_ORDER[n]
    if n == 3:
        return ["SB", "BB", "BTN"]
    core = [p for p in pf if p not in ("SB", "BB")]
    return ["SB", "BB"] + core


def assign_seats(n: int, hero_pos: str) -> dict[str, str]:
    cw = CW_FROM_BTN[n]
    if hero_pos not in cw:
        raise ValueError(f"hero {hero_pos} not in {cw}")
    idx = cw.index(hero_pos)
    return {str(k): cw[(idx + k) % n] for k in range(n)}


def num(x, default=0.0):
    if x is None or x == "":
        return default
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def classify(rest: str) -> str:
    r = rest or ""
    if "3bet" in r.lower():
        return "3bet"
    if "4bet" in r.lower():
        return "4bet"
    if "Open" in r:
        return "open"
    if "Limp" in r:
        return "limp"
    if "Allin" in r or "All-in" in r or "全押" in r:
        return "allin"
    if "Fold" in r or "弃牌" in r:
        return "fold"
    if "Check" in r or "让牌" in r:
        return "check"
    if "Raise" in r or "加注" in r:
        return "raise"
    if "Bet" in r or "主动下注" in r:
        return "bet"
    return "open"


class Builder:
    def __init__(self, orig: dict, n: int, hero_pos: str, sheet: str, kind: str):
        self.orig = orig
        self.n = n
        self.hero_pos = hero_pos
        self.sheet = sheet
        self.kind = kind
        self.seat_positions = assign_seats(n, hero_pos)
        self.pos_seat = {v: int(k) for k, v in self.seat_positions.items()}
        self.hero_seat = 0
        self.sb = float(orig["small_blind"])
        self.bb = float(orig["big_blind"])
        self.ante = float(orig.get("ante") or 0)
        self.init = {}
        for k in orig.get("final_stacks") or {}:
            self.init[int(k)] = num(orig["final_stacks"][k]) - num((orig.get("net_won_chips") or {}).get(k))
        if len(self.init) != n:
            self.init = {i: 200.0 * max(self.bb, 1.0) for i in range(n)}
        self.remaining = dict(self.init)
        if self.ante:
            for s in self.remaining:
                self.remaining[s] -= self.ante
                if self.remaining[s] < 0:
                    self.remaining[s] = 0
        self.actions: list[dict] = []
        self.folded: set[int] = set()
        self.allin: set[int] = set()
        self.street = None
        self.contrib = defaultdict(float)
        self.current = 0.0
        self.ts = 1778824769.0

    def can_act(self):
        return [s for s in range(self.n) if s not in self.folded and s not in self.allin]

    def emit(self, seat: int, action: str, amount: float | None, street: str, is_all_in=False):
        pos = self.seat_positions[str(seat)]
        self.actions.append(
            {
                "seat": seat,
                "action": action,
                "amount": 0.0 if amount is None else amount,
                "street": street,
                "position": pos,
                "timestamp": self.ts,
                "is_all_in": bool(is_all_in),
            }
        )
        self.ts += 0.05

    def post_blinds(self):
        sb_s = self.pos_seat["SB"]
        bb_s = self.pos_seat["BB"]
        sb_amt = min(self.sb, self.remaining[sb_s])
        self.remaining[sb_s] -= sb_amt
        self.contrib[sb_s] += sb_amt
        self.emit(sb_s, "POST_BLIND", sb_amt, "PRE_HAND", self.remaining[sb_s] <= EPS)
        if self.remaining[sb_s] <= EPS:
            self.remaining[sb_s] = 0
            self.allin.add(sb_s)
        bb_amt = min(self.bb, self.remaining[bb_s])
        self.remaining[bb_s] -= bb_amt
        self.contrib[bb_s] += bb_amt
        self.emit(bb_s, "POST_BLIND", bb_amt, "PRE_HAND", self.remaining[bb_s] <= EPS)
        if self.remaining[bb_s] <= EPS:
            self.remaining[bb_s] = 0
            self.allin.add(bb_s)
        self.street = "PREFLOP"
        self.current = max(self.contrib.values()) if self.contrib else self.bb

    def start_street(self, street: str):
        self.street = street
        self.contrib = defaultdict(float)
        self.current = 0.0

    def do_fold(self, seat: int):
        self.emit(seat, "FOLD", 0.0, self.street, False)
        self.folded.add(seat)

    def do_check(self, seat: int):
        self.emit(seat, "CHECK", 0.0, self.street, False)

    def _put(self, seat: int, want: float, action: str, force_allin=False):
        want = max(0.0, want)
        put = min(want, self.remaining[seat])
        is_ai = force_allin or put + EPS >= self.remaining[seat] or action == "ALL_IN"
        self.remaining[seat] -= put
        self.contrib[seat] += put
        if self.remaining[seat] <= EPS:
            self.remaining[seat] = 0
            is_ai = True
            self.allin.add(seat)
        if is_ai and action in ("RAISE", "BET", "CALL"):
            action = "ALL_IN"
        self.emit(seat, action, put, self.street, is_ai)
        if self.contrib[seat] > self.current + EPS:
            self.current = self.contrib[seat]
        return put

    def do_call(self, seat: int):
        need = max(0.0, self.current - self.contrib[seat])
        self._put(seat, need, "CALL")

    def do_bet(self, seat: int, size: float):
        self._put(seat, size, "BET")

    def do_raise_to(self, seat: int, street_total: float):
        need = max(0.0, street_total - self.contrib[seat])
        self._put(seat, need, "RAISE")

    def do_allin(self, seat: int):
        self._put(seat, self.remaining[seat], "ALL_IN", force_allin=True)

    def alive_order(self, order: list[str]):
        out = []
        for p in order:
            s = self.pos_seat[p]
            if s not in self.folded and s not in self.allin:
                out.append(s)
        return out

    def fold_rest_after(self, actor: int, order: list[str]):
        full = [self.pos_seat[p] for p in order]
        i = full.index(actor)
        for s in full[i + 1 :] + full[:i]:
            if s not in self.folded and s not in self.allin:
                self.do_fold(s)

    def preflop_fold_to(self, target: int):
        for s in self.alive_order(PREFLOP_ORDER[self.n]):
            if s == target:
                break
            if s not in self.folded:
                self.do_fold(s)

    def postflop_check_to(self, target: int, order: list[str]):
        for s in self.alive_order(order):
            if s == target:
                break
            self.do_check(s)

    def preflop_everyone_limp(self):
        """All complete BB, BB checks. Everyone sees next street."""
        self.street = "PREFLOP"
        for s in self.alive_order(PREFLOP_ORDER[self.n]):
            pos = self.seat_positions[str(s)]
            if pos == "BB":
                self.do_check(s)
            else:
                self.do_call(s)

    def street_all_check(self, street: str):
        self.start_street(street)
        for s in self.alive_order(postflop_order(self.n)):
            self.do_check(s)

    def finish_to_showdown_if_needed(self):
        """2+ still in after the named street: check down remaining streets. Do not truncate."""
        alive = [s for s in range(self.n) if s not in self.folded]
        if len(alive) < 2:
            return
        seq = ["PREFLOP", "FLOP", "TURN", "RIVER"]
        if self.street not in seq:
            return
        for st in seq[seq.index(self.street) + 1 :]:
            self.street_all_check(st)

    def build_preflop_named(self):
        self.street = "PREFLOP"
        hs = self.hero_seat
        order = PREFLOP_ORDER[self.n]
        bb_s = self.pos_seat["BB"]
        open_to = 3 * self.bb
        three_to = 9 * self.bb
        four_to = 27 * self.bb

        if self.kind == "open":
            if self.hero_pos == "BB":
                # limp from first, BB raises
                first = self.alive_order(order)[0]
                if first == hs:
                    first = self.alive_order(order)[1]
                self.preflop_fold_to(first)
                self.do_call(first)
                for s in self.alive_order(order):
                    if s not in (first, hs):
                        self.do_fold(s)
                self.do_raise_to(hs, open_to)
                if first not in self.folded:
                    self.do_fold(first)
            else:
                self.preflop_fold_to(hs)
                self.do_raise_to(hs, open_to)
                self.fold_rest_after(hs, order)

        elif self.kind == "3bet":
            pf = self.alive_order(order)
            if pf[0] == hs:
                # limp-reraise: legal UTG/BTN 3bet
                self.do_call(hs)
                raiser = None
                for s in pf[1:]:
                    if s != hs and s != bb_s:
                        raiser = s
                        break
                if raiser is None:
                    raiser = pf[-1] if pf[-1] != hs else pf[1]
                for s in pf[1:]:
                    if s == raiser:
                        self.do_raise_to(s, open_to)
                    elif s != hs:
                        self.do_fold(s)
                self.do_raise_to(hs, three_to)
                if raiser not in self.folded:
                    self.do_fold(raiser)
            else:
                opener = pf[0]
                self.do_raise_to(opener, open_to)
                for s in pf[1:]:
                    if s == hs:
                        self.do_raise_to(hs, three_to)
                    else:
                        self.do_fold(s)
                if opener not in self.folded:
                    self.do_fold(opener)

        elif self.kind == "4bet":
            before = []
            for s in self.alive_order(order):
                if s == hs:
                    break
                before.append(s)
            behind = [s for s in self.alive_order(order) if s not in before and s != hs]
            if len(before) < 2:
                for s in before:
                    self.do_fold(s)
                self.do_raise_to(hs, open_to)
                threebettor = behind[0]
                self.do_raise_to(threebettor, three_to)
                for s in behind[1:]:
                    self.do_fold(s)
                self.do_raise_to(hs, four_to)
                if threebettor not in self.folded:
                    self.do_fold(threebettor)
            else:
                self.do_raise_to(before[0], open_to)
                self.do_raise_to(before[1], three_to)
                for s in before[2:]:
                    self.do_fold(s)
                self.do_raise_to(hs, four_to)
                self.fold_rest_after(hs, order)

        elif self.kind == "limp":
            self.preflop_fold_to(hs)
            if self.hero_pos == "BB":
                self.do_check(hs)
            else:
                self.do_call(hs)
            self.fold_rest_after(hs, order)

        elif self.kind == "fold":
            if self.alive_order(order)[0] == hs:
                self.do_fold(hs)
                # Walk: everyone else folds, BB does not act and wins blinds/antes.
                bb_s = self.pos_seat["BB"]
                for s in self.alive_order(order):
                    if s == bb_s:
                        continue
                    self.do_fold(s)
            else:
                opener = self.alive_order(order)[0]
                self.do_raise_to(opener, open_to)
                for s in self.alive_order(order):
                    if s == hs:
                        self.do_fold(s)
                        break
                    if s != opener:
                        self.do_fold(s)
                self.fold_rest_after(opener, order)

        elif self.kind == "allin":
            if self.hero_pos == "BB":
                first = self.alive_order(order)[0]
                if first == hs:
                    first = self.alive_order(order)[1]
                self.preflop_fold_to(first)
                self.do_call(first)
                for s in self.alive_order(order):
                    if s not in (first, hs):
                        self.do_fold(s)
                self.do_allin(hs)
                if first not in self.folded:
                    self.do_fold(first)
            else:
                self.preflop_fold_to(hs)
                self.do_allin(hs)
                self.fold_rest_after(hs, order)
        else:
            self.preflop_fold_to(hs)
            self.do_raise_to(hs, open_to)
            self.fold_rest_after(hs, order)

    def build_postflop_named(self, street: str):
        self.start_street(street)
        hs = self.hero_seat
        order = postflop_order(self.n)
        size = max(self.bb, 0.5 * self.bb * 2)

        if self.kind == "check":
            for s in self.alive_order(order):
                self.do_check(s)
            return

        if self.kind == "bet":
            self.postflop_check_to(hs, order)
            self.do_bet(hs, size)
            self.fold_rest_after(hs, order)
            return

        if self.kind == "allin":
            self.postflop_check_to(hs, order)
            self.do_allin(hs)
            self.fold_rest_after(hs, order)
            return

        if self.kind == "fold":
            al = self.alive_order(order)
            if al[0] == hs:
                self.do_check(hs)
                bettor = al[1]
                self.do_bet(bettor, size)
            else:
                bettor = al[0]
                self.do_bet(bettor, size)
            self.fold_rest_after(bettor, order)
            return

        if self.kind == "raise":
            al = self.alive_order(order)
            if al[0] == hs:
                self.do_check(hs)
                bettor = al[1]
                self.do_bet(bettor, size)
                for s in al[2:]:
                    self.do_fold(s)
                self.do_raise_to(hs, size * 2 + self.contrib[hs])
                if bettor not in self.folded:
                    self.do_fold(bettor)
            else:
                bettor = al[0]
                self.do_bet(bettor, size)
                for s in al[1:]:
                    if s == hs:
                        self.do_raise_to(hs, size * 2)
                    else:
                        self.do_fold(s)
                if bettor not in self.folded:
                    self.do_fold(bettor)
            return

        # default bet
        self.postflop_check_to(hs, order)
        self.do_bet(hs, size)
        self.fold_rest_after(hs, order)

    def settle(self):
        invested = {s: float(self.init[s] - self.remaining[s]) for s in range(self.n)}
        # last street contribs for uncalled
        last = None
        last_amt = defaultdict(float)
        for a in self.actions:
            if a["action"] == "POST_BLIND":
                continue
            st = a["street"]
            if st != last:
                last = st
                last_amt = defaultdict(float)
            if a["action"] != "FOLD":
                last_amt[a["seat"]] += num(a["amount"])
        # include blinds in preflop last_amt already via POST not in last_amt if last is PREFLOP
        # uncalled from last street totals: use contrib if still on last street
        if last == "PREFLOP":
            street_tot = defaultdict(float)
            for a in self.actions:
                if a["street"] in ("PRE_HAND", "PREFLOP") and a["action"] != "FOLD":
                    street_tot[a["seat"]] += num(a["amount"])
            vals = list(street_tot.values()) or [0]
        else:
            vals = list(self.contrib.values()) or [0]
            street_tot = dict(self.contrib)
        ranked = sorted(street_tot.values(), reverse=True) if street_tot else [0]
        uncalled = max(0.0, ranked[0] - (ranked[1] if len(ranked) > 1 else 0))
        uncalled_seat = None
        if uncalled > EPS and street_tot:
            uncalled_seat = max(street_tot, key=street_tot.get)
            invested[uncalled_seat] -= uncalled
            self.remaining[uncalled_seat] += uncalled
        pot = sum(invested.values())
        alive = [s for s in range(self.n) if s not in self.folded]
        winner = self.hero_seat if self.hero_seat in alive else (alive[0] if alive else 0)
        nets = {s: -invested[s] for s in range(self.n)}
        nets[winner] += pot
        finals = {s: self.init[s] + nets[s] for s in range(self.n)}
        return invested, uncalled, winner, nets, finals, pot

    def to_json(self) -> dict:
        invested, uncalled, winner, nets, finals, pot = self.settle()
        j = deepcopy(self.orig)
        j["player_count"] = self.n
        j["seat_positions"] = self.seat_positions
        j["dealer_position"] = self.pos_seat["BTN"]
        j["hero_seat"] = 0
        # Seat 0–8 always the same nicknames (not Hero(BTN) / Player(SB)).
        nick_by_seat = {
            0: "Hero",
            1: "Ada",
            2: "Ben",
            3: "Cora",
            4: "Drew",
            5: "Eve",
            6: "Finn",
            7: "Gia",
            8: "Hank",
        }
        names = {str(i): nick_by_seat[i] for i in range(self.n)}
        j["seat_player_names"] = names
        j["seat_nicknames"] = dict(names)
        # remap hole cards: keep by seat index (hero still 0)
        holes = self.orig.get("hole_cards") or {}
        j["hole_cards"] = {}
        for i in range(self.n):
            j["hole_cards"][str(i)] = holes.get(str(i)) or holes.get(i)
        played = set()
        for a in self.actions:
            played.add("PREFLOP" if a["street"] == "PRE_HAND" else a["street"])
        holes = j["hole_cards"]
        flop = list(self.orig.get("flop") or [])
        turn = as_card(self.orig.get("turn"))
        river = as_card(self.orig.get("river"))
        used = collect_used_cards(holes, [], None, None)
        need_flop = bool(played & {"FLOP", "TURN", "RIVER"})
        need_turn = bool(played & {"TURN", "RIVER"})
        need_river = "RIVER" in played
        if need_flop:
            flop = [c for c in flop if c and c not in used]
            used.update(flop)
            while len(flop) < 3:
                flop.append(next_free_card(used))
            flop = flop[:3]
        else:
            flop = []
        if need_turn:
            if not turn or turn in used:
                turn = next_free_card(used)
            else:
                used.add(turn)
        else:
            turn = None
        if need_river:
            if not river or river in used:
                river = next_free_card(used)
            else:
                used.add(river)
        else:
            river = None
        j["flop"], j["turn"], j["river"] = flop, turn, river
        j["actions"] = self.actions
        j["net_won_chips"] = {str(s): round(nets[s], 6) for s in range(self.n)}
        j["final_stacks"] = {str(s): round(finals[s], 6) for s in range(self.n)}
        j["is_complete"] = True
        j["is_showdown"] = len([s for s in range(self.n) if s not in self.folded]) > 1 and self.kind == "check"
        j["winner_seat"] = winner
        j["fault_count"] = 0
        j.pop("total_pot", None)
        j.pop("total_invested", None)
        j.pop("uncalled_returned", None)
        return j


def hero_actions_count(j):
    hs = int(j["hero_seat"])
    return sum(1 for a in j["actions"] if int(a["seat"]) == hs and a["action"] != "POST_BLIND")


def last_hero_action_label(j, sheet):
    st = {"PREFLOP": "PREFLOP", "FLOP": "FLOP", "TURN": "TURN", "RIVER": "RIVER"}[sheet]
    hs = int(j["hero_seat"])
    acts = [a for a in j["actions"] if int(a["seat"]) == hs and a.get("street") == st]
    if not acts:
        return j.get("_kind", "open")
    a = acts[-1]
    if a.get("is_all_in") or a["action"] == "ALL_IN":
        return "all in"
    return {
        "RAISE": "raise",
        "BET": "bet",
        "FOLD": "fold",
        "CHECK": "check",
        "CALL": "call",
    }.get(a["action"], a["action"].lower())


def fill_sequence(ws, row, headers, actions):
    """Write street/position/action/bet_amount columns; clear old first."""
    street_idxs = []
    pos_idxs, act_idxs, amt_idxs = [], [], []
    for i, h in enumerate(headers):
        if not h:
            continue
        hs = str(h)
        if hs == "street" or hs.startswith("street") or hs.startswith("steet") or hs.startswith("strret"):
            street_idxs.append(i)
        elif hs == "position" or hs.startswith("position"):
            pos_idxs.append(i)
        elif hs == "action" or (hs.startswith("action.") and not hs.startswith("动作")):
            act_idxs.append(i)
        elif hs == "bet_amount" or hs.startswith("bet_amount"):
            amt_idxs.append(i)
    # column 15 动作 is named 动作 not action — already skipped
    for idxs in (street_idxs, pos_idxs, act_idxs, amt_idxs):
        for i in idxs:
            ws.cell(row, i + 1).value = None
    si = pi = 0
    last_st = object()
    for a in actions:
        st = a["street"]
        if st != last_st:
            if si < len(street_idxs):
                ws.cell(row, street_idxs[si] + 1).value = st
                si += 1
            last_st = st
        if pi >= len(pos_idxs):
            break
        ws.cell(row, pos_idxs[pi] + 1).value = a["position"]
        if pi < len(act_idxs):
            ws.cell(row, act_idxs[pi] + 1).value = a["action"]
        if pi < len(amt_idxs):
            amt = a.get("amount")
            if a["action"] in ("FOLD", "CHECK") or amt in (0, 0.0, None):
                ws.cell(row, amt_idxs[pi] + 1).value = None if a["action"] in ("FOLD", "CHECK") else amt
            else:
                ws.cell(row, amt_idxs[pi] + 1).value = amt
        pi += 1


def fix_headers(ws):
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h == "steet.2":
            ws.cell(1, c).value = "street.2"
        elif h == "strret.4":
            ws.cell(1, c).value = "street.4"


RANKS = "A23456789TJQK"
SUITS = "shdc"


def as_card(x):
    if not x:
        return None
    if isinstance(x, (list, tuple)):
        return x[0] if x else None
    return str(x)


def collect_used_cards(hole_cards, flop, turn, river) -> set[str]:
    used: set[str] = set()
    for cards in (hole_cards or {}).values():
        if cards:
            used.update(cards)
    used.update(flop or [])
    t, r = as_card(turn), as_card(river)
    if t:
        used.add(t)
    if r:
        used.add(r)
    return used


def next_free_card(used: set[str]) -> str:
    for rank in RANKS:
        for suit in SUITS:
            card = rank + suit
            if card not in used:
                used.add(card)
                return card
    raise RuntimeError("deck exhausted")


def blinds_text(j):
    sb, bb = j["small_blind"], j["big_blind"]
    def fmt(x):
        x = float(x)
        return str(int(x)) if abs(x - int(x)) < 1e-9 else ("%g" % x)
    return "%s/%s" % (fmt(sb), fmt(bb))


def main():
    shutil.copy2(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    n_rows = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        fix_headers(ws)
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        hidx = {h: i + 1 for i, h in enumerate(headers) if h}
        jcol = hidx["auto_json"]
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name:
                continue
            raw = ws.cell(r, jcol).value
            orig = json.loads(raw)
            name_s = str(name).replace("MP位置", "HJ位置")
            ws.cell(r, 1).value = name_s
            m = NAME_RE.match(name_s.strip())
            if not m:
                raise SystemExit("unparsed %s r%s %s" % (sheet, r, name_s))
            _tag, pos, _st, rest = m.groups()
            if pos == "MP":
                pos = "HJ"
            n = int(ws.cell(r, hidx["玩家人数"]).value)
            kind = classify(rest)
            b = Builder(orig, n, pos, sheet, kind)
            b.post_blinds()
            if sheet == "PREFLOP":
                b.build_preflop_named()
            else:
                b.preflop_everyone_limp()
                if sheet in ("TURN", "RIVER"):
                    b.street_all_check("FLOP")
                if sheet == "RIVER":
                    b.street_all_check("TURN")
                target = {"FLOP": "FLOP", "TURN": "TURN", "RIVER": "RIVER"}[sheet]
                b.build_postflop_named(target)
                b.finish_to_showdown_if_needed()
            j = b.to_json()
            ws.cell(r, hidx["英雄位"]).value = pos
            ws.cell(r, hidx["动作"]).value = ws.cell(r, hidx["动作"]).value  # keep
            # map 动作 to kind chinese/english original if still MP-era — keep original col
            ws.cell(r, hidx["Hero行动次数"]).value = hero_actions_count(j)
            cell_b = ws.cell(r, hidx["大小盲"])
            cell_b.value = blinds_text(j)
            cell_b.number_format = "@"
            fill_sequence(ws, r, headers, j["actions"])
            ws.cell(r, jcol).value = json.dumps(j, ensure_ascii=False)
            n_rows += 1
    wb.save(DST)
    print("wrote", DST, "rows", n_rows)


if __name__ == "__main__":
    main()
