#!/usr/bin/env python3
"""从 MCP 专项测试用例生成「步骤级」自动化 Excel。

核心：每个原用例的操作步骤拆成一行，在步骤后追加自动化字段
（api/ui 动作、工具、指标 code、完整 POST body、断言）。

10 月 26 源表格式：
  - 表头在第 2 行；不再有「正确结果(预期)」
  - 工具 / 指标 / 数据基准都写在「预期结果」里的
    「必须调用」「【核心】」「【数据基准】」
  - 「异常截图 / mcp返回原文 / …」是人工执行记录列，转换时忽略
  - 仍兼容旧表「正确结果(预期)」列

产出 sheet：
  - readme
  - config
  - cases      用例元数据 + 可自动化标记
  - steps      ★ 一步一行（人看步骤 + 机跑字段）
  - requests   仅 API 可执行行（供 pytest 直连）
  - metrics    从源表「指标」sheet 拷贝

用法：
  python3 scripts/build_mcp_auto_cases.py \\
    --source "/Users/xixi/Downloads/MCP专项测试用例_10月26.xlsx" \\
    --out "/Users/xixi/Downloads/MCP自动化步骤用例_10月26.xlsx"
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OPEN_ID = "1ad0c14143ae4d12894c41c89e603e94"
TODAY_START = "2026-10-22 00:00:00"
TODAY_END = "2026-10-22 23:59:59"
YDAY_START = "2026-10-21 00:00:00"
YDAY_END = "2026-10-21 23:59:59"
SESSION_ID = "1792639765135"
HAND_LOSS = "1792639765135_QHTH_1792640896015"
HAND_WIN = "1792639765135_9HQH_1792640971544"
HAND_2ND_LOSS = "1792639765135_3CAC_1792640325740"
OPP_MAIN = "Keppel"
OPP_LOW = "日月明F88"
OPP_MID = "Surfer99"

TOOL_RE = re.compile(
    r"(tool_[a-z_]+|opponent_profile|list_metrics|query_metrics|run_sql|describe_tables)"
)
METRIC_CODE_RE = re.compile(r"[（(]([a-z][a-z0-9_]*)[）)]")
STEP_RE = re.compile(r"^\s*(\d+)\.\s*(.+?)\s*$")
OPEN_ID_RE = re.compile(r"open\s*ID\s*[:：]\s*([0-9a-fA-F]{16,})", re.I)
SECTION_RE = re.compile(
    r"(【核心】|【数据基准】|【断言】|【红线】|【依据】|【测试点】)\s*(.*?)(?=\n【|\n●|$)",
    re.S,
)

# 人工执行记录列：转换自动化表时不读
EXEC_RECORD_HEADERS = {
    "异常截图",
    "mcp返回原文",
    "ai提炼mcp返回结果",
    "结果：mcp返回是否和正确结果一致",
    "执行日期",
}

SKIP_GROUP_PREFIXES = (
    "3.3 ",
    "3.5 ",
    "3.6 ",
    "3.8 ",
    "3.9 ",
    "3.11",
)

SKIP_STATUS_MARKERS = ("阻塞", "暂停", "skip", "不做")


def rpc_body(tool: str, arguments: dict[str, Any], req_id: int = 3) -> str:
    return json.dumps(
        {
            "jsonrpc": "2.0",
            "id": req_id,
            "method": "tools/call",
            "params": {"name": tool, "arguments": arguments},
        },
        ensure_ascii=False,
        indent=2,
    )


def overview(start: str, end: str) -> dict[str, Any]:
    return {"open_id": "{{open_id}}", "time_start": start, "time_end": end}


def sessions(start: str, end: str, limit: int = 20) -> dict[str, Any]:
    return {
        "open_id": "{{open_id}}",
        "time_start": start,
        "time_end": end,
        "limit": limit,
    }


def find_hands(
    *,
    start: str = TODAY_START,
    end: str = TODAY_END,
    hole_cards: str = "",
    order_by: str = "profit_bb",
    sort_order: str = "desc",
    limit: int = 10,
) -> dict[str, Any]:
    return {
        "open_id": "{{open_id}}",
        "time_start": start,
        "time_end": end,
        "hole_cards": hole_cards,
        "order_by": order_by,
        "sort_order": sort_order,
        "limit": limit,
        "offset": 0,
    }


def hand_detail(hand_id: str) -> dict[str, Any]:
    return {"open_id": "{{open_id}}", "hand_id": hand_id}


def opponent(name: str = OPP_MAIN) -> dict[str, Any]:
    return {"open_id": "{{open_id}}", "opponent_name": name}


def user_metrics(
    metrics: list[str],
    *,
    start: str = TODAY_START,
    end: str = TODAY_END,
    dimensions: list[str] | None = None,
) -> dict[str, Any]:
    return {
        "open_id": "{{open_id}}",
        "time_start": start,
        "time_end": end,
        "metrics": json.dumps(metrics, ensure_ascii=False),
        "ad_hoc_json": "[]",
        "dimensions": json.dumps(dimensions or [], ensure_ascii=False),
        "filters": "{}",
        "limit": 1000,
    }


def query_metrics(metrics: list[str], session_id: str = SESSION_ID) -> dict[str, Any]:
    return {
        "session_id": session_id,
        "metrics": json.dumps(metrics, ensure_ascii=False),
        "dimensions": "[]",
        "filters": "{}",
        "ad_hoc_json": "[]",
        "page": 0,
        "page_size": 0,
    }


def list_metrics(category: str = "打法画像") -> dict[str, Any]:
    return {"category": category}


def parse_metric_codes(text: str) -> list[str]:
    codes: list[str] = []
    for m in METRIC_CODE_RE.finditer(text or ""):
        code = m.group(1)
        if code not in codes:
            codes.append(code)
    return codes


def section_map(text: str) -> dict[str, str]:
    """从预期结果 / 备注里抽出【核心】【数据基准】等段落。"""
    out: dict[str, str] = {}
    for m in SECTION_RE.finditer(text or ""):
        key = m.group(1)
        val = re.sub(r"\s+", " ", (m.group(2) or "").strip())
        if val and key not in out:
            out[key] = val
    return out


def build_assert_hints(*, expected: str, correct: str = "", remark: str = "") -> str:
    """自动化断言描述：旧表用正确结果；新表用【核心】+【数据基准】。"""
    if (correct or "").strip():
        return correct.strip()[:800]

    # 只从「预期结果」抽段落，避免备注正文粘到【数据基准】后面
    sections = section_map(expected or "")
    parts: list[str] = []
    for key in ("【核心】", "【数据基准】"):
        if key in sections:
            parts.append(f"{key}{sections[key]}")
    if parts:
        return "\n".join(parts)[:800]

    # 预期结果没有分段时，再看备注里的同名段落
    sections = section_map(remark or "")
    for key in ("【核心】", "【数据基准】"):
        if key in sections:
            parts.append(f"{key}{sections[key]}")
    if parts:
        return "\n".join(parts)[:800]

    # 去掉「必须调用」行，剩下当提示
    body = re.sub(
        r"●\s*必须调用[：:].*?(?=\n\n|\n【|$)",
        "",
        expected or "",
        flags=re.S,
    ).strip()
    return (body or expected or "")[:800]


def extract_open_id(text: str) -> str:
    m = OPEN_ID_RE.search(text or "")
    return m.group(1) if m else ""


def parse_steps(steps_text: str) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    if not steps_text:
        return out
    for line in str(steps_text).splitlines():
        m = STEP_RE.match(line)
        if m:
            out.append((int(m.group(1)), m.group(2).strip()))
        elif line.strip() and out:
            # 续行并入上一步
            n, t = out[-1]
            out[-1] = (n, f"{t} {line.strip()}")
    return out


def extract_input(step_text: str) -> str:
    m = re.search(r"输入[:：]\s*(.+)$", step_text)
    return m.group(1).strip() if m else ""


def time_from_text(text: str) -> tuple[str, str, str]:
    """返回 label, start, end。"""
    t = text or ""
    if any(x in t for x in ("昨天", "昨日", "10-21", "10 月 21", "10月21")) and not any(
        x in t for x in ("今天", "今日", "10-22", "然后再问昨天")
    ):
        return "yesterday", YDAY_START, YDAY_END
    if "昨天" in t and ("今天" in t or "今日" in t):
        return "both", TODAY_START, TODAY_END
    return "today", TODAY_START, TODAY_END


def run_mode_for(group: str, status: str = "") -> tuple[str, str]:
    st = str(status or "")
    if any(m in st for m in SKIP_STATUS_MARKERS):
        return "skip", f"源表状态={st}"
    g = str(group or "")
    for p in SKIP_GROUP_PREFIXES:
        if g.startswith(p) or p.strip() in g:
            return "skip", f"分组偏故障/协议/Agent，步骤保留但 API 默认不跑"
    if "3.10" in g:
        return "manual", "工具定义专项，人工/半自动"
    return "auto", ""


def unique_tools(tools: list[str], module: str) -> list[str]:
    out: list[str] = []
    for t in tools:
        if t not in out:
            out.append(t)
    if out:
        return out
    m = str(module or "")
    for name in (
        "tool_get_user_overview",
        "tool_list_user_sessions",
        "tool_query_user_metrics",
        "tool_find_hands",
        "tool_get_hand_detail",
        "tool_get_snowie_advice",
        "opponent_profile",
        "list_metrics",
        "query_metrics",
    ):
        if name in m:
            out.append(name)
    return out


def args_for_tool(
    tool: str,
    *,
    start: str,
    end: str,
    question: str,
    metrics: list[str],
    blob: str,
) -> dict[str, Any]:
    q = f"{question}\n{blob}"
    if tool == "tool_get_user_overview":
        return overview(start, end)
    if tool == "tool_list_user_sessions":
        return sessions(start, end)
    if tool == "tool_find_hands":
        hole = ""
        so = "desc"
        if any(x in q for x in ("带 A", "带A", "所有 A")):
            hole = "A"
        elif "对子" in q:
            hole = "22"
        elif "同花" in q:
            hole = "AKs"
        if any(x in q for x in ("输最多", "最亏", "输超过", "亏损")):
            so = "asc"
        return find_hands(start=start, end=end, hole_cards=hole, sort_order=so, limit=10)
    if tool == "tool_get_hand_detail":
        hid = HAND_LOSS
        if any(x in q for x in ("赢最多", "9h", "9H", "9HQH")):
            hid = HAND_WIN
        return hand_detail(hid)
    if tool == "opponent_profile":
        name = OPP_MAIN
        if "日月明" in q:
            name = OPP_LOW
        elif "Surfer" in q:
            name = OPP_MID
        elif "查无" in q or "没有交手" in q:
            name = "不存在的对手XYZ"
        return opponent(name)
    if tool == "tool_query_user_metrics":
        mets = metrics or ["vpip_rate", "pfr_rate", "three_bet_rate"]
        dims: list[str] = []
        if any(x in q for x in ("位置", "BTN", "CO", "HJ")):
            dims = ["by_position"]
        return user_metrics(mets, start=start, end=end, dimensions=dims or None)
    if tool == "list_metrics":
        return list_metrics("打法画像")
    if tool == "query_metrics":
        return query_metrics(metrics or ["vpip_rate"])
    if tool == "tool_get_snowie_advice":
        return {"open_id": "{{open_id}}", "hand_id": HAND_LOSS, "action_index": 1}
    return overview(start, end)


def classify_step(step_text: str) -> tuple[str, str]:
    """返回 (auto_layer, auto_action)。"""
    t = step_text
    if re.search(r"输入[:：]", t) or t.startswith("再输入"):
        return "api+ui", "input_and_call"
    if t.strip() in {"发送", "发送。"} or t.startswith("发送"):
        return "ui", "click_send"
    if "执行流程" in t:
        return "ui+assert", "inspect_tool_flow"
    if any(x in t for x in ("记下", "记录", "抄下来")):
        return "assert", "capture_value"
    if any(x in t for x in ("核对", "对照表", "打勾", "比对", "对齐", "确认", "比")):
        return "assert", "compare_baseline"
    if "点开" in t or "展开" in t:
        return "ui", "expand_detail"
    return "manual", "manual_check"


def pick_tool_for_input(
    question: str,
    tools: list[str],
    input_index: int,
    expected: str,
) -> tuple[str, str]:
    """为第 N 次「输入」挑选工具。返回 (tool, role)。"""
    q = question or ""
    # OR 关系
    if "或" in expected and len(tools) >= 2:
        return tools[0], "preferred"

    if len(tools) == 1:
        return tools[0], "preferred"

    # 多工具：按问句语义
    if any(x in q for x in ("列一下", "哪几场", "每场", "场次清单", "几点开始")):
        if "tool_list_user_sessions" in tools:
            return "tool_list_user_sessions", "preferred"
    if any(x in q for x in ("怎么打的", "牌谱", "这手")):
        if "tool_get_hand_detail" in tools:
            return "tool_get_hand_detail", "chain"
    if any(x in q for x in ("找", "带 A", "对子", "输超过", "哪些手")):
        if "tool_find_hands" in tools:
            return "tool_find_hands", "preferred"
    if any(x in q for x in ("VPIP", "PFR", "3bet", "CBet", "指标")):
        if "tool_query_user_metrics" in tools:
            return "tool_query_user_metrics", "preferred"
        if "query_metrics" in tools:
            return "query_metrics", "preferred"
    if any(x in q for x in ("对手", "风格", "Keppel")):
        if "opponent_profile" in tools:
            return "opponent_profile", "preferred"

    # 按出现顺序轮询
    if input_index < len(tools):
        return tools[input_index], "preferred"
    return (tools[0] if tools else "tool_get_user_overview"), "preferred"


def style_header(ws, color: str = "1F4E79") -> None:
    fill = PatternFill("solid", fgColor=color)
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col[:60]:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value).split("\n")[0]) + 2))
        ws.column_dimensions[letter].width = width


def load_source(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], str]:
    """返回 (cases, metrics_rows, default_open_id)。"""
    wb = load_workbook(path, data_only=True)
    # 兼容表头在第 1 或第 2 行
    ws = wb["MCP专项"]
    header_row = 1
    if ws.cell(1, 1).value != "编号" and ws.cell(2, 1).value == "编号":
        header_row = 2
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    col = {
        str(h): i + 1
        for i, h in enumerate(headers)
        if h and str(h) not in EXEC_RECORD_HEADERS
    }

    def cell(r: int, name: str) -> Any:
        idx = col.get(name)
        return ws.cell(r, idx).value if idx else None

    cases: list[dict[str, Any]] = []
    default_open_id = ""
    for r in range(header_row + 1, ws.max_row + 1):
        cid = cell(r, "编号")
        if not cid:
            continue
        expected = str(cell(r, "预期结果") or "")
        # 旧表有「正确结果(预期)」；新表没有，断言描述从预期结果抽
        correct = str(cell(r, "正确结果(预期)") or "")
        remark = str(cell(r, "备注") or "")
        module = str(cell(r, "所属模块") or "")
        pre = str(cell(r, "前置条件") or "")
        hint_blob = build_assert_hints(expected=expected, correct=correct, remark=remark)
        tools = TOOL_RE.findall(expected + "\n" + correct + "\n" + module)
        metrics = parse_metric_codes(hint_blob + "\n" + expected + "\n" + correct)
        if not default_open_id:
            default_open_id = extract_open_id(pre)
        cases.append(
            {
                "case_id": str(cid).strip(),
                "group": cell(r, "分组") or "",
                "title": cell(r, "标题") or "",
                "module": module,
                "priority": cell(r, "优先级") or "",
                "status": cell(r, "状态") or "",
                "correct": correct,
                "expected": expected,
                "assert_hints": hint_blob,
                "pre": pre,
                "steps_text": cell(r, "操作步骤") or "",
                "remark": remark,
                "tools": unique_tools(tools, module),
                "metric_codes": metrics,
            }
        )

    metrics_rows: list[dict[str, Any]] = []
    if "指标" in wb.sheetnames:
        mws = wb["指标"]
        mh = [mws.cell(1, c).value for c in range(1, mws.max_column + 1)]
        for r in range(2, mws.max_row + 1):
            item = {str(mh[i]): mws.cell(r, i + 1).value for i in range(len(mh))}
            if item.get("指标 code") or item.get("工具英文名"):
                metrics_rows.append(item)
    return cases, metrics_rows, default_open_id or OPEN_ID


def build_step_rows(case: dict[str, Any]) -> list[dict[str, Any]]:
    """把一个 case 的操作步骤展开为带自动化字段的行。"""
    mode, skip_reason = run_mode_for(case["group"], case.get("status", ""))
    steps = parse_steps(str(case["steps_text"]))
    tools = case["tools"]
    metrics = case["metric_codes"]
    hints = str(case.get("assert_hints") or case.get("correct") or "")
    blob = f"{case['title']}\n{hints}\n{case['expected']}"
    rows: list[dict[str, Any]] = []
    input_idx = 0
    api_seq = 0

    if not steps:
        steps = [(1, "（原表无结构化步骤）")]

    for step_no, step_text in steps:
        layer, action = classify_step(step_text)
        question = extract_input(step_text)
        label, t0, t1 = time_from_text(question or step_text or blob)

        tool_name = ""
        role = ""
        body = ""
        assert_expr = ""
        metric_codes = ""
        enabled = "N"
        api_layer = ""

        if action == "input_and_call" and question:
            input_idx += 1
            # both → 第一次今天，若问句含昨天且是第二次则昨天
            if label == "both":
                if "昨天" in question and "今天" not in question:
                    t0, t1 = YDAY_START, YDAY_END
                    label = "yesterday"
                elif input_idx >= 2 and "昨天" in (question + blob):
                    t0, t1 = YDAY_START, YDAY_END
                    label = "yesterday"
                else:
                    t0, t1 = TODAY_START, TODAY_END
                    label = "today"

            tool_name, role = pick_tool_for_input(
                question, tools, input_idx - 1, case["expected"]
            )
            # 第二次输入且工具列表含 list_sessions
            if input_idx >= 2 and "tool_list_user_sessions" in tools and (
                "场次" in question or "列" in question
            ):
                tool_name = "tool_list_user_sessions"
                role = "preferred"

            args = args_for_tool(
                tool_name,
                start=t0,
                end=t1,
                question=question,
                metrics=metrics,
                blob=blob,
            )
            api_seq += 1
            body = rpc_body(tool_name, args, req_id=2 + api_seq)
            metric_codes = ",".join(metrics)
            enabled = "Y" if mode == "auto" else "N"
            api_layer = "api"
            # 按真实 MCP 顶层字段断言，避免把 A 工具字段套到 B 工具上
            if tool_name == "tool_get_user_overview":
                if label == "yesterday":
                    assert_expr = "session_count==0"
                elif any(x in metrics for x in ("hands", "total_hands")) or "手数" in hints:
                    assert_expr = "hands"
                elif any(x in metrics for x in ("profit_bb", "total_profit_bb")):
                    assert_expr = "profit_bb"
                else:
                    assert_expr = "session_count"
            elif tool_name == "tool_list_user_sessions":
                assert_expr = "sessions" if "sessions" in hints or "场次列表" in hints else "session_count"
            elif tool_name == "tool_find_hands":
                assert_expr = "hands"
            elif tool_name == "tool_get_hand_detail":
                assert_expr = "hand"
            elif tool_name == "opponent_profile":
                if "hands_vs_hero" in metrics or "交手" in hints:
                    assert_expr = "hands_vs_hero"
                else:
                    assert_expr = "vpip"
            elif tool_name in {"tool_query_user_metrics", "query_metrics"}:
                assert_expr = "data"
            elif metrics:
                # 只在指标码本身就是该工具常见顶层字段时才用
                top_ok = {
                    "session_count",
                    "sessions",
                    "hands",
                    "matched_count",
                    "profit_bb",
                    "hand",
                    "vpip",
                    "data",
                }
                assert_expr = metrics[0] if metrics[0] in top_ok else ""

        elif action == "inspect_tool_flow":
            tool_name = "|".join(tools)
            role = "assert_tool"
            assert_expr = "tool_called:" + (tools[0] if tools else "")
            enabled = "Y" if mode == "auto" else "N"
            api_layer = "ui_assert"

        elif action in {"capture_value", "compare_baseline"}:
            tool_name = "|".join(tools)
            role = "assert"
            metric_codes = ",".join(metrics)
            assert_expr = "compare_hints"
            enabled = "Y" if mode == "auto" else "N"
            api_layer = "assert"

        elif action == "click_send":
            # 合并语义：发送本身不单独调 API（已在 input 步）
            api_layer = "ui"
            enabled = "N"
            role = "ui_only"

        else:
            api_layer = "manual"
            enabled = "N"
            role = "manual"

        rows.append(
            {
                "case_id": case["case_id"],
                "group": case["group"],
                "title": case["title"],
                "priority": str(case["priority"]),
                "run_mode": mode,
                "skip_reason": skip_reason,
                "step_no": step_no,
                "step_text": step_text,
                # —— 以下为步骤后追加的自动化字段 ——
                "auto_layer": api_layer or layer,
                "auto_action": action,
                "user_input": question,
                "tool_name": tool_name,
                "role": role,
                "metric_codes": metric_codes,
                "time_label": label if question else "",
                "time_start": t0 if question else "",
                "time_end": t1 if question else "",
                "enabled": enabled,
                "request_body": body,
                "assert_expr": assert_expr,
                "assert_hints": hints[:800],
                "expect_tools": "|".join(tools),
                "notes": case["remark"] or "",
            }
        )
    return rows


def write_excel(
    cases: list[dict[str, Any]],
    metrics_rows: list[dict[str, Any]],
    out: Path,
    *,
    default_open_id: str = OPEN_ID,
    source_label: str = "10月26 源表",
) -> dict[str, int]:
    wb = Workbook()

    # readme
    ws0 = wb.active
    ws0.title = "readme"
    for row in [
        [f"MCP 自动化步骤用例（{source_label}）"],
        [],
        ["怎么看"],
        ["1. steps 表：每个「操作步骤」一行；右侧灰/紫列是自动化字段"],
        ["2. auto_action=input_and_call 的行带完整 request_body，可直接 POST"],
        ["3. requests 表：只收集 enabled=Y 且有 request_body 的 API 行，给 pytest 用"],
        ["4. cases 表：用例级 run_mode；skip 的分组步骤仍保留便于人工"],
        ["5. metrics 表：工具↔指标 code 对照（来自源表）"],
        ["6. assert_hints 来自源表「预期结果」的【核心】【数据基准】（旧表则用正确结果列）"],
        [],
        ["占位符"],
        ["{{open_id}}", "脚本替换为 config.default_open_id"],
        ["{{hand_id}}", "链式步骤从上一步 extract"],
    ]:
        ws0.append(row)
    ws0["A1"].font = Font(bold=True, size=14)

    # config
    ws_cfg = wb.create_sheet("config")
    ws_cfg.append(["key", "value", "note"])
    for row in [
        ("base_url", "https://d943-223-70-117-10.ngrok-free.app/mcp", "跑前改"),
        ("token_env", "MCP_TOKEN", "Bearer"),
        ("default_open_id", default_open_id or OPEN_ID, "源表前置条件 openID"),
        ("day_today_start", TODAY_START, ""),
        ("day_today_end", TODAY_END, ""),
        ("day_yesterday_start", YDAY_START, ""),
        ("day_yesterday_end", YDAY_END, ""),
        ("default_session_id", SESSION_ID, ""),
        ("sample_hand_loss", HAND_LOSS, ""),
        ("sample_opponent", OPP_MAIN, ""),
    ]:
        ws_cfg.append(list(row))
    style_header(ws_cfg, "2E7D32")
    autosize(ws_cfg)

    # cases
    ws_cases = wb.create_sheet("cases")
    case_headers = [
        "case_id",
        "group",
        "title",
        "priority",
        "run_mode",
        "skip_reason",
        "expect_tools",
        "metric_codes",
        "user_question_first",
        "assert_hints",
        "source_module",
        "precondition",
        "remark",
        "step_count",
        "api_step_count",
    ]
    ws_cases.append(case_headers)

    # steps ★
    ws_steps = wb.create_sheet("steps")
    step_headers = [
        "case_id",
        "group",
        "title",
        "priority",
        "run_mode",
        "step_no",
        "step_text",
        # automation columns after step
        "auto_layer",
        "auto_action",
        "user_input",
        "tool_name",
        "role",
        "metric_codes",
        "time_label",
        "time_start",
        "time_end",
        "enabled",
        "request_body",
        "assert_expr",
        "assert_hints",
        "expect_tools",
        "notes",
        "skip_reason",
    ]
    ws_steps.append(step_headers)

    # requests
    ws_req = wb.create_sheet("requests")
    req_headers = [
        "case_id",
        "step_no",
        "enabled",
        "tool_name",
        "role",
        "method",
        "request_body",
        "metric_codes",
        "assert_type",
        "assert_expr",
        "assert_hints",
        "user_question",
        "title",
        "priority",
        "group",
        "notes",
    ]
    ws_req.append(req_headers)

    stats = {
        "cases": 0,
        "auto": 0,
        "skip": 0,
        "manual": 0,
        "step_rows": 0,
        "api_rows": 0,
    }

    for case in cases:
        mode, reason = run_mode_for(case["group"], case.get("status", ""))
        step_rows = build_step_rows(case)
        api_n = sum(1 for s in step_rows if s["request_body"] and s["enabled"] == "Y")
        first_q = next((s["user_input"] for s in step_rows if s["user_input"]), "")
        hints = str(case.get("assert_hints") or case.get("correct") or "")

        ws_cases.append(
            [
                case["case_id"],
                case["group"],
                case["title"],
                str(case["priority"]),
                mode,
                reason,
                "|".join(case["tools"]),
                ",".join(case["metric_codes"]),
                first_q,
                hints[:800],
                case["module"],
                case["pre"],
                case["remark"],
                len(step_rows),
                api_n,
            ]
        )
        stats["cases"] += 1
        stats[mode] = stats.get(mode, 0) + 1

        for s in step_rows:
            ws_steps.append([s.get(h, "") for h in step_headers])
            stats["step_rows"] += 1
            if s["request_body"] and s["enabled"] == "Y":
                ws_req.append(
                    [
                        s["case_id"],
                        s["step_no"],
                        s["enabled"],
                        s["tool_name"],
                        s["role"],
                        "tools/call",
                        s["request_body"],
                        s["metric_codes"],
                        "expr",
                        s["assert_expr"],
                        s["assert_hints"],
                        s["user_input"],
                        s["title"],
                        s["priority"],
                        s["group"],
                        s["notes"],
                    ]
                )
                stats["api_rows"] += 1

    style_header(ws_cases)
    style_header(ws_steps, "6A1B9A")
    style_header(ws_req, "4527A0")
    autosize(ws_cases, 36)
    autosize(ws_steps, 28)
    ws_steps.column_dimensions["G"].width = 42  # step_text
    ws_steps.column_dimensions["R"].width = 48  # request_body
    for row in ws_steps.iter_rows(min_row=2, min_col=18, max_col=18):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # metrics
    ws_m = wb.create_sheet("metrics")
    if metrics_rows:
        keys = list(metrics_rows[0].keys())
        ws_m.append(keys)
        for item in metrics_rows:
            ws_m.append([item.get(k) for k in keys])
        style_header(ws_m, "00695C")
        autosize(ws_m, 28)
    else:
        ws_m.append(["提示", "源表无「指标」sheet"])

    ws0.append([])
    ws0.append(["生成统计"])
    for k, v in stats.items():
        ws0.append([k, v])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return stats


def convert_workbook(
    source: Path,
    out: Path,
    *,
    also_copy_to: Path | None = None,
) -> dict[str, int]:
    """把原始专项用例表转换成自动化 steps/requests Excel。可被 CLI / Web 复用。"""
    source = Path(source)
    out = Path(out)
    if not source.exists():
        raise FileNotFoundError(f"找不到源表: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx / .xlsm")
    cases, metrics_rows, open_id = load_source(source)
    if not cases:
        raise ValueError("未解析到任何用例（请确认含「MCP专项」sheet 且有编号列）")
    label = source.stem
    stats = write_excel(
        cases,
        metrics_rows,
        out,
        default_open_id=open_id,
        source_label=label,
    )
    if also_copy_to is not None:
        also_copy_to = Path(also_copy_to)
        also_copy_to.parent.mkdir(parents=True, exist_ok=True)
        write_excel(
            cases,
            metrics_rows,
            also_copy_to,
            default_open_id=open_id,
            source_label=label,
        )
    stats["default_open_id"] = open_id  # type: ignore[assignment]
    return stats


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--source",
        type=Path,
        default=Path("/Users/xixi/Downloads/MCP专项测试用例_10月26.xlsx"),
    )
    ap.add_argument(
        "--out",
        type=Path,
        default=Path("/Users/xixi/Downloads/MCP自动化步骤用例_10月26.xlsx"),
    )
    args = ap.parse_args()
    repo = Path(__file__).resolve().parents[1]
    alt = repo / "testdata" / "mcp" / args.out.name
    latest = repo / "automation" / "data" / "MCP自动化步骤用例_latest.xlsx"
    cases, metrics_rows, open_id = load_source(args.source)
    if not cases:
        raise SystemExit("未解析到任何用例")
    label = args.source.stem
    stats = write_excel(
        cases,
        metrics_rows,
        args.out,
        default_open_id=open_id,
        source_label=label,
    )
    for dest in (alt, latest):
        dest.parent.mkdir(parents=True, exist_ok=True)
        write_excel(
            cases,
            metrics_rows,
            dest,
            default_open_id=open_id,
            source_label=label,
        )
    print(f"wrote {args.out}")
    print(f"wrote {alt}")
    print(f"wrote {latest}")
    print(f"default_open_id={open_id}")
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
