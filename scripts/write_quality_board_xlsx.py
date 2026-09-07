#!/usr/bin/env python3
"""Readable, frequently-edited weekly quality board (Excel)."""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_weekly_quality_board import (  # noqa: E402
    BILL_FILL,
    BODY_FONT,
    EDIT_FILL,
    FACT_FILL,
    FAIL_FILL,
    GROUP_FILL,
    HINT_FONT,
    KPI_FILL,
    MODEL_FILL,
    PAGE_BANDS,
    PASS_FILL,
    ROOT,
    SKIP_FILL,
    THIN,
    TODO_FILL,
    collect_field_rows,
    load_cases,
    load_prd_fields,
    pages_missing_model,
    style_header,
    unique_metrics,
)

OUT = Path("/Users/xixi/Downloads/质量看板.xlsx")
OUT_PROJ = ROOT / "output/质量看板.xlsx"

BOLD = Font(name="Microsoft YaHei", bold=True, size=11)
KPI_FONT = Font(name="Microsoft YaHei", size=16, bold=True)
MID = Alignment(vertical="center")
MID_C = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="center", wrap_text=True)


def _title(cell, text: str, size: int = 14) -> None:
    cell.value = text
    cell.font = Font(name="Microsoft YaHei", size=size, bold=True)


def _input(cell, value="") -> None:
    cell.value = value if value != "" else None
    cell.fill = EDIT_FILL
    cell.alignment = MID_C
    cell.font = BODY_FONT
    cell.border = THIN


def _border_row(ws, r: int, cols: int) -> None:
    for c in range(1, cols + 1):
        ws.cell(r, c).border = THIN
        if ws.cell(r, c).font.name != "Microsoft YaHei":
            ws.cell(r, c).font = BODY_FONT


def _cf_equals(ws, rng: str, value: str, fill: PatternFill) -> None:
    ws.conditional_formatting.add(
        rng,
        CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill),
    )


def _widths(ws, mapping: dict[int, int]) -> None:
    for i, w in mapping.items():
        ws.column_dimensions[get_column_letter(i)].width = w


def _kpi_label(cell, text: str) -> None:
    cell.value = text
    cell.font = Font(name="Microsoft YaHei", bold=True, size=12)
    cell.alignment = MID_C
    cell.border = THIN


def _kpi_value(cell, formula: str, pct: bool = True) -> None:
    cell.value = formula
    cell.number_format = "0.0%" if pct else "0.0"
    cell.font = KPI_FONT
    cell.fill = KPI_FILL
    cell.alignment = MID_C
    cell.border = THIN


def _rate(ok: str, bad: str, empty: str) -> str:
    return f'=IF({ok}+{bad}=0,"{empty}",{ok}/({ok}+{bad}))'


def build(dest: Path) -> None:
    cases = load_cases()
    prd_fields = load_prd_fields()
    fact_rows, model_rows, bill_rows, _by_page = collect_field_rows(
        prd_fields, unique_metrics()
    )
    missing = pages_missing_model(prd_fields, model_rows)
    if missing:
        raise SystemExit("模型结论漏页: " + "、".join(f"{a} {b}" for a, b in missing))
    slices = [c for c in cases if c["module"] == "切局模块测试"]
    all_fact = sorted(
        fact_rows + bill_rows,
        key=lambda r: (r["page_id"], 0 if r["kind"] == "牌局事实" else 1, r["field"]),
    )
    by_page_rows: dict[str, list[dict]] = defaultdict(list)
    page_order: list[str] = []
    for row in all_fact:
        key = f'{row["page_id"]} {row["page_name"]}'
        if key not in by_page_rows:
            page_order.append(key)
        by_page_rows[key].append(row)
    model_by_page: dict[str, int] = defaultdict(int)
    for row in model_rows:
        model_by_page[f'{row["page_id"]} {row["page_name"]}'] += 1

    wb = Workbook()

    # —— 事实数据（字段在上，按页汇总在最底下）——
    ws3 = wb.active
    ws3.title = "事实数据"
    ws3.sheet_properties.tabColor = "166534"
    n_pages = len(page_order)
    field_start = 4
    last_field = field_start + len(all_fact) - 1
    fa, fb, fc = (
        f"A{field_start}:A{last_field}",
        f"B{field_start}:B{last_field}",
        f"C{field_start}:C{last_field}",
    )
    _title(ws3["A1"], "事实数据  ·  只列页面上的数据；按页汇总在最底下")
    ws3.merge_cells("A1:D1")
    _kpi_label(ws3["E1"], "M3")
    _kpi_value(ws3["F1"], f'=IF(COUNTA({fa})=0,"待标",COUNTIF({fa},"已实现")/COUNTA({fa}))')
    _kpi_label(ws3["G1"], "M4")
    ok = f'COUNTIFS({fa},"已实现",{fb},"对")'
    bad = f'COUNTIFS({fa},"已实现",{fb},"错")'
    _kpi_value(ws3["H1"], _rate(ok, bad, "待核"))
    ws3["A2"] = (
        "覆盖：已实现/未实现。准确：对/错/无真值。全部都要核。"
        "归类：牌局事实 / 非牌局业务字段。P12 已计入。筛「页面」一次改一页。拉到表底看按页汇总。"
    )
    ws3["A2"].font = HINT_FONT
    ws3.merge_cells("A2:H2")
    ws3.row_dimensions[2].height = 32

    for i, h in enumerate(["覆盖", "准确", "页面", "字段", "归类", "来源"], 1):
        ws3.cell(3, i, h)
    style_header(ws3, 3, 6)

    dv_cov = DataValidation(type="list", formula1='"已实现,未实现"', allow_blank=True)
    dv_acc = DataValidation(type="list", formula1='"对,错,无真值"', allow_blank=True)
    ws3.add_data_validation(dv_cov)
    ws3.add_data_validation(dv_acc)

    r = field_start
    prev = None
    band_i = -1
    for row in all_fact:
        page = f'{row["page_id"]} {row["page_name"]}'
        if page != prev:
            band_i = (band_i + 1) % 2
            prev = page
        band = PAGE_BANDS[band_i]
        _input(ws3.cell(r, 1))
        dv_cov.add(ws3.cell(r, 1))
        _input(ws3.cell(r, 2))
        dv_acc.add(ws3.cell(r, 2))
        ws3.cell(r, 3, page)
        ws3.cell(r, 4, row["field"]).alignment = WRAP
        ws3.cell(r, 5, row["kind"])
        ws3.cell(r, 6, row.get("origin") or "")
        ws3.cell(r, 6).font = HINT_FONT
        for c in range(3, 7):
            ws3.cell(r, c).fill = band
            ws3.cell(r, c).alignment = MID if c != 4 else WRAP
        if row["kind"] == "牌局事实":
            ws3.cell(r, 5).fill = FACT_FILL
        else:
            ws3.cell(r, 5).fill = BILL_FILL
        _border_row(ws3, r, 6)
        ws3.row_dimensions[r].height = 22
        r += 1

    _cf_equals(ws3, f"A{field_start}:A{last_field}", "已实现", PASS_FILL)
    _cf_equals(ws3, f"A{field_start}:A{last_field}", "未实现", SKIP_FILL)
    _cf_equals(ws3, f"B{field_start}:B{last_field}", "对", PASS_FILL)
    _cf_equals(ws3, f"B{field_start}:B{last_field}", "错", FAIL_FILL)
    ws3.auto_filter.ref = f"A3:F{last_field}"
    ws3.freeze_panes = "C4"

    sum_title = last_field + 2
    sum_header = last_field + 3
    sum_start = last_field + 4
    ws3.cell(sum_title, 1, "按页汇总（公式，不用手改）")
    ws3.cell(sum_title, 1).font = BOLD
    ws3.merge_cells(start_row=sum_title, start_column=1, end_row=sum_title, end_column=9)
    for i, h in enumerate(
        ["页面", "字段数", "已标覆盖", "已实现", "覆盖率", "已核准确", "对", "准确率", "模型字段"],
        1,
    ):
        ws3.cell(sum_header, i, h)
    style_header(ws3, sum_header, 9)
    for i, page in enumerate(page_order):
        rr = sum_start + i
        ws3.cell(rr, 1, page)
        ws3.cell(rr, 2, f"=COUNTIF({fc},A{rr})")
        ws3.cell(rr, 3, f'=COUNTIFS({fc},A{rr},{fa},"<>")')
        ws3.cell(rr, 4, f'=COUNTIFS({fc},A{rr},{fa},"已实现")')
        ws3.cell(rr, 5, f'=IF(B{rr}=0,"—",D{rr}/B{rr})')
        ws3.cell(rr, 5).number_format = "0.0%"
        ws3.cell(
            rr,
            6,
            f'=COUNTIFS({fc},A{rr},{fa},"已实现",{fb},"对")'
            f'+COUNTIFS({fc},A{rr},{fa},"已实现",{fb},"错")',
        )
        ws3.cell(rr, 7, f'=COUNTIFS({fc},A{rr},{fa},"已实现",{fb},"对")')
        ws3.cell(rr, 8, f'=IF(F{rr}=0,"—",G{rr}/F{rr})')
        ws3.cell(rr, 8).number_format = "0.0%"
        ws3.cell(rr, 9, f"=COUNTIF(模型结论!C:C,A{rr})")
        _border_row(ws3, rr, 9)
        ws3.row_dimensions[rr].height = 22
        ws3.cell(rr, 1).alignment = MID

    _widths(ws3, {1: 12, 2: 10, 3: 22, 4: 50, 5: 16, 6: 36, 7: 10, 8: 12, 9: 12})

    # —— 切局 ——
    ws2 = wb.create_sheet("切局")
    ws2.sheet_properties.tabColor = "B45309"
    _title(ws2["A1"], "M2 切局准确率  ·  通过了就是通过，按通过率算")
    ws2.merge_cells("A1:C1")
    _kpi_label(ws2["D1"], "已执行")
    ws2["E1"] = '=COUNTIF(A4:A500,"通过")+COUNTIF(A4:A500,"失败")+COUNTIF(A4:A500,"跳过")'
    ws2["E1"].alignment = MID_C
    ws2["E1"].border = THIN
    ws2["E1"].font = BODY_FONT
    _kpi_label(ws2["F1"], "M2")
    _kpi_value(
        ws2["H1"],
        '=IF(E1=0,"—",COUNTIF(A4:A500,"通过")/E1)',
    )
    ws2["A2"] = "结果填：通过 / 失败 / 跳过。M2 = 通过 / 已执行（不含未执行）。黄底可改。"
    ws2["A2"].font = HINT_FONT
    ws2.merge_cells("A2:H2")
    for i, h in enumerate(["结果", "用例ID", "标题", "场景"], 1):
        ws2.cell(3, i, h)
    style_header(ws2, 3, 4)
    dv_st = DataValidation(type="list", formula1='"通过,失败,跳过,未执行"', allow_blank=True)
    ws2.add_data_validation(dv_st)
    start = 4
    for i, s in enumerate(slices):
        rr = start + i
        _input(ws2.cell(rr, 1), s["status"])
        dv_st.add(ws2.cell(rr, 1))
        ws2.cell(rr, 2, s["id"])
        ws2.cell(rr, 3, s["title"]).alignment = WRAP
        ws2.cell(rr, 4, s["func"])
        _border_row(ws2, rr, 4)
        ws2.row_dimensions[rr].height = 28
    last2 = start + len(slices) - 1
    spare2 = last2 + 30
    for rr in range(last2 + 1, spare2 + 1):
        _input(ws2.cell(rr, 1))
        dv_st.add(ws2.cell(rr, 1))
        _border_row(ws2, rr, 4)
        ws2.row_dimensions[rr].height = 28
    last2 = spare2
    _cf_equals(ws2, "A4:A500", "通过", PASS_FILL)
    _cf_equals(ws2, "A4:A500", "失败", FAIL_FILL)
    _cf_equals(ws2, "A4:A500", "跳过", SKIP_FILL)
    _widths(ws2, {1: 12, 2: 24, 3: 56, 4: 22})
    ws2.freeze_panes = "A4"
    ws2.auto_filter.ref = f"A3:D{last2}"

    # —— 模型 ——
    ws4 = wb.create_sheet("模型结论")
    ws4.sheet_properties.tabColor = "3730A3"
    _title(ws4["A1"], "M5 模型数据覆盖率  ·  一层 页面/字段，只改覆盖，无准确率")
    ws4.merge_cells("A1:B1")
    _kpi_label(ws4["C1"], "M5")
    _kpi_value(ws4["D1"], '=IF(COUNTA(A4:A200)=0,"待标",COUNTIF(A4:A200,"已实现")/COUNTA(A4:A200))')
    ws4["A2"] = "覆盖填：已实现 / 未实现。无准确率。中间插行即可。"
    ws4["A2"].font = HINT_FONT
    for i, h in enumerate(["覆盖", "页面/字段", "页面", "来源"], 1):
        ws4.cell(3, i, h)
    style_header(ws4, 3, 4)
    dv4 = DataValidation(type="list", formula1='"已实现,未实现"', allow_blank=True)
    ws4.add_data_validation(dv4)
    start4 = 4
    for i, row in enumerate(model_rows):
        rr = start4 + i
        _input(ws4.cell(rr, 1))
        dv4.add(ws4.cell(rr, 1))
        ws4.cell(rr, 2, f'{row["page_id"]} {row["page_name"]} / {row["field"]}').alignment = WRAP
        ws4.cell(rr, 2).fill = MODEL_FILL
        ws4.cell(rr, 3, f'{row["page_id"]} {row["page_name"]}')
        ws4.cell(rr, 4, row.get("origin") or "").font = HINT_FONT
        _border_row(ws4, rr, 4)
        ws4.row_dimensions[rr].height = 24
    last4 = start4 + len(model_rows) - 1
    _cf_equals(ws4, "A4:A200", "已实现", PASS_FILL)
    _cf_equals(ws4, "A4:A200", "未实现", SKIP_FILL)
    _widths(ws4, {1: 12, 2: 64, 3: 22, 4: 36})
    ws4.freeze_panes = "A4"
    ws4.auto_filter.ref = f"A3:D{last4}"

    # —— 耗时并发 ——
    ws5 = wb.create_sheet("耗时并发")
    ws5.sheet_properties.tabColor = "0F172A"
    _title(ws5["A1"], "M6 报告耗时 · M8 用户并发")
    ws5.merge_cells("A1:E1")
    ws5["A2"] = "耗时：每周线上「停止录屏 → 报告可打开」。并发先手工填，后期压测替换。"
    ws5["A2"].font = HINT_FONT
    for i, h in enumerate(["编号", "指标", "本周值", "单位", "建议目标"], 1):
        ws5.cell(3, i, h)
    style_header(ws5, 3, 5)
    lat = [
        ("M6", "报告耗时", "分钟", "≤10 分钟"),
        ("M8", "用户并发数", "个", "20"),
    ]
    for i, (code, name, unit, target) in enumerate(lat, 4):
        ws5.cell(i, 1, code)
        ws5.cell(i, 2, name)
        _input(ws5.cell(i, 3))
        ws5.cell(i, 3).number_format = "0.0"
        ws5.cell(i, 4, unit)
        ws5.cell(i, 5, target).font = HINT_FONT
        _border_row(ws5, i, 5)
        ws5.row_dimensions[i].height = 26
    ws5["C5"].number_format = "0"
    _widths(ws5, {1: 10, 2: 36, 3: 12, 4: 10, 5: 28})
    ws5.freeze_panes = "A4"

    # —— 总览（按 002 的结构）——
    ws0 = wb.create_sheet("总览", 0)
    ws0.sheet_properties.tabColor = "1F4E79"
    _title(ws0["A1"], "每周质量看板", 18)
    ws0.merge_cells("A1:D1")
    ws0.row_dimensions[1].height = 28
    for i, h in enumerate(["编号", "指标", "建议目标", "0904"], 1):
        ws0.cell(3, i, h)
    style_header(ws0, 3, 4)

    def _ov_row(r, code, name, target, value, pct=True, indent=False):
        ws0.cell(r, 1, code)
        ws0.cell(r, 2, name)
        ws0.cell(r, 3, target)
        ws0.cell(r, 4, value)
        for c in range(1, 5):
            ws0.cell(r, c).font = BODY_FONT
            ws0.cell(r, c).border = THIN
            ws0.cell(r, c).alignment = WRAP
        ws0.cell(r, 4).font = Font(name="Microsoft YaHei", size=14, bold=True)
        ws0.cell(r, 4).fill = KPI_FILL
        ws0.row_dimensions[r].height = 28
        if indent:
            ws0.cell(r, 2).alignment = MID_C
        if pct:
            ws0.cell(r, 4).number_format = "0.0%"
        else:
            ws0.cell(r, 4).number_format = "0.0"

    _ov_row(4, "M2", "切局准确率", "≥95%", "=切局!H1")
    _ov_row(5, "M3", "事实数据覆盖率", 1, "=事实数据!F1")
    ws0["C5"].number_format = "0%"
    _ov_row(6, "M4", "事实数据准确率", "≥90%", "=事实数据!H1")
    _ov_row(7, "M5", "模型数据覆盖率", 1, "=模型结论!D1")
    ws0["C7"].number_format = "0%"
    _ov_row(8, "M6", "报告耗时", "≤10 分钟", "=耗时并发!C4", pct=False)
    _ov_row(9, "M8", "用户并发数", 20, "=耗时并发!C5", pct=False)
    ws0["D9"].number_format = "0"
    _widths(ws0, {1: 10, 2: 26, 3: 16, 4: 16})
    ws0.freeze_panes = "A4"

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(
        "wrote",
        dest,
        "sheets",
        wb.sheetnames,
        "pages",
        n_pages,
        "fields",
        len(all_fact),
        "field_start",
        field_start,
        "model",
        len(model_rows),
        "model_pages",
        sorted({r["page_id"] for r in model_rows}),
    )


if __name__ == "__main__":
    build(OUT)
    OUT_PROJ.write_bytes(OUT.read_bytes())
    print("copied", OUT_PROJ)
