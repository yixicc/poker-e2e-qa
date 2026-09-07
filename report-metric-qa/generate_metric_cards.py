#!/usr/bin/env python3
"""按指标字典一行一张对照卡，写出「指标用例」Excel。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from generate_test_points import (
    CONFLICT,
    OPPORTUNITY_RATES,
    SRC,
    calc_type,
    json_path,
    load_source,
)
from metric_card_expect import E, expect_for
from metric_fixtures import FIXTURES, fixture_for

ROOT = Path(__file__).resolve().parent

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
TITLE_FONT = Font(bold=True, name="Microsoft YaHei", size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
P0_FILL = PatternFill("solid", fgColor="FCE4D6")
P1_FILL = PatternFill("solid", fgColor="DEEBF7")
CONFLICT_FILL = PatternFill("solid", fgColor="F8CBAD")


def bullets(items: list[str]) -> str:
    return "\n".join(f"· {x}" for x in items if x)


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_intro(ws: Worksheet) -> None:
    ws.title = "00_怎么读"
    ws["A1"] = "指标用例（对照卡）· 一行一个指标，顺序与指标字典相同"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    rows = [
        ("读法", "不要从公式开始。先看有效素材，再看无效素材，再用「手算期望」对字段。"),
        ("和字典的关系", f"口径只认 {SRC.name} 的 01_指标字典。本表不改口径，只写怎么验。"),
        ("主卡", "Sheet 01：123 张。每张 = 计入 + 排除。P0 橙底，P1 蓝底。"),
        ("边界", "Sheet 02：只写分母为 0、源数据缺失、口径冲突。主卡通过不能代替这三行。"),
        ("数字从哪来", "默认用手算场，见下。换了录像先改场再改期望，不要空改数字。"),
        ("场A", "用户入座，时长 12:30=750秒，8 手。参与6、翻前弃3、进翻后5、摊牌3、VPIP5、PFR2、筹码−120、BB−0.1。H4 只交盲但有亏损，按清单算参与、不算 VPIP。"),
        ("场B / 场C", "场B=围观 20 分钟；场C=无关视频。都不得当作用户本场战绩。"),
        ("场D", "10 个 GTO 点：符合4、可接受2、高风险2、信息不足2。匹配度 50%。"),
        ("场E", "机会分母场。例：面对 open 4 次里 3bet 1 次 → 25%。SB 被加 2 次全 fold → 0%（有分母的 0，不是缺分母）。"),
        ("场F～I", "位置名 / 对手 / 白名单 / 连赢连亏。细则见同目录 sample_sessions.md。"),
        ("怎么跑", "按有效素材生成报告 → 打开「看哪个字段」→ 对「手算期望」。不做 App 点击验收。"),
        ("失败怎么记", "对不上手算、围观被算进去、或等于 Skill 而不是字典，记失败。"),
    ]
    ws.cell(2, 1, "章节")
    ws.cell(2, 2, "说明")
    style_header(ws, 2, 2)
    for i, (a, b) in enumerate(rows, 3):
        ws.cell(i, 1, a).alignment = WRAP
        ws.cell(i, 2, b).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 36
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 118
    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 24


MAIN_COLS = [
    "指标编码",
    "中文名称",
    "优先级",
    "数据层级",
    "是否用户可见",
    "含义（字典原文）",
    "统计方法 / 口径（字典原文）",
    "有效素材（必须准备、必须计入）",
    "无效素材（必须准备、必须排除）",
    "用哪一场 / 怎么配",
    "手算期望",
    "怎样算失败",
    "看哪个字段",
    "口径冲突",
]


def write_main(ws: Worksheet, rows: list[dict]) -> None:
    ws.title = "01_指标主卡"
    ws["A1"] = "每个指标一张对照卡 · 123 行 · 顺序=指标字典"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:N1")
    ws["A2"] = "左边 7 列来自字典，不改写。右边是用例。数字默认按 00_怎么读 的场A～场I。"
    ws.merge_cells("A2:N2")
    ws["A2"].alignment = WRAP
    for i, h in enumerate(MAIN_COLS, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(MAIN_COLS))
    for r, row in enumerate(rows, 5):
        vals = [
            row["code"],
            row["name"],
            row["priority"],
            row["level"],
            row["visible"],
            row["meaning"],
            row["method"],
            bullets(row["include"]),
            bullets(row["exclude"]),
            row["session"],
            row["expect"],
            row["fail"],
            row["json_path"],
            "是" if row["conflict"] else "否",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Microsoft YaHei", size=9)
            if row["priority"] == "P0":
                cell.fill = P0_FILL
            elif row["priority"] == "P1":
                cell.fill = P1_FILL
            if row["conflict"] and c == 14:
                cell.fill = CONFLICT_FILL
        ws.row_dimensions[r].height = 88
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:N{4+len(rows)}"
    widths = {
        1: 26,
        2: 18,
        3: 8,
        4: 12,
        5: 12,
        6: 28,
        7: 36,
        8: 42,
        9: 42,
        10: 28,
        11: 42,
        12: 32,
        13: 32,
        14: 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


BOUND_COLS = [
    "用例ID",
    "指标编码",
    "中文名称",
    "优先级",
    "边界类型",
    "准备什么素材",
    "手算期望",
    "怎样算失败",
    "看哪个字段",
]


def write_bounds(ws: Worksheet, rows: list[dict]) -> None:
    ws.title = "02_边界补充"
    ws["A1"] = "主卡之外必须测的边界 · 分母为 0 / 缺失 / 口径冲突"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:I1")
    ws["A2"] = "不要用围观视频冒充「分母为 0」。围观是整场不准入；分母为 0 是用户参与了、但没遇到该局面。"
    ws.merge_cells("A2:I2")
    ws["A2"].alignment = WRAP
    for i, h in enumerate(BOUND_COLS, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, len(BOUND_COLS))
    for r, row in enumerate(rows, 5):
        vals = [
            row["id"],
            row["code"],
            row["name"],
            row["priority"],
            row["kind"],
            row["prepare"],
            row["expect"],
            row["fail"],
            row["json_path"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Microsoft YaHei", size=9)
            if row["priority"] == "P0":
                cell.fill = P0_FILL
            if row["kind"] == "口径冲突":
                cell.fill = CONFLICT_FILL
        ws.row_dimensions[r].height = 56
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:I{4+len(rows)}"
    for col, w in {1: 12, 2: 26, 3: 18, 4: 8, 5: 12, 6: 48, 7: 42, 8: 36, 9: 32}.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def build_bounds(metrics: list[dict], obj_map: dict) -> list[dict]:
    out: list[dict] = []
    n = 0

    def add(m: dict, kind: str, prepare: str, expect: str, fail: str) -> None:
        nonlocal n
        n += 1
        out.append(
            {
                "id": f"BD-{n:03d}",
                "code": m["code"],
                "name": m["name"],
                "priority": m["priority"],
                "kind": kind,
                "prepare": prepare,
                "expect": expect,
                "fail": fail,
                "json_path": json_path(m["code"], obj_map),
            }
        )

    for m in metrics:
        code = m["code"]
        typ = calc_type(code)
        fx = fixture_for(code)
        path = json_path(code, obj_map)
        if typ == "T2" or code in OPPORTUNITY_RATES or (
            typ == "T7" and code.endswith("_rate")
        ):
            add(
                m,
                "分母为0",
                fx["zero_denom"]
                or "用户真实入座并参与了若干手，但从未出现该指标的机会（例如从未面对 open）。不要用围观视频。",
                f"`{path}` 为 null/省略，用户侧不展示 0% 或 100%。",
                "输出 0% / 100% 充数；或用围观场当本条素材。",
            )
        if typ in {"T0", "T3", "T5"} or fx.get("missing"):
            add(
                m,
                "源数据缺失",
                fx["missing"]
                or "缺该字段源数据（无视频/无结算/无识别）。另备一场齐全的对照，不得把对照值抄过来。",
                f"`{path}` 省略或 null；模块可 degraded/omit。",
                "填 0、空串、默认牌面、默认时长、假 0%。",
            )
        if code in CONFLICT:
            add(
                m,
                "口径冲突",
                f"准备能让两种口径算出不同数字的场。冲突点：{CONFLICT[code]}。可用场A（摊牌率 3/6=50% vs 3/8=37.5%）。",
                "必须等于字典/基准。场A 摊牌率=50%，参与=6 不是按 VPIP=5。",
                f"等于 Skill 口径则失败。{CONFLICT[code]}",
            )
    return out


def main() -> None:
    metrics, obj_map, _mod = load_source()
    missing = [m["code"] for m in metrics if m["code"] not in E or m["code"] not in FIXTURES]
    if missing:
        raise SystemExit(f"主卡缺素材或期望: {missing}")

    cards = []
    for m in metrics:
        fx = fixture_for(m["code"])
        ex = expect_for(m["code"])
        cards.append(
            {
                **m,
                "include": fx["include"],
                "exclude": fx["exclude"],
                "session": ex["session"] + "。 " + (fx.get("mix") or ""),
                "expect": ex["expect"],
                "fail": ex["fail"],
                "json_path": json_path(m["code"], obj_map),
                "conflict": m["code"] in CONFLICT,
            }
        )
    bounds = build_bounds(metrics, obj_map)

    wb = Workbook()
    write_intro(wb.active)
    write_main(wb.create_sheet(), cards)
    write_bounds(wb.create_sheet(), bounds)
    out = ROOT / "指标用例.xlsx"
    wb.save(out)
    print("cards", len(cards), "bounds", len(bounds), "wrote", out)


if __name__ == "__main__":
    main()
