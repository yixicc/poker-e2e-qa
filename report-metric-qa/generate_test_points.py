#!/usr/bin/env python3
"""生成复盘报告模块：测试框架 + 测试点（不含详细用例步骤）。"""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent
SRC = Path("/Users/xixi/Downloads/德州扑克复盘报告_数据指标需求清单.xlsx")

CONFLICT = {
    "showdown_rate": "分母：需求/基准=hands_played，Skill=total_hands",
    "passive_call_ratio": "分母：需求/基准含 check，Skill 不含 check",
    "postflop_call_rate": "分母：需求=翻后面对下注次数，Skill 未写清手数或次数",
    "postflop_raise_rate": "分母：需求=翻后面对下注次数，Skill 未写清手数或次数",
    "hands_played": "需求含盲注防守/进翻后/有盈亏；Skill 更接近自愿入池",
    "participation_rate": "随 hands_played 口径分叉，Skill 语义接近 VPIP",
    "profit_by_position": "需求=筹码净赢；Skill=BB（chips 另有字段）",
}

OPPORTUNITY_RATES = {
    "limp_rate",
    "call_open_rate",
    "three_bet_rate",
    "four_bet_rate",
    "fold_to_3bet_rate",
    "bb_defense_rate",
    "sb_defense_rate",
    "cbet_flop_rate",
    "turn_barrel_rate",
    "river_bet_rate",
    "river_call_vs_raise_rate",
    "postflop_call_rate",
    "postflop_raise_rate",
    "opponent_raise_rate",
    "opponent_call_rate",
    "opponent_fold_rate",
    "opponent_3bet_rate",
    "opponent_cbet_rate",
    "bb_defense_quality",
    "sb_defense_quality",
}

FACT_CODES = {
    "video_duration",
    "position_name",
    "actual_action",
    "gto_recommended_action",
    "gto_action_tendency",
    "opponent_id",
    "opponent_name",
    "hand_index",
    "video_timestamp",
    "hero_cards",
    "hero_position",
    "blinds",
    "effective_stack",
    "board_flop",
    "board_turn",
    "board_river",
    "pot_size_by_street",
    "action_sequence_preflop",
    "action_sequence_flop",
    "action_sequence_turn",
    "action_sequence_river",
    "hero_actual_actions",
    "showdown_result",
    "is_multiway_pot",
    "is_big_pot",
    "is_all_in_or_near_all_in",
    "main_opponent",
    "hand_order_sequence",
}

LIST_CODES = {
    "top_profit_hands",
    "top_loss_hands",
    "good_decision_bad_result_hands",
    "bad_process_good_result_hands",
    "emotion_node_hand_range",
    "hero_response_vs_opponent",
}

LABEL_CODES = {
    "ev_relation_type",
    "luck_direction",
    "decision_impact_direction",
    "decision_quality_label",
}

COMPOSITE_CODES = {
    "weak_top_pair_continue_rate",
    "marginal_hand_entry_rate",
    "profit_concentration_ratio",
    "marginal_entry_after_win",
    "marginal_entry_after_loss",
    "weak_AK_entry_count",
    "low_connector_entry_count",
    "oop_marginal_entry_count",
}

TYPE_NAME = {
    "T0": "事实字段",
    "T1": "计数",
    "T2": "比率",
    "T3": "金额",
    "T4": "分类标签",
    "T5": "列表/TopN",
    "T6": "分组",
    "T7": "复合/白名单",
}

S_NAME = {
    "S1": "计入/正例",
    "S2": "不计入/反例",
    "S3": "分母机会",
    "S4": "分母为0",
    "S5": "聚合/自洽",
    "S6": "单位与精度",
    "S7": "缺失/降级",
    "S8": "口径冲突",
}

PACKS = [
    ("PACK-COUNT", "基础计数", "参与/翻前弃/进翻后/摊牌互斥结果齐全，验 Session 计数"),
    ("PACK-ENTRY", "入池边界", "自愿入池 vs 只交盲、limp vs raise，拆 hands_played 与 VPIP"),
    ("PACK-OPP", "机会分母", "面对 open/3bet/BB被加/可c-bet/河牌被加：做了 vs 有机会没做"),
    ("PACK-MONEY", "金额与换算", "赢/亏/平；至少两档大盲，验按手换算 BB"),
    ("PACK-GTO", "GTO 分界", "首选一致、频率刚好≥/<30%、加注尺度刚好在±30%内外、信息不足"),
    ("PACK-SEAT", "位置映射", "2～9 人桌各覆盖，验位置名不是写死六人桌"),
    ("PACK-VILLAIN", "对手", "单挑 vs 多人、同一对手多手，验交手手数与盈亏归属"),
    ("PACK-DEFINE", "定义白名单", "弱顶对/边缘起手/边缘连张：名单内 vs 看起来像但不算"),
    ("PACK-MISS", "缺失降级", "无视频、无 GTO、无盈亏、无对手，验 omit/degraded"),
    ("PACK-CALIBER", "口径冲突", "专门打分母或单位分叉的指标"),
    ("PACK-JSON", "整包 E2E", "任意完整场次的报告 JSON，跑结构/取值/自洽/status 四道门"),
]

MODULE_STATUS = [
    ("smart_summary_status", "智能总结"),
    ("emotion_module_status", "情绪与心态波动节点"),
    ("key_hand_module_status", "关键手牌深度复盘"),
    ("opponent_position_module_status", "对手与位置表现分析"),
    ("pnl_ev_module_status", "盈亏与期望分析"),
    ("hand_type_warning_status", "重点注意的牌局类型"),
    ("improvement_suggestion_status", "整体改进建议"),
]

JSON_OBJECTS = [
    ("session_summary", "Session基础"),
    ("profit_ev_summary", "盈亏EV"),
    ("gto_decision_points", "GTO决策"),
    ("style_profile_stats", "打法画像"),
    ("position_stats", "位置数据"),
    ("opponent_stats", "对手数据"),
    ("hands", "单手牌"),
    ("timeline_stats", "心态波动"),
]


def load_source() -> tuple[list[dict], dict[str, dict], dict[str, list]]:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["01_指标字典"]
    metrics = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        metrics.append(
            {
                "code": str(r[0]).strip(),
                "name": str(r[1] or "").strip(),
                "meaning": str(r[2] or "").strip(),
                "method": str(r[3] or "").strip(),
                "level": str(r[4] or "").strip(),
                "visible": str(r[5] or "").strip(),
                "support": str(r[6] or "").strip(),
                "priority": str(r[7] or "P1").strip(),
            }
        )
    obj_map: dict[str, dict] = {}
    ws = wb["04_数据包结构"]
    start = 16
    for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if row and row[0] == "数据包对象" and row[1] == "指标编码":
            start = i + 1
            break
    for r in ws.iter_rows(min_row=start, values_only=True):
        if not r[1]:
            continue
        obj_map[str(r[1]).strip()] = {
            "object": str(r[0] or "").strip(),
            "method": str(r[6] or "").strip(),
        }
    mod_map: dict[str, list] = defaultdict(list)
    ws = wb["02_报告模块映射"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[2]:
            continue
        mod_map[str(r[2]).strip()].append(
            {
                "report_module": str(r[0] or "").strip(),
                "status_field": str(r[7] or "").strip(),
            }
        )
    return metrics, obj_map, dict(mod_map)


def calc_type(code: str) -> str:
    if code in FACT_CODES:
        return "T0"
    if code in LIST_CODES:
        return "T5"
    if code in LABEL_CODES:
        return "T4"
    if code in COMPOSITE_CODES:
        return "T7"
    if (
        "_by_" in code
        or code.startswith("segment_")
        or code.startswith("risk_by_")
        or code.endswith("_by_position")
        or code == "opponent_position_distribution"
    ):
        return "T6"
    if code.endswith("_rate") or code.endswith("_ratio"):
        return "T2"
    if any(k in code for k in ("profit", "ev", "chips")) and "rate" not in code:
        return "T3"
    if (
        code.endswith("_count")
        or code.endswith("_hands")
        or code.endswith("_points")
        or code.endswith("_streak_count")
        or "count" in code
    ):
        return "T1"
    return "T0"


def json_path(code: str, obj_map: dict[str, dict]) -> str:
    obj = (obj_map.get(code) or {}).get("object") or "unknown"
    if obj == "hands":
        return f"hands[].{code}"
    if obj == "gto_decision_points" and code in {
        "actual_action",
        "gto_recommended_action",
        "gto_action_tendency",
        "decision_quality_label",
    }:
        return f"gto_decision_points[].{code}"
    if obj == "opponent_stats" or code in {
        "hands_vs_hero",
        "profit_vs_opponent",
        "ev_vs_opponent",
        "hero_response_vs_opponent",
    }:
        if code in {"opponent_id", "opponent_name"}:
            return f"opponent_stats[].{code}"
        return f"opponent_stats.{code}"
    if obj == "position_stats":
        listed = "by_position" in code or code == "position_name"
        return f"position_stats[].{code}" if listed else f"position_stats.{code}"
    return f"{obj}.{code}"


def default_pack(code: str, typ: str, sid: str) -> str:
    if sid == "S8" or code in CONFLICT:
        return "PACK-CALIBER"
    if sid == "S7" or sid == "S4":
        return "PACK-MISS" if sid == "S7" else "PACK-OPP"
    if typ == "T0":
        if code.startswith("board_") or code.startswith("action_sequence") or code.startswith("hero_"):
            return "PACK-COUNT"
        if "position" in code:
            return "PACK-SEAT"
        if code.startswith("opponent") or code == "main_opponent":
            return "PACK-VILLAIN"
        if code == "video_duration":
            return "PACK-MISS"
        return "PACK-COUNT"
    if typ == "T7":
        return "PACK-DEFINE"
    if typ == "T4" or code.startswith("gto_") or "decision" in code or "risk" in code:
        return "PACK-GTO"
    if typ == "T3" or "profit" in code or code.endswith("_ev"):
        return "PACK-MONEY"
    if typ == "T6" and "opponent" in code:
        return "PACK-VILLAIN"
    if typ == "T6" and ("position" in code or code.startswith("segment_")):
        return "PACK-SEAT" if "position" in code else "PACK-COUNT"
    if code in OPPORTUNITY_RATES or sid == "S3":
        return "PACK-OPP"
    if code in {"vpip_rate", "hands_played", "participation_rate", "limp_rate", "pfr_rate"}:
        return "PACK-ENTRY"
    if typ == "T5":
        return "PACK-MONEY"
    return "PACK-COUNT"


def scenarios_for(m: dict) -> list[dict]:
    code, name, method, typ = m["code"], m["name"], m["method"], m["calc_type"]
    out: list[dict] = []

    def add(sid: str, title: str, criteria: str, dim: list[str]) -> None:
        out.append(
            {
                "scenario": sid,
                "scenario_name": S_NAME[sid],
                "title": title,
                "pass_criteria": criteria,
                "coverage_dimensions": dim,
            }
        )

    if typ == "T0":
        add(
            "S1",
            f"{name}有值时路径与格式正确",
            f"JSON `{m['json_path']}` 存在且格式/枚举符合口径：{method}",
            ["normal", "consistency"],
        )
        add(
            "S2",
            f"{name}在规则外或不可识别时不得编造",
            f"不满足口径时 `{m['json_path']}` 省略或为空，不得填默认牌面/动作/位置。",
            ["negative", "consistency"],
        )
        add(
            "S5",
            f"{name}与同源事实自洽",
            f"`{m['json_path']}` 能从 `hands[]` 行动链/牌面/座位反推，不得与 `action_sequence_*` / 公牌 / 人数位置表冲突。",
            ["consistency"],
        )
        add(
            "S7",
            f"{name}缺失时按报告规则省略",
            f"源数据缺失时字段不出现或为 null，对应模块可 degraded/omit，用户侧不出现残缺假值。",
            ["recovery"],
        )
    elif typ == "T1":
        add(
            "S1",
            f"{name}按口径计入",
            f"满足「{method}」的手/动作必须进入 `{m['json_path']}`，计数值=独立标记条数。",
            ["normal"],
        )
        add(
            "S2",
            f"{name}按口径排除",
            f"不满足计入条件的手/动作不得进入分子；看似相近的情况（只交盲、未摊牌、信息不足等）必须排除。",
            ["negative"],
        )
        add(
            "S5",
            f"{name}可从明细加总还原",
            f"由 `hands[]` 或决策点列表按同一口径求和，必须等于 `{m['json_path']}`。",
            ["consistency"],
        )
    elif typ == "T2":
        add(
            "S1",
            f"{name}分子正例计入",
            f"满足分子条件的手/动作计入。口径：{method}",
            ["normal"],
        )
        add(
            "S2",
            f"{name}分子反例不计入",
            f"不满足分子条件（含盲注被动、无机会却发生的动作等）不得计入分子。",
            ["negative"],
        )
        add(
            "S3",
            f"{name}分母机会定义正确",
            f"只有口径规定的机会进入分母。公式：{method}。机会未出现的手不得稀释或虚增。",
            ["normal", "boundary"],
        )
        if code in OPPORTUNITY_RATES or code.endswith("_rate"):
            add(
                "S4",
                f"{name}分母为 0 时不造假百分比",
                f"分母为 0 时 `{m['json_path']}` 为 null/省略，不得输出 0% 或 100% 充数；用户模块按缺失省略。",
                ["boundary", "recovery"],
            )
        add(
            "S6",
            f"{name}单位与精度符合契约",
            f"内部计算保留原始精度；对外若为百分比，按需求：过程小数、展示约定（GTO 三类为整数百分比，其余按目录）。",
            ["consistency"],
        )
    elif typ == "T3":
        add(
            "S1",
            f"{name}单点取值来源正确",
            f"单手或整场取值来自结算字段或约定公式：{method}，不得用起始-结束筹码的未核验差值替代已核验净赢。",
            ["normal", "consistency"],
        )
        add(
            "S5",
            f"{name}聚合加总得上",
            f"单手金额之和（或分组之和）等于整场 `{m['json_path']}`；不得漏手、重复手。",
            ["consistency"],
        )
        add(
            "S6",
            f"{name}筹码与 BB 换算按手大盲",
            f"BB = 该手 chips ÷ 该手 bigBlind，禁止用场均大盲；精度按内部≥4 位、汇总展示 2 位（若文档约定）。",
            ["boundary", "consistency"],
        )
        add(
            "S7",
            f"{name}未过结算准入不得进入金额统计",
            f"结算冲突/不可用的手不进入 `{m['json_path']}`，模块可降级，不得用 0 冒充已核验盈亏。",
            ["negative", "recovery"],
        )
    elif typ == "T4":
        add(
            "S1",
            f"{name}枚举值全集各有正例",
            f"口径声明的每一类标签（如符合GTO/基本可接受/风险偏高/信息不足，或高于/接近/低于期望）都必须能被独立打出。口径：{method}",
            ["normal"],
        )
        add(
            "S2",
            f"{name}分界点归类正确",
            f"频率 30%、加注尺度 ±30%、EV 阈值等分界必须落在正确一类；刚好在边界上的样本可判定且稳定。",
            ["boundary"],
        )
        add(
            "S5",
            f"{name}与计数字段互斥加总",
            f"分类计数互斥；可判断类之和=judgeable；不得一张决策点进两类。",
            ["consistency"],
        )
    elif typ == "T5":
        add(
            "S1",
            f"{name}入选条件正确",
            f"只有满足「{method}」的手进入列表。",
            ["normal"],
        )
        add(
            "S2",
            f"{name}不应入选的手被排除",
            f"未过结算/质量准入，或不符合盈亏/质量组合条件的手不得出现在 `{m['json_path']}`。",
            ["negative"],
        )
        add(
            "S5",
            f"{name}排序键与 TopN 正确",
            f"排序字段与方向符合口径；N 不足时返回实际条数；并列规则稳定。",
            ["boundary", "consistency"],
        )
        add(
            "S7",
            f"{name}空列表时省略或为空数组",
            f"无入选样本时不得编造占位手牌；模块可 omit/degraded。",
            ["recovery"],
        )
    elif typ == "T6":
        add(
            "S1",
            f"{name}分组键完整且命名正确",
            f"按位置/对手/街/时段分组；位置名必须走人数映射表（2–9 人），不得写死六人桌。口径：{method}",
            ["normal"],
        )
        add(
            "S4",
            f"{name}空组处理正确",
            f"某位置/对手/时段无样本时，该组省略或比率为 null，不得用 0 冒充“该位置没入池”。",
            ["boundary", "recovery"],
        )
        add(
            "S5",
            f"{name}分组加总等于全集",
            f"各组计数/金额之和=未分组总计；比率类按组内分子分母重算，不得拿全场比率复用到每组。",
            ["consistency"],
        )
        if "rate" in code or "ratio" in code or "vpip" in code or "pfr" in code:
            add(
                "S6",
                f"{name}分组比率单位正确",
                f"每组比率独立计算，单位与全场同类指标一致。",
                ["consistency"],
            )
    elif typ == "T7":
        add(
            "S1",
            f"{name}白名单内计入",
            f"落在《06_定义补充》白名单（弱顶对/边缘起手/边缘连张等）且满足动作条件的样本必须计入。口径：{method}",
            ["normal"],
        )
        add(
            "S2",
            f"{name}白名单外不计入",
            f"看起来相似但不在白名单（如 ATs 不当弱顶对踢脚、T9s 是否算边缘连张按表执行）不得计入。",
            ["negative"],
        )
        add(
            "S3",
            f"{name}复合条件只满足一半不计入",
            f"仅满足牌力或仅满足动作/时间窗口（如连亏后 N 手）之一时，不得计入。",
            ["boundary"],
        )
        add(
            "S7",
            f"{name}无法识别牌力或窗口时省略",
            f"手牌/公共牌/时间序不足时不计入、不编造；比率分母为 0 时走 S4 规则。",
            ["recovery"],
        )

    if code in CONFLICT:
        add(
            "S8",
            f"{name}口径冲突必须按需求清单/基准详解判定",
            f"冲突点：{CONFLICT[code]}。期望以需求清单（及已确认的基准详解）为准；实现若等于 Skill 口径，记冲突，不判通过。",
            ["consistency"],
        )
    return out


def build_points(metrics: list[dict], obj_map: dict, mod_map: dict) -> list[dict]:
    points: list[dict] = []
    n = 0

    def next_id() -> str:
        nonlocal n
        n += 1
        return f"TP-RPT-{n:03d}"

    for m in metrics:
        code = m["code"]
        typ = calc_type(code)
        path = json_path(code, obj_map)
        rec = {
            **m,
            "calc_type": typ,
            "calc_type_name": TYPE_NAME[typ],
            "json_path": path,
            "json_object": (obj_map.get(code) or {}).get("object", ""),
            "conflict": CONFLICT.get(code),
            "report_modules": [x["report_module"] for x in mod_map.get(code, [])],
            "status_fields": sorted({x["status_field"] for x in mod_map.get(code, []) if x.get("status_field")}),
        }
        for sc in scenarios_for(rec):
            pid = next_id()
            points.append(
                {
                    "id": pid,
                    "family": "metric_scenario",
                    "title": sc["title"],
                    "metric_code": code,
                    "metric_name": m["name"],
                    "priority": m["priority"] if m["priority"] in {"P0", "P1", "P2"} else "P1",
                    "calc_type": typ,
                    "calc_type_name": TYPE_NAME[typ],
                    "scenario": sc["scenario"],
                    "scenario_name": sc["scenario_name"],
                    "json_object": rec["json_object"],
                    "json_path": path,
                    "data_layer": m["level"],
                    "visible": m["visible"],
                    "report_modules": rec["report_modules"],
                    "status_fields": rec["status_fields"],
                    "formula": m["method"],
                    "pass_criteria": sc["pass_criteria"],
                    "coverage_dimensions": sc["coverage_dimensions"],
                    "data_pack": default_pack(code, typ, sc["scenario"]),
                    "caliber_conflict": code in CONFLICT,
                    "conflict_note": CONFLICT.get(code, ""),
                    "status": "designed",
                }
            )

    # JSON gates
    for obj, layer in JSON_OBJECTS:
        points.append(
            {
                "id": next_id(),
                "family": "json_gate",
                "title": f"整包存在对象 {obj} 且粒度正确",
                "metric_code": "",
                "metric_name": obj,
                "priority": "P0",
                "calc_type": "JSON-G1",
                "calc_type_name": "JSON结构",
                "scenario": "G1",
                "scenario_name": "结构契约",
                "json_object": obj,
                "json_path": obj,
                "data_layer": layer,
                "visible": "包级",
                "report_modules": [],
                "status_fields": [],
                "formula": "04_数据包结构",
                "pass_criteria": f"`{obj}` 存在；列表粒度正确（hands 一手一行，决策点一层一行，对手/位置按组一行）；类型与单位按指标目录。",
                "coverage_dimensions": ["consistency"],
                "data_pack": "PACK-JSON",
                "caliber_conflict": False,
                "conflict_note": "",
                "status": "designed",
            }
        )

    invariants = [
        ("hands.length == session_summary.total_hands", "hands[] 条数等于总手数"),
        ("sum(hands.hand_profit_chips) == profit_ev_summary.total_profit_chips", "单手筹码盈亏之和等于整场筹码盈亏"),
        ("sum(hands.hand_profit_bb) == profit_ev_summary.total_profit_bb", "单手 BB 盈亏之和等于整场 BB（按手换算后加总）"),
        ("gto_matched+acceptable+high_risk == judgeable", "GTO 三类互斥且加总等于可判断数"),
        ("gto_match_rate == matched/judgeable", "GTO 匹配度由计数还原，整数百分比按舍入规则"),
        ("acceptable_decision_rate + high_risk_decision_rate == 100%", "四舍五入前可接受+高风险=100%，展示允许 1 个百分点误差"),
        ("sum(position profit chips) == total_profit_chips", "各位置筹码盈亏之和等于整场"),
        ("rate * denominator == numerator", "任一比率可用分子分母还原（按声明舍入）"),
        ("user-visible value == internal same code", "用户可见字段与内部同名指标同一值"),
        ("winning+losing+breakeven <= total_hands", "盈亏手数不大于总手数，且与单手符号一致"),
    ]
    for formula, title in invariants:
        points.append(
            {
                "id": next_id(),
                "family": "json_gate",
                "title": title,
                "metric_code": "",
                "metric_name": "cross_field",
                "priority": "P0",
                "calc_type": "JSON-G3",
                "calc_type_name": "JSON自洽",
                "scenario": "G3",
                "scenario_name": "交叉一致性",
                "json_object": "report_bundle",
                "json_path": formula,
                "data_layer": "整包",
                "visible": "包级",
                "report_modules": [],
                "status_fields": [],
                "formula": formula,
                "pass_criteria": f"整包 JSON 必须满足：{formula}。单字段通过但本式失败则整包失败。",
                "coverage_dimensions": ["consistency"],
                "data_pack": "PACK-JSON",
                "caliber_conflict": False,
                "conflict_note": "",
                "status": "designed",
            }
        )

    points.append(
        {
            "id": next_id(),
            "family": "json_gate",
            "title": "P0 用户可见字段均能从 hands[] 或决策点反推",
            "metric_code": "",
            "metric_name": "traceability",
            "priority": "P0",
            "calc_type": "JSON-G2",
            "calc_type_name": "JSON取值",
            "scenario": "G2",
            "scenario_name": "取值可追溯",
            "json_object": "report_bundle",
            "json_path": "*",
            "data_layer": "整包",
            "visible": "是",
            "report_modules": ["智能总结"],
            "status_fields": ["smart_summary_status"],
            "formula": "用户可见指标 ⊂ 由明细可计算集合",
            "pass_criteria": "每个用户可见指标都能从 hands[] / gto_decision_points[] 按同一口径算回；对不上则失败。",
            "coverage_dimensions": ["consistency"],
            "data_pack": "PACK-JSON",
            "caliber_conflict": False,
            "conflict_note": "",
            "status": "designed",
        }
    )

    for field, mod in MODULE_STATUS:
        for state, crit in (
            ("normal", "具备该模块 normal 条件时 status=normal，对应小节字段齐全。"),
            ("degraded", "只满足部分条件时 status=degraded，只保留有数据支撑的字段，不编造缺失项。"),
            ("omit", "不满足最小条件时 status=omit，对应小节缺席或为空，用户不可见残缺结论。"),
        ):
            points.append(
                {
                    "id": next_id(),
                    "family": "module_status",
                    "title": f"{mod} 在 {state} 条件下 status 与内容一致",
                    "metric_code": field,
                    "metric_name": mod,
                    "priority": "P0" if state in {"normal", "omit"} else "P1",
                    "calc_type": "JSON-G4",
                    "calc_type_name": "模块降级",
                    "scenario": "G4",
                    "scenario_name": f"模块{state}",
                    "json_object": "module_status",
                    "json_path": field,
                    "data_layer": "模块可用性",
                    "visible": "是" if state != "omit" else "省略",
                    "report_modules": [mod],
                    "status_fields": [field],
                    "formula": "05_模块可用性",
                    "pass_criteria": crit + " 不得向用户暴露字段缺失或解析失败原因。",
                    "coverage_dimensions": ["state", "recovery"],
                    "data_pack": "PACK-MISS" if state != "normal" else "PACK-JSON",
                    "caliber_conflict": False,
                    "conflict_note": "",
                    "status": "designed",
                }
            )
    return points


def coverage_matrix(metrics: list[dict], points: list[dict]) -> list[dict]:
    by = defaultdict(set)
    for p in points:
        if p["family"] == "metric_scenario" and p["metric_code"]:
            by[p["metric_code"]].add(p["scenario"])
    rows = []
    for m in metrics:
        typ = calc_type(m["code"])
        have = by.get(m["code"], set())
        applicable = {s["scenario"] for s in scenarios_for({**m, "calc_type": typ, "json_path": ""})}
        row = {
            "metric_code": m["code"],
            "metric_name": m["name"],
            "priority": m["priority"],
            "calc_type": typ,
        }
        for s in S_NAME:
            if s not in applicable:
                row[s] = "不适用"
            elif s in have:
                row[s] = "已设计"
            else:
                row[s] = "缺失"
        row["designed_n"] = sum(1 for s in applicable if s in have)
        row["applicable_n"] = len(applicable)
        rows.append(row)
    return rows


FRAMEWORK_MD = r"""# 复盘报告模块 · 测试框架

本目录是**新的一套**测试资产，只覆盖「AI 复盘报告」的指标计算与报告 JSON，不接 App 页面用例批次。

当前交付：**测试框架 + 测试点**。详细用例步骤（TC）尚未写。

来源：`德州扑克复盘报告_数据指标需求清单.xlsx`（01 字典 / 02 模块映射 / 04 数据包 / 05 降级 / 06 定义补充）。

---

## 1. 覆盖合同（怎样才算全）

五张网都要闭合。测试点已按这五张网展开。

| 网 | 测试点家族 | 闭合标准 |
|---|---|---|
| 指标字典全量 | `metric_scenario` | 123 个 metric_code 都有适用场景点 |
| 公式场景全量 | S1–S8 | 适用格为「已设计」；不适用必须标明，禁止空白 |
| JSON 契约 | `json_gate` G1 | 8 个数据包对象都有结构点 |
| 交叉一致性 | `json_gate` G3 | 加总、比例还原、列表长度、可见=内部 |
| 模块降级 | `module_status` G4 | 7 个模块 × normal/degraded/omit |

---

## 2. 指标计算类型（T0–T7）

| 类型 | 含义 | 场景怎么铺 |
|---|---|---|
| T0 事实字段 | 手牌、公牌、行动链、位置名 | 有值格式、不编造、与明细自洽、缺失省略 |
| T1 计数 | 手数、决策次数 | 计入、排除、明细加总 |
| T2 比率 | VPIP/PFR/3bet… | 分子正反例、分母机会、分母为 0、单位 |
| T3 金额 | chips / BB / EV | 来源、加总、按手换算 BB、未准入不进 |
| T4 分类标签 | GTO 四类、运气/EV 关系 | 枚举全集、分界点、与计数互斥 |
| T5 列表/TopN | 盈亏最大手、打对但输 | 入选、排除、排序 N、空列表 |
| T6 分组 | 位置/对手/街/时段 | 分组键、空组、组加总=全集 |
| T7 复合/白名单 | 弱顶对、边缘牌、连亏后入池 | 名单内、名单外、半条件、无法识别 |

同一类型共用模板，避免按牌桌剧情笛卡尔积。

---

## 3. 公式场景类（S1–S8）

| 类 | 要打到的边界 |
|---|---|
| S1 计入/正例 | 必须进分子或必须写出的值 |
| S2 不计入/反例 | 看起来像、其实不算 |
| S3 分母机会 | 什么才算一次机会 |
| S4 分母为 0 | 禁止假 0% / 100% |
| S5 聚合/自洽 | 单手/分组能还原整场 |
| S6 单位与精度 | chips vs BB，小数 vs 整数百分比 |
| S7 缺失/降级 | 省略，不编造 |
| S8 口径冲突 | 需求清单 vs Skill，以需求/基准为准 |

每个指标只生成**适用**的 S，不适用在覆盖矩阵写「不适用」。

---

## 4. 数据包（为边界服务，不是为手数服务）

| 包 ID | 用途 |
|---|---|
| PACK-COUNT | 四种 Session 结果互斥齐全 |
| PACK-ENTRY | 自愿入池 vs 只交盲，拆参与和 VPIP |
| PACK-OPP | 各类「面对 X」机会：做了 vs 没做 |
| PACK-MONEY | 赢/亏/平 + 至少两档大盲 |
| PACK-GTO | 30% 频率与 ±30% 尺度分界 |
| PACK-SEAT | 2–9 人位置名映射 |
| PACK-VILLAIN | 单挑/多人、同一对手多手 |
| PACK-DEFINE | 弱顶对/边缘牌白名单 |
| PACK-MISS | 无视频/无 GTO/无盈亏/无对手 |
| PACK-CALIBER | 分母或单位分叉的指标 |
| PACK-JSON | 完整报告包，跑四道门 |

原则：一包只改变一个决策条件。后续写 TC 时按包注入事实层，再抽报告 JSON。

---

## 5. E2E JSON 四道门

1. **G1 结构**：对象在、粒度对、类型对。  
2. **G2 取值**：每个出现的字段 = 该指标独立计算结果；用户可见字段必须能从 `hands[]` / 决策点反推。  
3. **G3 自洽**：加总、比率还原、三类 GTO 互斥、位置之和=整场。  
4. **G4 状态**：`05_模块可用性` 的 7 个 status，三种状态各有点。

执行形态（TC 阶段才写步骤）：注入数据包 → 取整包 JSON → 跑本目录断言表。不做页面点击。

---

## 6. 编号与状态

| 资产 | 规则 |
|---|---|
| 测试点 | `TP-RPT-001` 起，发布后不复用 |
| 用例（未写） | 预留 `TC-RPT-001`，一条只绑一个测试点结论 |
| 批次 | `BATCH-RPT-METRIC-001` |
| 测试点状态 | 现为 `designed`；补齐数据包与期望数字后可 `ready_for_review` |

优先级沿用指标字典 P0/P1/P2。JSON 四道门与模块 omit/normal 记 P0。

---

## 7. 口径冲突（S8）

以下指标期望**以需求清单为准**（若另有已确认基准详解，与清单冲突时先会审再改点，不在执行时改口径）：

- `showdown_rate` 分母  
- `passive_call_ratio` 是否含 check  
- `postflop_call_rate` / `postflop_raise_rate` 分母粒度  
- `hands_played` / `participation_rate` 是否宽于 VPIP  
- `profit_by_position` 单位 chips vs BB  

实现若走出 Skill 指南口径，记「口径冲突」，不判通过。

---

## 8. 下一步（尚未做）

1. 按数据包造最小事实 JSON。  
2. 把「已设计」测试点展开为 TC：给定包 + JSON 路径 + 期望 + 判定。  
3. 期望数字在有事实包之后回填，不在本阶段编造。
"""


def write_framework() -> None:
    (ROOT / "00_测试框架.md").write_text(FRAMEWORK_MD, encoding="utf-8")


def write_json(points: list[dict], matrix: list[dict], metrics: list[dict]) -> None:
    payload = {
        "batchId": "BATCH-RPT-METRIC-001",
        "status": "designed",
        "scope": "复盘报告模块 · 指标计算与报告 JSON",
        "source": str(SRC),
        "counts": {
            "metrics": len(metrics),
            "test_points": len(points),
            "metric_scenario": sum(1 for p in points if p["family"] == "metric_scenario"),
            "json_gate": sum(1 for p in points if p["family"] == "json_gate"),
            "module_status": sum(1 for p in points if p["family"] == "module_status"),
            "caliber_conflict_points": sum(1 for p in points if p["caliber_conflict"]),
        },
        "calc_types": TYPE_NAME,
        "scenarios": S_NAME,
        "data_packs": [{"id": a, "name": b, "purpose": c} for a, b, c in PACKS],
        "test_points": points,
        "coverage_matrix": matrix,
    }
    (ROOT / "test-points.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
TITLE_FONT = Font(bold=True, name="Microsoft YaHei", size=14)
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
P0_FILL = PatternFill("solid", fgColor="FCE4D6")
CONFLICT_FILL = PatternFill("solid", fgColor="F8CBAD")
NA_FILL = PatternFill("solid", fgColor="F2F2F2")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_sheet_intro(ws: Worksheet, title: str, lines: list[str], header_row: int) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    for i, line in enumerate(lines, 2):
        ws.cell(i, 1, line)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.cell(i, 1).alignment = WRAP
        ws.cell(i, 1).font = Font(name="Microsoft YaHei", size=10)
    ws.row_dimensions[1].height = 22
    for i in range(2, header_row):
        ws.row_dimensions[i].height = 18


def autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def write_excel(points: list[dict], matrix: list[dict]) -> None:
    wb = openpyxl.Workbook()

    # 00
    ws = wb.active
    ws.title = "00_使用说明"
    write_sheet_intro(
        ws,
        "复盘报告模块 · 测试点清单（不含详细用例）",
        [
            "批次 BATCH-RPT-METRIC-001。只覆盖报告指标与报告 JSON，不接 App 页面用例。",
            "先测测试点与框架；TC 步骤尚未编写。期望数字待数据包落地后回填。",
            "执行时：按 data_pack 注入事实 → 取报告 JSON → 用 json_path 对照 pass_criteria。",
            "S8 口径冲突以需求清单为准；实现若等于 Skill 口径记冲突，不判通过。",
        ],
        6,
    )
    headers = ["章节", "内容"]
    ws.cell(6, 1, "章节")
    ws.cell(6, 2, "内容")
    style_header(ws, 6, 2)
    rows = [
        ("测试框架", "见同目录 00_测试框架.md"),
        ("指标场景点", "Sheet 02，每个指标的适用 S 类各一条"),
        ("JSON/模块点", "Sheet 03，G1 结构 / G2 追溯 / G3 自洽 / G4 降级"),
        ("覆盖矩阵", "Sheet 04，指标 × S，已设计 / 不适用，不允许空白"),
        ("数据包", "Sheet 05，后续造数与执行按包，不按手数堆"),
        ("机器可读", "test-points.json"),
    ]
    for i, (a, b) in enumerate(rows, 7):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
    ws.freeze_panes = "A7"
    autosize(ws, {1: 18, 2: 72})

    # 02 metric points
    ws = wb.create_sheet("02_指标场景测试点")
    write_sheet_intro(
        ws,
        "指标场景测试点",
        [
            "一行 = 一个可独立判定的结论。标题即场景，不是「指标正常」。",
            "后续写 TC 时：一条用例绑定一个测试点，步骤只做取 JSON + 判定。",
        ],
        5,
    )
    cols = [
        "测试点ID",
        "优先级",
        "指标编码",
        "中文名称",
        "计算类型",
        "场景类",
        "测试点",
        "JSON路径",
        "数据包",
        "通过标准",
        "口径公式",
        "是否用户可见",
        "报告模块",
        "覆盖维度",
        "口径冲突",
        "冲突说明",
        "状态",
    ]
    for i, h in enumerate(cols, 1):
        ws.cell(5, i, h)
    style_header(ws, 5, len(cols))
    r = 6
    for p in points:
        if p["family"] != "metric_scenario":
            continue
        vals = [
            p["id"],
            p["priority"],
            p["metric_code"],
            p["metric_name"],
            f"{p['calc_type']} {p['calc_type_name']}",
            f"{p['scenario']} {p['scenario_name']}",
            p["title"],
            p["json_path"],
            p["data_pack"],
            p["pass_criteria"],
            p["formula"],
            p["visible"],
            "、".join(p["report_modules"]),
            "、".join(p["coverage_dimensions"]),
            "是" if p["caliber_conflict"] else "否",
            p["conflict_note"],
            p["status"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Microsoft YaHei", size=9)
            if p["priority"] == "P0":
                cell.fill = P0_FILL
            if p["caliber_conflict"] and c == 15:
                cell.fill = CONFLICT_FILL
        ws.row_dimensions[r].height = 42
        r += 1
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{r-1}"
    autosize(
        ws,
        {
            1: 14,
            2: 8,
            3: 28,
            4: 22,
            5: 16,
            6: 16,
            7: 36,
            8: 36,
            9: 16,
            10: 56,
            11: 40,
            12: 12,
            13: 22,
            14: 18,
            15: 10,
            16: 36,
            17: 12,
        },
    )

    # 03 json
    ws = wb.create_sheet("03_JSON与模块测试点")
    write_sheet_intro(
        ws,
        "E2E JSON 与模块降级测试点",
        ["G1 结构 · G2 取值追溯 · G3 交叉一致性 · G4 七模块 × 三状态。"],
        5,
    )
    cols3 = [
        "测试点ID",
        "优先级",
        "门",
        "测试点",
        "JSON路径/不变量",
        "数据包",
        "通过标准",
        "模块/对象",
        "状态",
    ]
    for i, h in enumerate(cols3, 1):
        ws.cell(5, i, h)
    style_header(ws, 5, len(cols3))
    r = 6
    for p in points:
        if p["family"] == "metric_scenario":
            continue
        vals = [
            p["id"],
            p["priority"],
            f"{p['scenario']} {p['scenario_name']}",
            p["title"],
            p["json_path"],
            p["data_pack"],
            p["pass_criteria"],
            "、".join(p["report_modules"]) or p["json_object"],
            p["status"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = THIN
            cell.font = Font(name="Microsoft YaHei", size=9)
            if p["priority"] == "P0":
                cell.fill = P0_FILL
        ws.row_dimensions[r].height = 36
        r += 1
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(cols3))}{r-1}"
    autosize(ws, {1: 14, 2: 8, 3: 16, 4: 42, 5: 48, 6: 14, 7: 56, 8: 24, 9: 12})

    # 04 matrix
    ws = wb.create_sheet("04_覆盖矩阵")
    write_sheet_intro(
        ws,
        "指标 × 场景覆盖矩阵",
        ["已设计 = 已有测试点；不适用 = 该指标无此场景类；不允许空白。缺失格表示生成失败。"],
        5,
    )
    mh = ["指标编码", "中文名称", "优先级", "计算类型", *S_NAME.values(), "已设计/适用"]
    for i, h in enumerate(mh, 1):
        ws.cell(5, i, h)
    style_header(ws, 5, len(mh))
    r = 6
    for row in matrix:
        vals = [
            row["metric_code"],
            row["metric_name"],
            row["priority"],
            row["calc_type"],
            *[row[s] for s in S_NAME],
            f"{row['designed_n']}/{row['applicable_n']}",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = THIN
            cell.font = Font(name="Microsoft YaHei", size=9)
            if v == "已设计":
                cell.fill = OK_FILL
            elif v == "不适用":
                cell.fill = NA_FILL
            elif v == "缺失":
                cell.fill = CONFLICT_FILL
        r += 1
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(mh))}{r-1}"
    autosize(ws, {1: 28, 2: 22, 3: 8, 4: 8, **{i: 12 for i in range(5, 14)}, 13: 12})

    # 05 packs
    ws = wb.create_sheet("05_数据包目录")
    write_sheet_intro(ws, "数据包目录", ["写 TC 前按包造最小事实；一包一类边界。"], 5)
    for i, h in enumerate(["数据包ID", "名称", "用途", "关联测试点数"], 1):
        ws.cell(5, i, h)
    style_header(ws, 5, 4)
    pack_n = Counter(p["data_pack"] for p in points)
    r = 6
    for pid, name, purpose in PACKS:
        for c, v in enumerate([pid, name, purpose, pack_n.get(pid, 0)], 1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[r].height = 32
        r += 1
    ws.freeze_panes = "A6"
    autosize(ws, {1: 16, 2: 16, 3: 64, 4: 16})

    wb.save(ROOT / "测试点清单.xlsx")


def main() -> None:
    metrics, obj_map, mod_map = load_source()
    points = build_points(metrics, obj_map, mod_map)
    matrix = coverage_matrix(metrics, points)
    missing = [r for r in matrix if any(r[s] == "缺失" for s in S_NAME)]
    if missing:
        raise SystemExit(f"覆盖矩阵存在缺失格: {[r['metric_code'] for r in missing]}")
    write_framework()
    write_json(points, matrix, metrics)
    write_excel(points, matrix)
    print("metrics", len(metrics))
    print("points", len(points), Counter(p["family"] for p in points))
    print("types", Counter(p["calc_type"] for p in points if p["family"] == "metric_scenario"))
    print("wrote", ROOT)


if __name__ == "__main__":
    main()
