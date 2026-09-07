#!/usr/bin/env python3
"""生成可执行用例：每条写清有效素材、无效素材、怎么测、期望看到什么。"""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from metric_fixtures import (
    JSON_OBJECT_FIXTURES,
    MODULE_FIXTURES,
    fixture_for,
)

ROOT = Path(__file__).resolve().parent
POINTS = json.loads((ROOT / "test-points.json").read_text(encoding="utf-8"))

# S1 主用例已经带有效+无效对照，这两类不再单独成篇（避免四条模板重复）
FOLD_INTO_S1 = {
    ("T0", "S2"),
    ("T0", "S5"),
    ("T1", "S2"),
    ("T1", "S5"),
    ("T2", "S2"),
    ("T3", "S5"),
    ("T4", "S5"),
    ("T5", "S2"),
    ("T5", "S5"),
    ("T6", "S5"),
    ("T7", "S2"),
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
TITLE_FONT = Font(bold=True, name="Microsoft YaHei", size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
P0_FILL = PatternFill("solid", fgColor="FCE4D6")
CONFLICT_FILL = PatternFill("solid", fgColor="F8CBAD")


def bullets(items: list[str]) -> str:
    return "\n".join(f"· {x}" for x in items if x)


def case_type(tp: dict) -> str:
    s = tp["scenario"]
    if s in {"S2", "S7"} or str(tp.get("scenario_name", "")).startswith("模块"):
        return "negative"
    if s in {"S3", "S4", "S6", "G4"}:
        return "boundary"
    return "functional"


def env() -> list[str]:
    return [
        "测试环境能用指定素材生成复盘报告，并取出整包 JSON（或等价数据接口）",
        "当前用户 = Hero；围观、无关视频不得当成该用户本场战绩",
        "口径以《德州扑克复盘报告_数据指标需求清单》为准；标了口径冲突的条目不认 Skill 指南",
    ]


def title_for(tp: dict, fx: dict) -> str:
    name = tp["metric_name"] or tp["title"]
    sid = tp["scenario"]
    if tp["family"] != "metric_scenario":
        return tp["title"]
    mapping = {
        "S1": f"{name}：有效素材必须计入，无效素材必须排除",
        "S2": f"{name}：分界样本必须稳定落在正确一类",
        "S3": f"{name}：有机会但没做，只进分母、不进分子",
        "S4": f"{name}：没有任何机会时不得输出假百分比",
        "S6": f"{name}：单位和换算按口径",
        "S7": f"{name}：缺少源数据时省略，不得编造",
        "S8": f"{name}：按需求清单算，不按 Skill 指南",
        "S5": f"{name}：明细加总必须等于汇总",
    }
    return mapping.get(sid, tp["title"])


def metric_main_steps(tp: dict, fx: dict) -> list[dict]:
    path = tp["json_path"]
    name = tp["metric_name"]
    return [
        {
            "action": f"准备有效素材，并与无效素材配在对照场次里。\n有效：\n{bullets(fx['include'])}\n配场：{fx['mix']}",
            "expectedResult": "素材能生成报告（或明确被拒）；测试员能指出哪几手/哪段视频是有效、哪几手是无效。",
        },
        {
            "action": f"准备无效素材（必须真正用上，不能只写在纸上）：\n{bullets(fx['exclude'])}",
            "expectedResult": "无效素材已进入对照场次或单独对照场，不会被误当成有效样本。",
        },
        {
            "action": f"用有效场次生成复盘报告，打开 JSON 的 `{path}`（{name}）。",
            "expectedResult": (
                f"有效素材全部计入。`{path}` 对得上有效素材手算值"
                + (f"；{fx['unit']}" if fx.get("unit") else "")
                + f"。{tp['pass_criteria']}"
            ),
        },
        {
            "action": "用无效场次（或同一场里的无效手）再看同一字段。",
            "expectedResult": "无效素材不得改变结果：围观/无关不得生成该用户战绩；看起来像但不满足口径的手不得进分子。不得填默认牌面、默认时长、假 0%。",
        },
        {
            "action": fx["consistency"]
            and f"自洽核对：{fx['consistency']}"
            or f"若有 hands[] / 决策点明细，按同一口径手算一遍，与 `{path}` 比较。",
            "expectedResult": "手算与接口值一致；对不上则本条失败。",
        },
    ]


def metric_scene_steps(tp: dict, fx: dict) -> list[dict]:
    path = tp["json_path"]
    sid = tp["scenario"]
    if sid == "S1":
        return metric_main_steps(tp, fx)
    if sid == "S2":
        return [
            {
                "action": f"准备卡在分界上的样本（不要用明显正例）。\n有效侧：\n{bullets(fx['include'][:2] or fx['include'])}",
                "expectedResult": "分界样本能稳定判定，重复跑结果不变。",
            },
            {
                "action": f"对照无效/另一侧：\n{bullets(fx['exclude'][:3] or fx['exclude'])}",
                "expectedResult": f"`{path}` 落在口径规定的那一类。{tp['pass_criteria']}",
            },
        ]
    if sid == "S3":
        opp = fx["opportunity"] or [
            "准备「有机会但 Hero 没做该动作」的手：分母要计、分子不计"
        ]
        return [
            {
                "action": f"准备有机会但没做的素材：\n{bullets(opp)}\n同时保留至少一手做了该动作的有效素材。",
                "expectedResult": "测试员能数出：机会几次、做成几次。",
            },
            {
                "action": f"生成报告，看 `{path}`，用手算 做成次数/机会次数。",
                "expectedResult": f"没做的机会只进分母、不进分子。{tp['pass_criteria']}",
            },
        ]
    if sid == "S4":
        zero = fx["zero_denom"] or "整场没有任何该指标的分母机会（用户参与了，但从未遇到该局面）。"
        return [
            {
                "action": f"准备分母为 0 的场次：{zero}\n不要用围观视频冒充「分母为 0」（围观应整场不准入）。",
                "expectedResult": "场次能生成，但该指标没有一次机会。",
            },
            {
                "action": f"看 `{path}` 和对应报告小节。",
                "expectedResult": f"字段为 null/省略，不得输出 0% 或 100%。用户侧不展示假百分比。{tp['pass_criteria']}",
            },
        ]
    if sid == "S6":
        unit = fx["unit"] or tp["formula"]
        return [
            {
                "action": f"准备能暴露换算错误的素材（例：两档不同大盲、或同时有 check 与 call）。\n有效：\n{bullets(fx['include'][:2] or fx['include'])}",
                "expectedResult": "素材里单位条件已经拉开（不能全场同一个大盲还指望测出换算错）。",
            },
            {
                "action": f"看 `{path}` 的单位与数值。口径：{unit}",
                "expectedResult": f"{tp['pass_criteria']}",
            },
        ]
    if sid == "S7":
        miss = fx["missing"] or "源数据缺失：无视频/无该字段/识别失败。"
        return [
            {
                "action": f"准备缺失素材：{miss}",
                "expectedResult": "系统仍能出报告或按规则整场 omit，不得崩溃后填默认值。",
            },
            {
                "action": f"看 `{path}` 与相关模块 status。",
                "expectedResult": f"字段省略或 null；不得用 0、空串、默认牌面、默认时长冒充。{tp['pass_criteria']}",
            },
        ]
    if sid == "S8":
        note = tp.get("conflict_note") or fx.get("conflict") or ""
        return [
            {
                "action": f"准备一份能让两种口径算出不同数字的场次。冲突点：{note}\n有效素材仍按需求清单勾选。",
                "expectedResult": "手算两个候选值，两者必须不同，否则本条无法验收。",
            },
            {
                "action": f"看 `{path}`。",
                "expectedResult": f"必须等于需求清单/基准详解的值。若等于 Skill 口径，记「口径冲突」不通过。{tp['pass_criteria']}",
            },
        ]
    return metric_main_steps(tp, fx)


def json_steps(tp: dict) -> tuple[list[str], list[str], list[dict]]:
    sid = tp["scenario"]
    obj = tp.get("json_object") or ""
    if sid == "G1":
        inc = JSON_OBJECT_FIXTURES.get(obj, "用户真实参与的完整一场。")
        return (
            [inc],
            ["围观/无关视频", "接口失败返回 HTML/报错页", "本版本声明不交付该对象却静默缺失还判通过"],
            [
                {
                    "action": f"用用户参与的完整场次拉整包 JSON，查 `{tp['json_path']}` 在不在、类型对不对。",
                    "expectedResult": tp["pass_criteria"],
                },
                {
                    "action": "用围观或无关视频对照。",
                    "expectedResult": "不得把围观产出当成该用户完整战绩包；失败响应不得当 JSON 通过。",
                },
            ],
        )
    if sid == "G2":
        return (
            ["一场用户参与、hands[] 与汇总字段都有的完整报告"],
            ["只有汇总、没有明细的包（无法反推）"],
            [
                {
                    "action": "列出所有「用户可见=是」的指标，从 hands[] / 决策点按同一口径重算。",
                    "expectedResult": tp["pass_criteria"],
                }
            ],
        )
    if sid == "G3":
        return (
            ["一场同时有手数、盈亏、GTO 三类计数、位置盈亏的完整报告"],
            ["缺某一侧仍声称加总得上"],
            [
                {
                    "action": f"在同一份 JSON 内验：{tp['formula']}",
                    "expectedResult": tp["pass_criteria"],
                }
            ],
        )
    # G4
    state = tp["scenario_name"].replace("模块", "")
    spec = MODULE_FIXTURES.get(tp["metric_code"], {})
    inc = spec.get(state, f"按 05_模块可用性 准备能把「{tp['metric_name']}」置于 {state} 的素材。")
    others = [spec[k] for k in ("normal", "degraded", "omit") if k != state and k in spec]
    return (
        [inc],
        others or ["不满足该状态却期望 status 变成该值"],
        [
            {
                "action": f"只改变该模块所需条件，生成报告，读 `{tp['json_path']}`。",
                "expectedResult": f"`{tp['json_path']}` = {state}。{tp['pass_criteria']}",
            }
        ],
    )


def to_case(tp: dict, seq: int, extra_tp_ids: list[str] | None = None) -> dict:
    fx = fixture_for(tp["metric_code"]) if tp["family"] == "metric_scenario" else {}
    if tp["family"] == "metric_scenario":
        include = fx["include"]
        exclude = fx["exclude"]
        mix = fx["mix"]
        steps = metric_scene_steps(tp, fx)
        if tp["scenario"] == "S1":
            pre = [
                "本条是对照主用例：有效素材必须计入，无效素材必须排除。",
                mix,
            ]
        elif tp["scenario"] == "S3":
            include = fx["opportunity"] or include
            pre = ["本条只验分母机会：有机会没做的手必须在分母、不在分子。"]
        elif tp["scenario"] == "S4":
            include = [fx["zero_denom"] or "用户参与了，但从未出现该指标的机会"]
            exclude = [
                "用围观/无关视频冒充「分母为 0」（那是整场不准入，不是分母为 0）",
                "其实有一次机会却被漏识别，然后输出 0%",
            ]
            pre = ["不要用围观视频代替「分母为 0」。用户必须真实参与，只是没遇到该局面。"]
        elif tp["scenario"] == "S7":
            include = [fx["missing"] or "源数据缺失"]
            exclude = [
                "对照场：一场源数据齐全的用户参与牌局（本条不得把对照场的值抄到缺失场）",
                "用 0、空串、默认牌面、默认时长、假 0% 填上缺失字段",
            ]
            pre = ["本条只验缺失：不得编造默认值。"]
        elif tp["scenario"] == "S8":
            pre = [f"口径冲突：{tp.get('conflict_note') or fx.get('conflict')}", "只认需求清单。"]
        else:
            pre = [mix]
    else:
        include, exclude, steps = json_steps(tp)
        mix = "有效场次与对照场次分开跑，不要混成一次结果。"
        pre = [mix]

    ids = [tp["id"], *(extra_tp_ids or [])]
    return {
        "id": f"TC-RPT-{seq:03d}",
        "title": title_for(tp, fx),
        "testPointIds": ids,
        "priority": tp["priority"],
        "type": case_type(tp),
        "level": "api",
        "preconditions": pre + env()[:1],
        "valid_materials": include,
        "invalid_materials": exclude,
        "how_to_mix": mix if tp["family"] == "metric_scenario" else "有效场次与对照场次分开跑。",
        "testData": [
            f"metric_code={tp['metric_code'] or '-'}",
            f"json_path={tp['json_path']}",
            f"formula={tp['formula']}",
        ],
        "environment": env(),
        "steps": steps,
        "status": "ready_for_review",
        "version": "2.0",
        "batchId": "BATCH-RPT-METRIC-001",
        "automationCandidate": True,
        "family": tp["family"],
        "scenario": f"{tp['scenario']} {tp['scenario_name']}",
        "metric_code": tp["metric_code"],
        "metric_name": tp["metric_name"],
        "json_path": tp["json_path"],
        "formula": tp["formula"],
        "caliber_conflict": tp["caliber_conflict"],
        "pass_criteria": tp["pass_criteria"],
    }


def select_points() -> list[tuple[dict, list[str]]]:
    """主用例带上被折叠的 S2/S5 测试点 ID。"""
    folded: dict[str, list[str]] = defaultdict(list)
    keep: list[dict] = []
    for tp in POINTS["test_points"]:
        if tp["family"] == "metric_scenario" and (tp["calc_type"], tp["scenario"]) in FOLD_INTO_S1:
            folded[tp["metric_code"]].append(tp["id"])
            continue
        keep.append(tp)
    out = []
    for tp in keep:
        extra = folded.get(tp["metric_code"], []) if tp.get("scenario") == "S1" else []
        out.append((tp, extra))
    return out


def flatten_steps(steps: list[dict]) -> tuple[str, str]:
    acts, exps = [], []
    for i, s in enumerate(steps, 1):
        acts.append(f"{i}. {s['action']}")
        exps.append(f"{i}. {s['expectedResult']}")
    return "\n".join(acts), "\n".join(exps)


def style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


COLS = [
    "用例ID",
    "优先级",
    "指标",
    "场景",
    "要验证什么",
    "有效素材（必须准备、必须计入）",
    "无效素材（必须准备、必须排除）",
    "怎么配场",
    "操作步骤",
    "预期结果",
    "看哪个字段",
    "口径公式",
    "关联测试点",
    "口径冲突",
    "状态",
]


def write_excel(cases: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "00_怎么读这份表"
    ws["A1"] = "复盘报告指标用例 v2.0 — 按「有效素材 / 无效素材」执行"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    guide = [
        ("读法", "不要从公式开始。先看「有效素材」准备什么，再看「无效素材」必须混入什么，最后看字段是否只吃有效、吐掉无效。"),
        (
            "例子：视频时长",
            "有效：用户本人坐在桌上打德州、视频能读出时长（如 12 分 30 秒）。无效：用户只围观、或与德州无关的视频、或读不到时长。期望：有效场次写出元数据时长；围观/无关不得当作用户本场时长；损坏元数据则省略，不得填 0:00。",
        ),
        (
            "例子：参与手数",
            "有效：自愿入池、盲注防守、进翻后、本手有盈亏。无效：围观；以及「只交盲即弃」是否算参与要按清单（有盈亏则算）。同一场必须两种都有。",
        ),
        ("一条用例验一件事", "主用例 = 计入+排除对照。另外还有：有机会没做、分母为 0、缺数据、口径冲突。不要用主用例通过代替分母为 0。"),
        ("怎么跑", "用素材生成报告 → 打开 JSON 对应字段 → 对照预期。不做 App 点击验收。"),
        ("口径冲突", "标「是」的行：必须按需求清单。算出来若等于 Skill 指南，记冲突，不通过。"),
        ("数量", f"本份 {len(cases)} 条。测试点仍是 539（对照矩阵没丢）；S2/S5 折进主用例，避免四条套话。"),
        ("同目录", "素材细则在 metric_fixtures.py；测试点在 测试点清单.xlsx。"),
    ]
    ws.cell(2, 1, "章节")
    ws.cell(2, 2, "说明")
    style_header(ws, 2, 2)
    for i, (a, b) in enumerate(guide, 3):
        ws.cell(i, 1, a).alignment = WRAP
        ws.cell(i, 2, b).alignment = WRAP
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        ws.row_dimensions[i].height = 48
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A3"

    def dump(sheet: str, rows: list[dict], note: str) -> None:
        w = wb.create_sheet(sheet)
        w["A1"] = sheet
        w["A1"].font = TITLE_FONT
        w.merge_cells("A1:J1")
        w["A2"] = note
        w.merge_cells("A2:J2")
        w["A2"].alignment = WRAP
        w.row_dimensions[2].height = 28
        for i, h in enumerate(COLS, 1):
            w.cell(4, i, h)
        style_header(w, 4, len(COLS))
        r = 5
        for c in rows:
            acts, exps = flatten_steps(c["steps"])
            vals = [
                c["id"],
                c["priority"],
                c["metric_name"] or c["title"],
                c["scenario"],
                c["title"],
                bullets(c["valid_materials"]),
                bullets(c["invalid_materials"]),
                c["how_to_mix"],
                acts,
                exps,
                c["json_path"],
                c["formula"],
                "、".join(c["testPointIds"]),
                "是" if c["caliber_conflict"] else "否",
                c["status"],
            ]
            for col, v in enumerate(vals, 1):
                cell = w.cell(r, col, v)
                cell.alignment = WRAP
                cell.border = THIN
                cell.font = Font(name="Microsoft YaHei", size=9)
                if c["priority"] == "P0":
                    cell.fill = P0_FILL
                if c["caliber_conflict"] and col == 14:
                    cell.fill = CONFLICT_FILL
            w.row_dimensions[r].height = 96
            r += 1
        w.freeze_panes = "C5"
        w.auto_filter.ref = f"A4:{get_column_letter(len(COLS))}{r-1}"
        widths = {
            1: 12,
            2: 8,
            3: 16,
            4: 16,
            5: 36,
            6: 46,
            7: 46,
            8: 32,
            9: 42,
            10: 42,
            11: 32,
            12: 28,
            13: 18,
            14: 10,
            15: 14,
        }
        for col, width in widths.items():
            w.column_dimensions[get_column_letter(col)].width = width

    dump(
        "01_P0用例",
        [c for c in cases if c["priority"] == "P0"],
        "先测这些。打开第一行「视频时长」看读法：有效=用户参与的牌局视频，无效=围观或无关视频。",
    )
    dump(
        "02_P1P2用例",
        [c for c in cases if c["priority"] != "P0" and c["family"] == "metric_scenario"],
        "打法画像、位置、对手、心态、白名单。仍然是有效/无效对照，不是「展示正确」。",
    )
    dump(
        "03_口径冲突用例",
        [c for c in cases if c["caliber_conflict"] and c["scenario"].startswith("S8")],
        "必须按需求清单。与 Skill 算出同一个数则本条无法验收（先换能分叉的场次）。",
    )
    dump(
        "04_JSON与模块用例",
        [c for c in cases if c["family"] != "metric_scenario"],
        "整包对象在不在、加总得上、七个模块 normal/degraded/omit。素材见「有效/无效」列。",
    )

    wb.save(ROOT / "测试用例清单.xlsx")


def main() -> None:
    selected = select_points()
    cases = [to_case(tp, i, extra) for i, (tp, extra) in enumerate(selected, 1)]
    payload = {
        "batchId": "BATCH-RPT-METRIC-001",
        "version": "2.0",
        "status": "ready_for_review",
        "sourceTestPoints": str(ROOT / "test-points.json"),
        "counts": {
            "total": len(cases),
            "P0": sum(1 for c in cases if c["priority"] == "P0"),
            "P1": sum(1 for c in cases if c["priority"] == "P1"),
            "P2": sum(1 for c in cases if c["priority"] == "P2"),
            "metric_scenario": sum(1 for c in cases if c["family"] == "metric_scenario"),
            "json_or_module": sum(1 for c in cases if c["family"] != "metric_scenario"),
            "caliber_conflict": sum(1 for c in cases if c["caliber_conflict"] and str(c["scenario"]).startswith("S8")),
        },
        "notes": [
            "v2：每条用例写有效素材和无效素材，不再套「按公式独立计算」空话。",
            "S2/S5 折进各指标主用例，测试点 ID 仍挂在主用例上。",
            "不接 App 页面用例批次。",
        ],
        "cases": cases,
    }
    (ROOT / "test-cases.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_excel(cases)
    print("cases", len(cases), Counter(c["priority"] for c in cases))
    print("wrote", ROOT / "测试用例清单.xlsx")


if __name__ == "__main__":
    main()
