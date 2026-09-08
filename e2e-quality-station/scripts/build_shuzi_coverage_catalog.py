#!/usr/bin/env python3
"""一批次源数据（Excel auto_json 或 jsonl）→ e2e-quality 站点。"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, str(Path(__file__).resolve().parent))
from e2e_quality_pages import legacy_redirect, write_site  # noqa: E402
from e2e_source import load_source  # noqa: E402
from shuzi_field_schema import PAIRS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(os.environ["E2E_SRC"]).expanduser() if os.environ.get("E2E_SRC") else None
OUT_DIR = ROOT / "output"
SITE_DIR = OUT_DIR / "e2e-quality"
COPY_DIR = Path(os.environ["E2E_COPY_DIR"]).expanduser() if os.environ.get("E2E_COPY_DIR") else None
COPY_SITE = (COPY_DIR / "e2e-quality") if COPY_DIR else None
MARKER = "poker-gto-history-qa/"


def stored_source_path(src: Path) -> str:
    src = Path(src).resolve()
    try:
        return str(src.relative_to(ROOT))
    except ValueError:
        return str(src)


def resolve_source_path(raw: str, batch_id: str = "") -> Path:
    """本机绝对路径、相对仓库路径、uploads/ 下按批次名，拷到别的机器也能找到。"""
    raw = str(raw or "").strip()
    candidates: list[Path] = []
    if raw:
        p = Path(raw)
        candidates.append(p)
        if not p.is_absolute():
            candidates.append(ROOT / p)
        norm = raw.replace("\\", "/")
        if MARKER in norm:
            candidates.append(ROOT / norm.split(MARKER, 1)[1])
        name = Path(norm).name
        if name:
            candidates.append(SITE_DIR / "uploads" / name)
    if batch_id:
        uploads = SITE_DIR / "uploads"
        if uploads.exists():
            for f in uploads.iterdir():
                if f.is_file() and (f.stem == batch_id or f.name.startswith(batch_id)):
                    candidates.append(f)
    seen: set[str] = set()
    for c in candidates:
        key = str(c)
        if key in seen:
            continue
        seen.add(key)
        if c.exists():
            return c
    return Path(raw) if raw else Path()

ACTIONS = ["open", "limp", "3bet", "4bet", "bet", "raise", "check", "fold", "all in"]
HERO_POS = ["BTN", "SB", "BB", "UTG", "UTG+1", "UTG+2", "LJ", "HJ", "CO"]
PLAYERS = list(range(3, 10))

FONT = "Microsoft YaHei"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name=FONT, size=10)
TITLE_FONT = Font(bold=True, name=FONT, size=16)
H2_FONT = Font(bold=True, name=FONT, size=12)
BODY = Font(name=FONT, size=10)
MUTED = Font(name=FONT, size=10, color="666666")
KPI_FILL = PatternFill("solid", fgColor="F3EFE7")
NOTE_FILL = PatternFill("solid", fgColor="FFF8DC")
GAP_FILL = PatternFill("solid", fgColor="FCE4D6")
ZERO_FILL = PatternFill("solid", fgColor="F7F7F7")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
MID = Alignment(vertical="center")
MID_C = Alignment(horizontal="center", vertical="center")


def batch_id_of(src: Path) -> str:
    return src.stem.strip() or "batch"


def load_cases(src: Path) -> tuple[list[dict], list[dict]]:
    cases, fields, _trees, _hist = load_source(src)
    return cases, fields


def _style_header(ws: Worksheet, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def _put(ws: Worksheet, r: int, c: int, value, *, font=BODY, fill=None, align=MID, fmt=None):
    cell = ws.cell(r, c, value)
    cell.font = font
    cell.alignment = align
    cell.border = THIN
    if fill is not None:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def _widths(ws: Worksheet, mapping: dict[int, int]) -> None:
    for i, w in mapping.items():
        ws.column_dimensions[get_column_letter(i)].width = w


def matrix(cases: list[dict], row_key: str, col_key: str, rows: list, cols: list) -> list[list]:
    cnt: dict[tuple, int] = Counter((c[row_key], c[col_key]) for c in cases)
    return [[cnt.get((rk, ck), 0) for ck in cols] for rk in rows]


def write_cover(wb: Workbook, cases: list[dict], fields: list[dict], generated: str, source_name: str) -> None:
    ws = wb.active
    ws.title = "说明"
    n = len(cases)
    streets = Counter(c["street"] for c in cases)
    actions = Counter(c["action"] for c in cases)
    showdown_n = sum(1 for c in cases if c["showdown"])
    blinds_n = len({c["blinds"] for c in cases})
    by_cat = Counter(f["cat"] for f in fields)

    ws.merge_cells("A1:I1")
    ws["A1"] = "数字好友桌 e2e · 覆盖目录 + 字段匹配（第 1 期）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:I2")
    ws["A2"] = f"生成 {generated} · 源数据 {source_name} · 匹配率待批跑，不是 0%"
    ws["A2"].font = MUTED

    ws.merge_cells("A4:I4")
    ws["A4"] = "给老板看什么"
    ws["A4"].font = H2_FONT
    ws.merge_cells("A5:I7")
    ws["A5"] = (
        f"覆盖：{n} 手测了哪些街/动作/人数/英雄位。"
        "质量：三对字段匹配率（源数据→视觉、视觉→规则、源数据→规则）。"
        "三份 JSON 先译成同一套界面树（P0=Excel 座位，0=Hero），再递归比。"
        "主指标是字段匹配率，不再报整手全对/救回改坏。昵称按原值比对，批跑后名字类可能接近 0%。"
    )
    ws["A5"].alignment = WRAP
    ws["A5"].fill = NOTE_FILL

    labels = [
        ("用例手数", n),
        ("街", f"翻前{streets['翻前']} / 翻牌{streets['翻牌']} / 转牌{streets['转牌']} / 河牌{streets['河牌']}"),
        ("动作类型", f"{len(actions)} 种"),
        ("人数", "3–9 人"),
        ("盲注档", blinds_n),
        ("摊牌手数", showdown_n),
        ("源数据界面字段路径", len(fields)),
        ("其中 卡牌/行动/文字/图像/数字", f"{by_cat['卡牌']}/{by_cat['行动']}/{by_cat['文字']}/{by_cat['图像']}/{by_cat['数字']}"),
        ("Straddle / 预期边池", "0 / 0"),
    ]
    ws["A9"] = "本期覆盖数字"
    ws["A9"].font = H2_FONT
    _put(ws, 10, 1, "项", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    _put(ws, 10, 2, "值", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, (k, v) in enumerate(labels, start=11):
        _put(ws, i, 1, k)
        _put(ws, i, 2, v, fill=KPI_FILL, font=Font(name=FONT, size=11, bold=True))

    ws["A21"] = "第 2 期待填 · 三对字段匹配率（本期不是 0%）"
    ws["A21"].font = H2_FONT
    headers = ["环节", "完全一致手数", "字段总匹配率", "卡牌类", "文字类", "图像类", "数字类"]
    for col, title in enumerate(headers, start=1):
        _put(ws, 22, col, title, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, pair in enumerate(PAIRS):
        _put(ws, 23 + i, 1, pair["title"])
        for col in range(2, 8):
            _put(ws, 23 + i, col, "待批跑", fill=NOTE_FILL, align=MID_C, font=MUTED)

    ws["A27"] = "比对口径（批跑后生效）"
    ws["A27"].font = H2_FONT
    bullets = [
        "三对独立：源数据→视觉、视觉→规则、源数据→规则。左侧是分母。视觉缺整手时，源数据有的字段记缺失并计入总数。",
        f"对齐：fusion sessionId 列表；批跑序号 1–{n} 对 session 手序。默认翻前→翻牌→转牌→河牌。P0=Excel 座位，视觉/规则按位置映射。",
        "匹配率=匹配/总数。值不同=值不匹配；左侧有右侧无=缺失；数组长短不一=列表长度不匹配。数字 \"20\" 与 20 视为相同。",
        "分类：卡牌=手牌/公牌；文字=name/action/comment/position；图像=dealer/status/avatar（没有叶子则匹配率 —）；数字=筹码/下注/座位/底池/盲注。",
        "session_id、hand_id、时间戳、video_url 不进界面树。昵称、图像类按原值比对。",
    ]
    for i, t in enumerate(bullets, start=28):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
        _put(ws, i, 1, t, font=MUTED, align=WRAP)
        ws.row_dimensions[i].height = 28

    ws["A34"] = "本批边界（汇报时必须说）"
    ws["A34"].font = H2_FONT
    ws.merge_cells("A35:I36")
    ws["A35"] = (
        "好友桌、数字记分牌、无 Straddle、用例预期无边池。"
        "不代表抓头、边池、BB 单位、真实微扑克桌。BB.xlsx 未纳入。"
        "源数据→视觉 / 源数据→规则 的「总数」已按 Excel 界面树盘点；视觉→规则 的总数要等视觉落地。"
    )
    ws["A35"].alignment = WRAP
    ws["A35"].fill = GAP_FILL
    ws["A35"].font = BODY
    ws["A35"].border = THIN

    ws.row_dimensions[5].height = 64
    ws.row_dimensions[35].height = 40
    _widths(ws, {1: 22, 2: 36, 3: 16, 4: 12, 5: 12, 6: 12, 7: 12, 8: 12, 9: 14})
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def write_field_pair(wb: Workbook, pair: dict, fields: list[dict], stats: dict | None = None) -> None:
    ws = wb.create_sheet(pair["title"])
    looked = bool(stats)
    ws["A1"] = f"{pair['title']}（左侧 {pair['left']} → 右侧 {pair['right']}）" + (
        "" if looked else " · 匹配率待批跑"
    )
    ws["A1"].font = H2_FONT
    headers = ["类", "字段", "总数", "匹配", "值不匹配", "缺失", "列表长度不匹配", "匹配率", "出错手数"]
    for i, h in enumerate(headers, start=1):
        _put(ws, 3, i, h, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    use_total = pair["totals"] == "gold"
    for i, f in enumerate(fields, start=4):
        _put(ws, i, 1, f["cat"], align=MID_C)
        _put(ws, i, 2, f["path"])
        row = (stats or {}).get(f["path"]) if looked else None
        if row:
            total = row.get("total")
            _put(ws, i, 3, total, align=MID_C)
            _put(ws, i, 4, row.get("match"), align=MID_C)
            _put(ws, i, 5, row.get("mismatch"), align=MID_C)
            _put(ws, i, 6, row.get("missing"), align=MID_C)
            _put(ws, i, 7, row.get("len_mismatch"), align=MID_C)
            rate = row.get("rate")
            _put(ws, i, 8, None if rate is None else rate / 100.0, align=MID_C, fmt="0.0%")
            _put(ws, i, 9, row.get("n_error"), align=MID_C)
        else:
            if use_total:
                _put(ws, i, 3, f["total"], align=MID_C)
            else:
                _put(ws, i, 3, "", fill=NOTE_FILL, align=MID_C)
            for col in range(4, 10):
                _put(ws, i, col, "", fill=NOTE_FILL, align=MID_C)
    ws.auto_filter.ref = f"A3:I{3 + len(fields)}"
    ws.freeze_panes = "A4"
    _widths(ws, {1: 10, 2: 42, 3: 10, 4: 10, 5: 12, 6: 10, 7: 18, 8: 12, 9: 12})
    ws.row_dimensions[1].height = 22


def write_overview(wb: Workbook, cases: list[dict]) -> None:
    ws = wb.create_sheet("覆盖总览")
    ws["A1"] = "按单一维度计手数（一手只计一次）"
    ws["A1"].font = H2_FONT

    def dump_counter(start_col: int, title: str, items: list[tuple], total: int):
        _put(ws, 3, start_col, title, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
        _put(ws, 3, start_col + 1, "手数", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
        _put(ws, 3, start_col + 2, "占全部", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
        r = 4
        for k, v in items:
            _put(ws, r, start_col, k, align=MID)
            _put(ws, r, start_col + 1, v, align=MID_C)
            _put(ws, r, start_col + 2, v / total, align=MID_C, fmt="0.0%")
            r += 1
        _put(ws, r, start_col, "合计", font=Font(name=FONT, bold=True, size=10))
        _put(ws, r, start_col + 1, total, font=Font(name=FONT, bold=True, size=10), align=MID_C)
        _put(ws, r, start_col + 2, 1, font=Font(name=FONT, bold=True, size=10), align=MID_C, fmt="0.0%")

    n = len(cases)
    dump_counter(1, "街", [(s, sum(1 for c in cases if c["street"] == s)) for s in ["翻前", "翻牌", "转牌", "河牌"]], n)
    dump_counter(5, "动作", [(a, sum(1 for c in cases if c["action"] == a)) for a in ACTIONS], n)
    dump_counter(9, "人数", [(f"{p}人", sum(1 for c in cases if c["n_players"] == p)) for p in PLAYERS], n)
    dump_counter(13, "英雄位", [(h, sum(1 for c in cases if c["hero_pos"] == h)) for h in HERO_POS], n)

    ws["A16"] = "盲注档"
    ws["A16"].font = H2_FONT
    blinds = Counter(c["blinds"] for c in cases).most_common()
    _put(ws, 17, 1, "大小盲", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    _put(ws, 17, 2, "手数", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, (k, v) in enumerate(blinds, start=18):
        _put(ws, i, 1, k, align=MID_C)
        _put(ws, i, 2, v, align=MID_C)

    ws["D16"] = "开关项（开/关手数）"
    ws["D16"].font = H2_FONT
    flags = [("延迟看牌", "delay"), ("Allin两次", "allin2"), ("限制GPS和IP", "gps"), ("随机入座", "random_seat")]
    _put(ws, 17, 4, "项", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    _put(ws, 17, 5, "开", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    _put(ws, 17, 6, "关", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, (label, key) in enumerate(flags, start=18):
        _put(ws, i, 4, label)
        _put(ws, i, 5, sum(1 for c in cases if c[key] == "开"), align=MID_C)
        _put(ws, i, 6, sum(1 for c in cases if c[key] == "关"), align=MID_C)

    _widths(ws, {1: 14, 2: 10, 3: 10, 4: 14, 5: 12, 6: 10, 7: 10, 8: 10, 9: 10, 10: 10, 11: 10, 13: 12, 14: 10, 15: 10})
    ws.freeze_panes = "A3"


def write_cross(wb: Workbook, title: str, cases: list[dict], row_key: str, col_key: str, rows: list, cols: list, row_label: str) -> None:
    ws = wb.create_sheet(title)
    ws["A1"] = f"{row_label} × {col_key}（格子=手数；灰格=结构上不会出现或本批为 0）"
    ws["A1"].font = H2_FONT
    _put(ws, 3, 1, row_label, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, ck in enumerate(cols, start=2):
        _put(ws, 3, i, ck, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    _put(ws, 3, len(cols) + 2, "合计", font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)

    mat = matrix(cases, row_key, col_key, rows, cols)
    for ri, rk in enumerate(rows):
        r = 4 + ri
        _put(ws, r, 1, rk, align=MID)
        row_sum = 0
        for ci, val in enumerate(mat[ri]):
            fill = ZERO_FILL if val == 0 else None
            _put(ws, r, 2 + ci, val, align=MID_C, fill=fill)
            row_sum += val
        _put(ws, r, len(cols) + 2, row_sum, align=MID_C, font=Font(name=FONT, bold=True, size=10))

    total_r = 4 + len(rows)
    _put(ws, total_r, 1, "合计", font=Font(name=FONT, bold=True, size=10))
    for ci in range(len(cols)):
        s = sum(mat[ri][ci] for ri in range(len(rows)))
        _put(ws, total_r, 2 + ci, s, align=MID_C, font=Font(name=FONT, bold=True, size=10))
    _put(ws, total_r, len(cols) + 2, len(cases), align=MID_C, font=Font(name=FONT, bold=True, size=10), fill=KPI_FILL)

    widths = {1: 12}
    widths.update({i: 10 for i in range(2, len(cols) + 3)})
    _widths(ws, widths)
    ws.freeze_panes = "B4"


def write_list(wb: Workbook, cases: list[dict]) -> None:
    ws = wb.create_sheet("用例清单")
    headers = [
        "批跑序号",
        "街",
        "表内行",
        "用例名称",
        "动作",
        "人数",
        "英雄位",
        "大小盲",
        "前注",
        "Hero行动次数",
        "自愿动作数",
        "是否摊牌",
        "总底池",
        "hand_id",
        "延迟看牌",
        "Allin两次",
        "Straddle",
        "预期边池",
        "Excel表",
        "Excel行号",
    ]
    for i, h in enumerate(headers, start=1):
        _put(ws, 1, i, h, font=HEADER_FONT, fill=HEADER_FILL, align=MID_C)
    for i, c in enumerate(cases, start=2):
        vals = [
            c["seq"],
            c["street"],
            c["sheet_row"],
            c["name"],
            c["action"],
            c["n_players"],
            c["hero_pos"],
            c["blinds"],
            c["ante"],
            c["hero_acts"],
            c["n_vol_actions"],
            "是" if c["showdown"] else "否",
            c["total_pot"],
            c["hand_id"],
            c["delay"],
            c["allin2"],
            c["straddle"],
            c["sidepot"],
            c["sheet"],
            c["excel_row"],
        ]
        for col, v in enumerate(vals, start=1):
            _put(ws, i, col, v, align=MID if col == 4 else MID_C)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(cases)+1}"
    ws.freeze_panes = "E2"
    _widths(
        ws,
        {
            1: 10,
            2: 8,
            3: 8,
            4: 48,
            5: 10,
            6: 8,
            7: 10,
            8: 12,
            9: 8,
            10: 14,
            11: 12,
            12: 10,
            13: 10,
            14: 18,
            15: 10,
            16: 12,
            17: 12,
            18: 12,
            19: 12,
            20: 10,
        },
    )
    ws.row_dimensions[1].height = 22


def build_xlsx(
    path: Path,
    cases: list[dict],
    fields: list[dict],
    generated: str,
    source_name: str,
    field_stats: dict | None = None,
) -> None:
    wb = Workbook()
    write_cover(wb, cases, fields, generated, source_name)
    for pair in PAIRS:
        write_field_pair(wb, pair, fields, (field_stats or {}).get(pair["id"]))
    write_overview(wb, cases)
    write_cross(wb, "街×动作", cases, "street", "action", ["翻前", "翻牌", "转牌", "河牌"], ACTIONS, "街")
    write_cross(wb, "人数×英雄位", cases, "n_players", "hero_pos", PLAYERS, HERO_POS, "人数")
    write_cross(wb, "人数×街", cases, "n_players", "street", PLAYERS, ["翻前", "翻牌", "转牌", "河牌"], "人数")
    write_list(wb, cases)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def ingest_source(
    src: Path,
    uploader: str = "",
    *,
    batch_id: str | None = None,
    source_name: str | None = None,
    session_id: str = "",
    note: str = "",
    fetch_fusion: bool = True,
) -> dict:
    src = Path(src)
    if not src.exists():
        raise FileNotFoundError(f"source not found: {src}")

    display_name = source_name or src.name
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    cases, fields, trees, histories = load_source(src)
    if len(cases) < 1:
        raise ValueError("no cases")
    ids = [c["hand_id"] for c in cases]
    if len(set(ids)) != len(ids):
        raise ValueError("duplicate hand_id")
    if not fields:
        raise ValueError("empty field inventory")

    match_info = None
    session_ids: list[str] = []
    stored_sid = session_id.strip()
    if stored_sid and fetch_fusion:
        from e2e_fusion import align_to_cases, fetch_session_results, parse_session_ids
        from e2e_match import match_batch

        session_ids = parse_session_ids(session_id)
        if not session_ids:
            raise ValueError("请填写 fusion sessionId")
        fusion_hands = fetch_session_results(session_ids)
        aligned = align_to_cases(cases, fusion_hands)
        match_info = match_batch(trees, aligned, histories)
        match_info["n_fusion"] = len(fusion_hands)
        match_info["session_ids"] = session_ids

    bid = batch_id or batch_id_of(Path(display_name))
    title = Path(display_name).stem or bid
    batch_dir = SITE_DIR / "batches" / bid
    xlsx_path = batch_dir / "覆盖目录.xlsx"
    field_stats = (match_info or {}).get("field_stats")
    pair_rates = (match_info or {}).get("pair_rates")
    build_xlsx(xlsx_path, cases, fields, generated, display_name, field_stats)
    snapshot = {
        "generated_at": generated,
        "source": stored_source_path(src),
        "batch_id": bid,
        "n": len(cases),
        "session_ids": session_ids,
        "note": note or "",
        "join_order": "PREFLOP -> FLOP -> TURN -> RIVER, then sheet row order",
        "pairs": list(PAIRS),
        "fields": fields,
        "cases": cases,
        "match": {
            k: match_info[k]
            for k in ("pair_rates", "n_vis", "n_rule", "n_gold", "n_fusion", "session_ids")
            if match_info and k in match_info
        }
        if match_info
        else None,
    }
    batch = write_site(
        site_dir=SITE_DIR,
        batch_id=bid,
        batch_title=title,
        source_name=display_name,
        source_path=stored_source_path(src),
        generated=generated,
        cases=cases,
        fields=fields,
        actions=ACTIONS,
        players=PLAYERS,
        hero_pos=HERO_POS,
        snapshot=snapshot,
        uploader=uploader,
        pair_rates=pair_rates,
        field_stats=field_stats,
        session_id=",".join(session_ids) or stored_sid,
        trees=trees,
        histories=histories,
        note=note,
    )
    rel_from_out = f"e2e-quality/{batch['detail']}"
    (OUT_DIR / "数字e2e_覆盖目录.html").write_text(
        legacy_redirect(rel_from_out, "上传记录详情"),
        encoding="utf-8",
    )
    if COPY_DIR is not None:
        COPY_DIR.mkdir(parents=True, exist_ok=True)
        if COPY_SITE.exists():
            shutil.rmtree(COPY_SITE)
        shutil.copytree(SITE_DIR, COPY_SITE)
        (COPY_DIR / "数字e2e_覆盖目录.html").write_text(
            legacy_redirect(rel_from_out, "上传记录详情"),
            encoding="utf-8",
        )
        (COPY_DIR / "数字e2e_覆盖目录.xlsx").write_bytes(xlsx_path.read_bytes())
    batch["xlsx"] = str(xlsx_path)
    return batch


def lookup_existing_batch(batch_id: str, session_id: str, uploader: str = "") -> dict:
    """已有批次补填 sessionId，重新拉 fusion 算出匹配率。"""
    sid = (session_id or "").strip()
    if not sid:
        raise ValueError("请填写 fusion sessionId")
    bid = (batch_id or "").strip()
    index_path = SITE_DIR / "batches.json"
    if not bid and index_path.exists():
        try:
            data = json.loads(index_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            data = {}
        batches = data.get("batches") or []
        if batches:
            bid = str(batches[0].get("id") or "")
    if not bid:
        raise ValueError("还没有上传记录，请先上传 Excel 或 jsonl")
    meta_path = SITE_DIR / "batches" / bid / "meta.json"
    if not meta_path.exists():
        raise ValueError(f"找不到批次 {bid}，请先上传源数据")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    src = resolve_source_path(meta.get("source") or "", bid)
    if not src.exists():
        raise ValueError(f"源文件不在了：{meta.get('source') or src}。请重新上传 Excel 或 jsonl。")
    return ingest_source(
        src,
        uploader or meta.get("uploader") or "",
        batch_id=bid,
        source_name=meta.get("source_name") or src.name,
        session_id=sid,
        note=meta.get("note") or "",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="从源数据 Excel / jsonl 生成 e2e 质量站（匹配率 + 覆盖分布）")
    parser.add_argument(
        "--src",
        type=Path,
        default=SRC,
        help="源数据：Excel（读 auto_json 列）或 jsonl / json",
    )
    parser.add_argument("--uploader", default="", help="上传人，缺省列表里留空")
    parser.add_argument("--session-id", default="", help="fusion sessionId，可逗号分隔多个")
    args = parser.parse_args()
    if not args.src:
        raise SystemExit("请用 --src 指定 Excel 或 jsonl")
    try:
        batch = ingest_source(args.src, args.uploader, session_id=args.session_id)
    except (FileNotFoundError, ValueError) as e:
        raise SystemExit(str(e)) from e
    print(f"wrote {batch.get('xlsx')}")
    print(f"wrote {SITE_DIR / 'index.html'}")
    print(f"wrote {SITE_DIR / batch['detail']}")
    print(f"fields {batch.get('n_fields')} cases {batch.get('n')}")
    if COPY_SITE:
        print(f"copied to {COPY_SITE}")


if __name__ == "__main__":
    main()
