"""把某次执行的 result.xlsx 渲染成可在浏览器里看的页面。"""

from __future__ import annotations

import html
import json
from pathlib import Path
from typing import Any

_SECTIONS = {"结果一览", "失败原因", "按工具", "按优先级", "按分组"}
_RATE_HEADERS = {"占失败", "通过率"}
_DETAIL_SKIP = {"response_raw", "notes", "role", "expect_tools", "assert_hints"}
_DETAIL_ORDER = [
    "case_id",
    "step_no",
    "title",
    "group",
    "priority",
    "tool_name",
    "result",
    "失败说明",
    "assert_expr",
    "http_status",
    "elapsed_ms",
    "user_question",
    "response_summary",
]


def _cells(row: tuple[Any, ...] | None) -> list[Any]:
    if not row:
        return []
    vals: list[Any] = []
    for value in row:
        if value is None:
            vals.append("")
        elif isinstance(value, float):
            vals.append(value)
        else:
            vals.append(str(value).strip())
    while vals and vals[-1] == "":
        vals.pop()
    return vals


def _fmt(value: Any, header: str = "") -> str:
    if isinstance(value, float):
        if header in _RATE_HEADERS or header == "通过率":
            if 0 <= value <= 1:
                return f"{value * 100:.1f}%"
        if value.is_integer():
            return str(int(value))
        return f"{value:.1f}"
    return "" if value is None else str(value)


def _esc(value: Any) -> str:
    return html.escape(str(value), quote=True)


def parse_summary(ws) -> tuple[list[tuple[str, str]], str, list[dict[str, Any]]]:
    kpis: list[tuple[str, str]] = []
    callout = ""
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for row in ws.iter_rows(values_only=True):
        cells = _cells(row)
        if not cells:
            continue
        head = str(cells[0])
        if head in _SECTIONS:
            current = {"title": head, "rows": []}
            sections.append(current)
            if any(c != "" for c in cells[1:]):
                current["rows"].append(cells)
            continue
        if head in {"MCP 执行汇总", "项"} or head.startswith("导出时间"):
            if head == "项" and len(cells) > 1:
                current = {"title": "结果", "rows": [cells]}
                sections.append(current)
            continue
        if current is None:
            current = {"title": "汇总", "rows": []}
            sections.append(current)
        if current["title"] == "结果一览":
            if len(cells) >= 3 and isinstance(cells[2], str) and len(cells[2]) > 12:
                callout = cells[2]
            label = head
            raw = cells[1] if len(cells) > 1 else ""
            kpis.append((label, _fmt(raw, label)))
            continue
        current["rows"].append(cells)
    return kpis, callout, sections


def parse_detail(ws) -> tuple[list[str], list[dict[str, str]]]:
    row_iter = ws.iter_rows(values_only=True)
    header = _cells(next(row_iter, None))
    headers = [str(h) for h in header if str(h)]
    show = [h for h in _DETAIL_ORDER if h in headers and h not in _DETAIL_SKIP]
    for name in headers:
        if name not in show and name not in _DETAIL_SKIP:
            show.append(name)
    records: list[dict[str, str]] = []
    for row in row_iter:
        cells = list(row) if row else []
        if not any(c is not None and str(c).strip() for c in cells):
            continue
        item: dict[str, str] = {}
        for i, name in enumerate(headers):
            raw = cells[i] if i < len(cells) else ""
            text = _fmt(raw, name)
            if name in {"response_summary", "request_body", "原始报错", "response_raw"} and len(text) > 280:
                text = text[:280] + "…"
            item[name] = text
        if item.get("case_id") or item.get("result") or item.get("title"):
            records.append(item)
    return show, records


def _table(headers: list[Any], rows: list[list[Any]]) -> str:
    if not headers:
        return ""
    head_html = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    body = []
    for row in rows:
        padded = list(row) + [""] * (len(headers) - len(row))
        tds = []
        for i, cell in enumerate(padded[: len(headers)]):
            header = str(headers[i])
            text = _fmt(cell, header)
            cls = ""
            if header in {"条数", "失败"} and text not in {"", "0"}:
                cls = ' class="num bad"'
            elif header in {"通过", "通过率"}:
                cls = ' class="num ok"'
            elif header in {"条数", "总数", "占失败", "跳过"}:
                cls = ' class="num"'
            tds.append(f"<td{cls}>{_esc(text)}</td>")
        body.append("<tr>" + "".join(tds) + "</tr>")
    return f"<table><thead><tr>{head_html}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _summary_html(kpis: list[tuple[str, str]], callout: str, sections: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    if kpis:
        cards = []
        for label, value in kpis:
            tone = ""
            if label == "通过":
                tone = " ok"
            elif label == "失败":
                tone = " bad"
            cards.append(
                f'<div class="kpi{tone}"><span>{_esc(label)}</span><b>{_esc(value)}</b></div>'
            )
        parts.append('<div class="kpis">' + "".join(cards) + "</div>")
    if callout:
        parts.append(f'<p class="callout">{_esc(callout)}</p>')
    for section in sections:
        if section["title"] == "结果一览" or not section["rows"]:
            continue
        rows = section["rows"]
        headers = [str(c) for c in rows[0]]
        data = rows[1:]
        parts.append(f'<h2>{_esc(section["title"])}</h2>')
        parts.append(_table(headers, data))
    if not parts:
        parts.append('<p class="muted">这份结果没有汇总表。</p>')
    return "".join(parts)


def _detail_html(columns: list[str], records: list[dict[str, str]]) -> str:
    if not columns:
        return '<p class="muted">这份结果没有执行明细。</p>'
    head = "".join(f"<th>{_esc(c)}</th>" for c in columns)
    body = []
    for rec in records:
        result = (rec.get("result") or "").lower()
        tds = []
        for col in columns:
            text = rec.get(col, "")
            cls = ""
            if col == "result":
                cls = f' class="tag { _esc(result) }"'
            elif col == "失败说明":
                cls = ' class="note"'
            tds.append(f"<td{cls}>{_esc(text)}</td>")
        blob = _esc(" ".join(rec.get(c, "") for c in columns))
        body.append(f'<tr data-result="{_esc(result)}" data-q="{blob}">' + "".join(tds) + "</tr>")
    return f"""
    <div class="toolbar">
      <button type="button" class="on" data-filter="all">全部 {len(records)}</button>
      <button type="button" data-filter="failed">失败</button>
      <button type="button" data-filter="passed">通过</button>
      <button type="button" data-filter="skipped">跳过</button>
      <input id="q" placeholder="搜索用例、工具、失败说明" />
    </div>
    <div class="scroll">
      <table class="detail">
        <thead><tr>{head}</tr></thead>
        <tbody id="rows">{''.join(body)}</tbody>
      </table>
    </div>
    <script>
    const rows = [...document.querySelectorAll('#rows tr')];
    let filter = 'all';
    let query = '';
    function apply() {{
      const q = query.trim().toLowerCase();
      for (const tr of rows) {{
        const okf = filter === 'all' || tr.dataset.result === filter;
        const okq = !q || (tr.dataset.q || '').toLowerCase().includes(q);
        tr.hidden = !(okf && okq);
      }}
    }}
    document.querySelectorAll('.toolbar button').forEach(btn => {{
      btn.addEventListener('click', () => {{
        filter = btn.dataset.filter;
        document.querySelectorAll('.toolbar button').forEach(b => b.classList.toggle('on', b === btn));
        apply();
      }});
    }});
    document.getElementById('q').addEventListener('input', (ev) => {{
      query = ev.target.value || '';
      apply();
    }});
    </script>
    """


def render_result_page(*, xlsx: Path, run_id: str, download_href: str, back_href: str = "/index.html") -> str:
    from openpyxl import load_workbook

    meta_path = xlsx.parent / "meta.json"
    finished = run_id
    cases = ""
    if meta_path.exists():
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            finished = str(meta.get("finished_at") or finished)
            cases = Path(str(meta.get("cases_xlsx") or "")).name
        except json.JSONDecodeError:
            pass

    wb = load_workbook(xlsx, data_only=True)
    try:
        summary_ws = wb["汇总"] if "汇总" in wb.sheetnames else (
            wb["汇总统计"] if "汇总统计" in wb.sheetnames else None
        )
        detail_ws = wb["执行明细"] if "执行明细" in wb.sheetnames else None
        kpis, callout, sections = parse_summary(summary_ws) if summary_ws is not None else ([], "", [])
        columns, records = parse_detail(detail_ws) if detail_ws is not None else ([], [])
    finally:
        wb.close()

    summary = _summary_html(kpis, callout, sections)
    detail = _detail_html(columns, records)
    sub = finished + (f" · {cases}" if cases else "")
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>执行结果 { _esc(run_id) }</title>
<style>
  :root {{
    --bg:#0f1419; --panel:#1a2332; --line:#2d3a4d; --text:#e7ecf3;
    --muted:#8b9bb4; --accent:#3d9cf0; --ok:#3ecf8e; --bad:#f07178;
  }}
  * {{ box-sizing:border-box; }}
  body {{
    margin:0; font-family:"SF Pro Text","PingFang SC","Segoe UI",sans-serif;
    background:radial-gradient(1200px 600px at 10% -10%, #1b2a40, var(--bg));
    color:var(--text);
  }}
  a {{ color:var(--accent); text-decoration:none; }}
  header {{
    display:flex; justify-content:space-between; gap:16px; align-items:center;
    padding:18px 22px; border-bottom:1px solid var(--line);
  }}
  h1 {{ margin:0; font-size:20px; }}
  .sub {{ color:var(--muted); font-size:12px; margin-top:4px; }}
  .actions {{ display:flex; gap:10px; align-items:center; }}
  .btn {{
    background:var(--accent); color:#041018; border-radius:10px; padding:8px 14px;
    font-weight:700; font-size:13px;
  }}
  .btn.ghost {{ background:transparent; color:var(--text); border:1px solid var(--line); }}
  main {{ padding:18px 22px 40px; }}
  .tabs {{ display:flex; gap:8px; margin-bottom:14px; }}
  .tabs button {{
    background:#121a24; color:var(--text); border:1px solid var(--line); border-radius:999px;
    padding:6px 12px; cursor:pointer;
  }}
  .tabs button.on {{ background:var(--accent); color:#041018; border-color:transparent; font-weight:700; }}
  .kpis {{ display:flex; flex-wrap:wrap; gap:10px; margin-bottom:12px; }}
  .kpi {{
    background:var(--panel); border:1px solid var(--line); border-radius:12px;
    min-width:108px; padding:10px 12px;
  }}
  .kpi span {{ display:block; color:var(--muted); font-size:12px; }}
  .kpi b {{ font-size:22px; }}
  .kpi.ok b {{ color:var(--ok); }}
  .kpi.bad b {{ color:var(--bad); }}
  .callout {{
    background:#121a24; border:1px solid var(--line); border-radius:12px;
    padding:12px 14px; color:var(--text); line-height:1.5;
  }}
  h2 {{ font-size:15px; margin:18px 0 8px; }}
  table {{ width:100%; border-collapse:collapse; background:var(--panel); border:1px solid var(--line); }}
  th, td {{ padding:8px 10px; border-bottom:1px solid var(--line); text-align:left; font-size:12px; vertical-align:top; }}
  th {{ color:var(--muted); background:rgba(0,0,0,.2); position:sticky; top:0; }}
  td.num {{ text-align:right; font-variant-numeric:tabular-nums; }}
  td.bad {{ color:var(--bad); font-weight:700; }}
  td.ok {{ color:var(--ok); font-weight:700; }}
  td.note {{ white-space:pre-wrap; min-width:280px; max-width:520px; }}
  td.tag.passed {{ color:var(--ok); font-weight:700; }}
  td.tag.failed, td.tag.error {{ color:var(--bad); font-weight:700; }}
  .toolbar {{ display:flex; gap:8px; align-items:center; margin-bottom:10px; flex-wrap:wrap; }}
  .toolbar button, .toolbar input {{
    background:#121a24; color:var(--text); border:1px solid var(--line); border-radius:999px;
    padding:6px 12px;
  }}
  .toolbar button.on {{ background:var(--accent); color:#041018; font-weight:700; }}
  .toolbar input {{ border-radius:10px; min-width:240px; }}
  .scroll {{ overflow:auto; max-height:calc(100vh - 210px); border-radius:12px; }}
  .panel {{ display:none; }}
  .panel.on {{ display:block; }}
  .muted {{ color:var(--muted); }}
</style>
</head>
<body>
  <header>
    <div>
      <h1>执行结果</h1>
      <div class="sub">{_esc(sub)}</div>
    </div>
    <div class="actions">
      <a class="btn ghost" href="{_esc(back_href)}">返回列表</a>
      <a class="btn" href="{_esc(download_href)}" download>下载 Excel</a>
    </div>
  </header>
  <main>
    <div class="tabs">
      <button type="button" class="on" data-tab="summary">汇总</button>
      <button type="button" data-tab="detail">执行明细</button>
    </div>
    <section id="summary" class="panel on">{summary}</section>
    <section id="detail" class="panel">{detail}</section>
  </main>
  <script>
  document.querySelectorAll('.tabs button').forEach(btn => {{
    btn.addEventListener('click', () => {{
      document.querySelectorAll('.tabs button').forEach(b => b.classList.toggle('on', b === btn));
      document.querySelectorAll('.panel').forEach(p => p.classList.toggle('on', p.id === btn.dataset.tab));
    }});
  }});
  </script>
</body>
</html>
"""
