"""每次跑测归档到 reports/runs/{run_id}/，供门户历史列表回溯。"""

from __future__ import annotations

import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from common.result_store import ExecRecord


def new_run_id() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def summarize(records: list[ExecRecord]) -> dict[str, Any]:
    c = Counter(str(r.result or "unknown") for r in records)
    total = len(records)
    passed = c.get("passed", 0)
    failed = c.get("failed", 0) + c.get("error", 0)
    skipped = c.get("skipped", 0)
    return {
        "total": total,
        "passed": passed,
        "failed": failed,
        "skipped": skipped,
        "pass_rate": f"{(passed / total * 100):.1f}%" if total else "N/A",
    }


def _is_ui_record(r: ExecRecord) -> bool:
    tool = str(r.tool_name or "").strip().lower()
    role = str(r.role or "").strip().lower()
    return role == "ui" or tool in {"ui/chat", "ui", "production_chat"} or tool.startswith("ui/")


def detect_suite(records: list[ExecRecord], *, hint: str = "") -> str:
    """返回 api | ui | both | unknown。"""
    h = (hint or "").strip().lower()
    if h in {"api", "ui", "both"}:
        return h
    if not records:
        return "unknown"
    ui_n = sum(1 for r in records if _is_ui_record(r))
    api_n = len(records) - ui_n
    if ui_n and api_n:
        return "both"
    if ui_n:
        return "ui"
    if api_n:
        return "api"
    return "unknown"


def suite_from_result_xlsx(path: Path) -> str | None:
    """从执行明细 tool_name / role 推断 api|ui|both。"""
    path = Path(path)
    if not path.exists():
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "执行明细" not in wb.sheetnames:
            return None
        ws = wb["执行明细"]
        row_iter = ws.iter_rows(values_only=True)
        header = next(row_iter, None)
        if not header:
            return None
        headers = [str(h).strip() if h is not None else "" for h in header]
        tool_i = headers.index("tool_name") if "tool_name" in headers else -1
        role_i = headers.index("role") if "role" in headers else -1
        if tool_i < 0 and role_i < 0:
            return None
        ui_n = api_n = 0
        for row in row_iter:
            if not row:
                continue
            tool = str(row[tool_i] or "").strip().lower() if tool_i >= 0 and tool_i < len(row) else ""
            role = str(row[role_i] or "").strip().lower() if role_i >= 0 and role_i < len(row) else ""
            is_ui = role == "ui" or tool in {"ui/chat", "ui", "production_chat"} or tool.startswith("ui/")
            if is_ui:
                ui_n += 1
            else:
                api_n += 1
        if ui_n and api_n:
            return "both"
        if ui_n:
            return "ui"
        if api_n:
            return "api"
        return None
    finally:
        wb.close()


def refresh_blank_run_suites(reports_dir: Path) -> int:
    """历史 meta 无 suite 时，从 result.xlsx 回填。"""
    runs_root = Path(reports_dir) / "runs"
    if not runs_root.exists():
        return 0
    updated = 0
    for run_dir in runs_root.iterdir():
        if not run_dir.is_dir():
            continue
        meta_f = run_dir / "meta.json"
        xlsx = run_dir / "result.xlsx"
        if not meta_f.exists():
            continue
        try:
            meta = json.loads(meta_f.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        if meta.get("suite") in {"api", "ui", "both"}:
            continue
        suite = suite_from_result_xlsx(xlsx) if xlsx.exists() else None
        if not suite:
            suite = "api" if meta.get("legacy") else "unknown"
        meta["suite"] = suite
        meta_f.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        updated += 1
    return updated


def archive_run(
    *,
    reports_dir: Path,
    records: list[ExecRecord],
    excel_path: Path,
    html_report: Path | None = None,
    allure_report_dir: Path | None = None,
    cases_xlsx: str = "",
    exitstatus: int = 0,
    run_id: str | None = None,
    suite: str = "",
) -> dict[str, Any]:
    """写入 runs/{run_id}/meta.json + 结果文件副本。"""
    reports_dir = Path(reports_dir)
    rid = run_id or new_run_id()
    run_dir = reports_dir / "runs" / rid
    run_dir.mkdir(parents=True, exist_ok=True)

    stats = summarize(records)
    # Excel
    excel_dst = run_dir / "result.xlsx"
    if excel_path.exists():
        shutil.copy2(excel_path, excel_dst)

    # pytest-html
    html_dst = run_dir / "report.html"
    if html_report and html_report.exists():
        shutil.copy2(html_report, html_dst)

    # Allure（整目录复制，便于历史回看）
    allure_dst = run_dir / "allure-report"
    has_allure = False
    if allure_report_dir and (allure_report_dir / "index.html").exists():
        if allure_dst.exists():
            shutil.rmtree(allure_dst)
        shutil.copytree(allure_report_dir, allure_dst)
        has_allure = True

    meta = {
        "run_id": rid,
        "finished_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "cases_xlsx": cases_xlsx,
        "exitstatus": exitstatus,
        "suite": detect_suite(records, hint=suite),
        "has_excel": excel_dst.exists(),
        "has_html": html_dst.exists(),
        "has_allure": has_allure,
        **stats,
    }
    (run_dir / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    # 最新指针
    (reports_dir / "runs" / "latest.txt").write_text(rid, encoding="utf-8")
    return meta


def stats_from_result_xlsx(path: Path) -> dict[str, Any] | None:
    """从执行明细的 result 列数出通过/失败，补全历史列表。"""
    path = Path(path)
    if not path.exists():
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "执行明细" not in wb.sheetnames:
            return None
        ws = wb["执行明细"]
        row_iter = ws.iter_rows(values_only=True)
        header = next(row_iter, None)
        if not header:
            return None
        headers = [str(h).strip() if h is not None else "" for h in header]
        if "result" not in headers:
            return None
        idx = headers.index("result")
        counts: Counter[str] = Counter()
        for row in row_iter:
            if not row or idx >= len(row):
                continue
            val = str(row[idx] or "").strip().lower()
            if val == "passed":
                counts["passed"] += 1
            elif val in {"failed", "error"}:
                counts["failed"] += 1
            elif val == "skipped":
                counts["skipped"] += 1
        total = counts["passed"] + counts["failed"] + counts["skipped"]
        if total == 0:
            return None
        passed = counts["passed"]
        return {
            "total": total,
            "passed": passed,
            "failed": counts["failed"],
            "skipped": counts["skipped"],
            "pass_rate": f"{(passed / total * 100):.1f}%",
        }
    finally:
        wb.close()


def refresh_blank_run_stats(reports_dir: Path) -> int:
    """历史补录的 meta 没有通过/失败时，从 result.xlsx 回填。"""
    runs_root = Path(reports_dir) / "runs"
    if not runs_root.exists():
        return 0
    updated = 0
    for run_dir in runs_root.iterdir():
        if not run_dir.is_dir():
            continue
        meta_f = run_dir / "meta.json"
        xlsx = run_dir / "result.xlsx"
        if not meta_f.exists() or not xlsx.exists():
            continue
        try:
            meta = json.loads(meta_f.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        if meta.get("total") not in ("", None) and meta.get("passed") not in ("", None):
            continue
        stats = stats_from_result_xlsx(xlsx)
        if not stats:
            continue
        meta.update(stats)
        meta_f.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        updated += 1
    return updated


def list_runs(reports_dir: Path) -> list[dict[str, Any]]:
    runs_root = Path(reports_dir) / "runs"
    if not runs_root.exists():
        return []
    items: list[dict[str, Any]] = []
    for d in sorted(runs_root.iterdir(), reverse=True):
        if not d.is_dir():
            continue
        meta_f = d / "meta.json"
        if not meta_f.exists():
            continue
        try:
            meta = json.loads(meta_f.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        meta.setdefault("run_id", d.name)
        items.append(meta)
    return items


def enrich_latest_run(
    reports_dir: Path,
    *,
    html_report: Path | None = None,
    allure_report_dir: Path | None = None,
    suite: str = "",
) -> dict[str, Any] | None:
    """pytest 结束后由 shell 调用：把 HTML/Allure 补进最近一次 run 目录。"""
    reports_dir = Path(reports_dir)
    latest_f = reports_dir / "runs" / "latest.txt"
    if not latest_f.exists():
        return None
    rid = latest_f.read_text(encoding="utf-8").strip()
    run_dir = reports_dir / "runs" / rid
    meta_f = run_dir / "meta.json"
    if not meta_f.exists():
        return None
    meta = json.loads(meta_f.read_text(encoding="utf-8"))

    if html_report and html_report.exists():
        shutil.copy2(html_report, run_dir / "report.html")
        meta["has_html"] = True

    if allure_report_dir and (allure_report_dir / "index.html").exists():
        dst = run_dir / "allure-report"
        if dst.exists():
            shutil.rmtree(dst)
        shutil.copytree(allure_report_dir, dst)
        meta["has_allure"] = True

    hint = (suite or "").strip().lower()
    if hint in {"api", "ui", "both"}:
        meta["suite"] = hint
    elif meta.get("suite") not in {"api", "ui", "both"}:
        inferred = suite_from_result_xlsx(run_dir / "result.xlsx")
        if inferred:
            meta["suite"] = inferred

    meta_f.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta

