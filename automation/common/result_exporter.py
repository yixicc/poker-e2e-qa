"""把执行结果导出为 Excel：执行明细 + 汇总。"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from common.result_store import ExecRecord

# 左：原用例字段（与 cases/requests 对齐）；右：本次执行结果，不混排
CASE_HEADERS = [
    "case_id",
    "group",
    "title",
    "priority",
    "step_no",
    "user_question",
    "assert_hints",
    "expect_tools",
    "notes",
    "tool_name",
    "role",
    "metric_codes",
    "assert_expr",
]

EXEC_HEADERS = [
    "请求报文",
    "响应答复",
    "http_status",
    "执行时间ms",
    "执行结果",
    "结果分析",
]

DETAIL_HEADERS = CASE_HEADERS + EXEC_HEADERS

DETAIL_WIDTHS = {
    "case_id": 14,
    "group": 22,
    "title": 28,
    "priority": 9,
    "step_no": 9,
    "user_question": 32,
    "assert_hints": 28,
    "expect_tools": 24,
    "notes": 16,
    "tool_name": 22,
    "role": 10,
    "metric_codes": 16,
    "assert_expr": 22,
    "请求报文": 40,
    "响应答复": 48,
    "http_status": 12,
    "执行时间ms": 12,
    "执行结果": 12,
    "结果分析": 48,
}

_TEAL = "0F766E"

# 断言名和返回里真实字段的对应，只写能对上的
_FIELD_ALIAS = {
    "hand_history": "hands",
    "position_hand_history": "hands",
    "blind_recent_sessions": "hands",
    "hand_detail": "hand",
    "hand_actions": "hand",
    "hand_opponents": "hand",
    "opponent_vpip": "vpip",
    "opponent_pfr": "pfr",
    "opponent_cbet_rate": "cbet_flop",
    "opponent_fold_rate": "fold_to_3bet",
}

_NAVY = "1F4E79"
_NAVY_2 = "2E5A88"
_INK = "1F2937"
_MUTED = "64748B"
_LINE = "D6DEE8"
_ZEBRA = "F4F7FB"
_PAPER = "F8FAFC"
_GREEN = "166534"
_GREEN_BG = "E8F6EE"
_RED = "B42318"
_RED_BG = "FDECEC"
_AMBER = "92400E"
_WHITE = "FFFFFF"

_THIN = Border(
    left=Side(style="thin", color=_LINE),
    right=Side(style="thin", color=_LINE),
    top=Side(style="thin", color=_LINE),
    bottom=Side(style="thin", color=_LINE),
)
_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)
_LEFT = Alignment(vertical="center", horizontal="left", wrap_text=True)
_LEFT_TOP = Alignment(vertical="top", horizontal="left", wrap_text=True)


@dataclass
class FailureNote:
    category: str = ""
    field: str = ""
    text: str = ""
    group_text: str = ""


@dataclass
class _Node:
    kind: str
    text: str = ""
    children: dict[str, "_Node"] | None = None
    items: list["_Node"] | None = None


def _trunc(text: Any, n: int = 3000) -> str:
    s = "" if text is None else str(text)
    return s if len(s) <= n else s[:n] + f"...(截断共{len(s)}字)"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(size: int = 11, *, bold: bool = False, color: str = _INK) -> Font:
    return Font(size=size, bold=bold, color=color)


def _exc_message(longrepr: str) -> str:
    lines: list[str] = []
    for line in (longrepr or "").splitlines():
        stripped = line.strip()
        if stripped.startswith("E "):
            lines.append(re.sub(r"^E\s+", "", stripped))
    text = " ".join(lines) if lines else (longrepr or "").strip()
    return re.sub(
        r"^(?:AssertionError|Failed|McpError|Exception|Error):\s*",
        "",
        text,
        count=1,
    ).strip()


def _skip_ws(s: str, i: int) -> int:
    while i < len(s) and s[i] in " \t\r\n":
        i += 1
    return i


def _read_string(s: str, i: int) -> tuple[str, int]:
    quote = s[i]
    i += 1
    chars: list[str] = []
    while i < len(s):
        ch = s[i]
        if ch == "\\":
            if i + 1 >= len(s):
                return "".join(chars), len(s)
            chars.append(s[i + 1])
            i += 2
            continue
        if ch == quote:
            return "".join(chars), i + 1
        chars.append(ch)
        i += 1
    return "".join(chars), len(s)


def _read_value(s: str, i: int) -> tuple[_Node, int]:
    i = _skip_ws(s, i)
    if i >= len(s):
        return _Node("trunc", "截断"), i
    if s.startswith(("None", "null"), i):
        n = 4 if s.startswith("None", i) else 4
        return _Node("null", "null"), i + n
    if s.startswith(("True", "true"), i):
        return _Node("bool", "true"), i + 4
    if s.startswith("False", i):
        return _Node("bool", "false"), i + 5
    if s.startswith("false", i):
        return _Node("bool", "false"), i + 5
    if s[i] in "'\"":
        text, j = _read_string(s, i)
        shown = text if len(text) <= 48 else text[:48] + "…"
        return _Node("str", shown), j
    if s[i] in "-0123456789":
        j = i + 1
        while j < len(s) and s[j] in "0123456789.eE+-":
            j += 1
        return _Node("num", s[i:j]), j
    if s[i] == "{":
        children, j = _read_dict(s, i)
        return _Node("dict", "{…}", children=children), j
    if s[i] == "[":
        if s.startswith("[]", i):
            return _Node("empty_list", "空列表"), i + 2
        items, j = _read_list(s, i)
        kind = "empty_list" if not items else "list"
        text = "空列表" if not items else "列表"
        return _Node(kind, text, items=items), j
    return _Node("trunc", "截断"), i + 1


def _read_dict(s: str, i: int) -> tuple[dict[str, _Node], int]:
    i += 1
    out: dict[str, _Node] = {}
    while True:
        i = _skip_ws(s, i)
        if i >= len(s) or s[i] == "}":
            return out, min(i + 1, len(s))
        if s[i] not in "'\"":
            return out, i
        key, i = _read_string(s, i)
        i = _skip_ws(s, i)
        if i >= len(s) or s[i] != ":":
            return out, i
        node, i = _read_value(s, i + 1)
        out[key] = node
        i = _skip_ws(s, i)
        if i < len(s) and s[i] == ",":
            i += 1
            continue
        if i < len(s) and s[i] == "}":
            return out, i + 1
        return out, i


def _read_list(s: str, i: int) -> tuple[list[_Node], int]:
    i += 1
    items: list[_Node] = []
    while True:
        i = _skip_ws(s, i)
        if i >= len(s) or s[i] == "]":
            return items, min(i + 1, len(s))
        node, i = _read_value(s, i)
        items.append(node)
        i = _skip_ws(s, i)
        if i < len(s) and s[i] == ",":
            i += 1
            continue
        if i < len(s) and s[i] == "]":
            return items, i + 1
        return items, i


def _parse_mapping(blob: str) -> dict[str, _Node]:
    text = (blob or "").strip()
    start = text.find("{")
    if start < 0:
        return {}
    try:
        children, _ = _read_dict(text, start)
    except Exception:  # noqa: BLE001
        return {}
    return children


def _payload_blob(message: str, response_summary: str) -> str:
    marker = message.find("payload=")
    if marker >= 0:
        return message[marker + len("payload=") :]
    summary = (response_summary or "").strip()
    if summary.startswith("{") or summary.startswith("["):
        return summary
    return ""


def _keys(mapping: dict[str, _Node]) -> str:
    names = list(mapping.keys())
    if not names:
        return "（没解析出字段）"
    shown = names[:8]
    text = "、".join(shown)
    if len(names) > len(shown):
        text += " 等"
    return text


def _scalar(node: _Node | None) -> str:
    if node is None:
        return ""
    if node.kind == "null":
        return "null"
    if node.kind in {"num", "bool", "str", "empty_list"}:
        return node.text
    if node.kind == "list":
        return "列表"
    if node.kind == "dict":
        return "对象"
    return node.text or ""


def _pairs(mapping: dict[str, _Node], *, skip: set[str] | None = None, limit: int = 6) -> str:
    skip = skip or set()
    parts: list[str] = []
    for key, node in mapping.items():
        if key in skip or node.kind not in {"null", "num", "bool", "str"}:
            continue
        if node.kind == "str" and len(node.text) > 24:
            continue
        parts.append(f"{key}={_scalar(node)}")
        if len(parts) >= limit:
            break
    return "、".join(parts)


def _first_row(mapping: dict[str, _Node]) -> dict[str, _Node]:
    data = mapping.get("data")
    if data is None or data.kind != "list" or not data.items:
        return {}
    first = data.items[0]
    if first.kind == "dict" and first.children:
        return first.children
    return {}


def _hand_id(mapping: dict[str, _Node]) -> str:
    hand = mapping.get("hand")
    if hand is None or not hand.children:
        return ""
    return _scalar(hand.children.get("hand_id"))


def _alias_phrase(field_name: str, mapping: dict[str, _Node]) -> tuple[str, str]:
    """明细里带具体值，归类说明里不带，避免同一原因被数值拆开。"""
    alias = _FIELD_ALIAS.get(field_name)
    if not alias or alias not in mapping:
        return "", ""
    node = mapping[alias]
    if alias == "hands":
        count = _scalar(mapping.get("matched_count"))
        detail = "牌在 hands 里" + (f"，matched_count={count}" if count else "")
        return detail, "牌在 hands 里，没有这个字段。"
    if alias == "hand" or node.kind == "dict":
        hid = _hand_id(mapping)
        detail = "牌谱在 hand 对象里" + (f"，hand_id={hid}" if hid else "")
        return detail, "牌谱在 hand 对象里，没有这个字段。"
    value = _scalar(node)
    detail = f"实际字段名是 {alias}" + (f"={value}" if value else "")
    return detail, f"实际字段名是 {alias}，不是这个断言名。"


def explain_failure(rec: ExecRecord) -> FailureNote:
    """把 pytest 原文收成一两句人话。通过的用例返回空说明。"""
    result = str(rec.result or "")
    if result == "passed" or not result:
        return FailureNote()
    if result == "skipped":
        reason = _exc_message(rec.error_message) or str(rec.error_message or "").strip()
        reason = re.sub(r"\s+", " ", reason)
        text = f"跳过：{reason[:180]}" if reason else "跳过"
        return FailureNote("跳过", "", text, text)

    message = _exc_message(rec.error_message)
    tool = str(rec.tool_name or "").strip() or "接口"
    field_name = str(rec.assert_expr or "").strip()

    if message.startswith("数据不符"):
        text = re.sub(r"\s+", " ", message).strip()
        return FailureNote("数据不符", "", text[:500], "回答里的数据和预期描述对不上")
    if "tool isError" in message or message.startswith("tool isError"):
        return _explain_tool_error(message, tool)
    missing = re.search(r"缺少字段\s+([A-Za-z0-9_]+)", message)
    if missing:
        return _explain_missing_field(missing.group(1), message, rec.response_summary, tool)
    return _explain_other(message, tool, field_name)


def _explain_tool_error(message: str, tool: str) -> FailureNote:
    body = re.sub(r"^tool isError:\s*", "", message)
    body = re.sub(r"For further information visit \S+", "", body)
    body = re.sub(r"\s+", " ", body).strip()
    if "input_value=None" in body or "NoneType" in body:
        text = f"工具 {tool} 执行失败：服务端校验不通过，实际返回 None，不是合法结果。"
    elif body.lower().startswith("error executing tool"):
        text = f"工具 {tool} 执行失败：{body[:160]}"
    else:
        text = f"工具 {tool} 返回错误：{body[:180] or '工具标记 isError'}"
    return FailureNote("工具报错", "", text, text)


def _explain_missing_field(field_name: str, message: str, summary: str, tool: str) -> FailureNote:
    mapping = _parse_mapping(_payload_blob(message, summary))
    top_keys = _keys(mapping)
    row = _first_row(mapping)
    if field_name in row and row[field_name].kind == "null":
        siblings = _pairs(row, skip={field_name})
        extra = f"\n同一行还有 {siblings}。" if siblings else ""
        text = (
            f"断言字段「{field_name}」在返回的 data 第一行里，但值是 null。{extra}\n"
            "接口有结果，这个指标没有样本；断言把 null 当成了缺少字段。"
        )
        group = "data 第一行有这个字段，但值是 null。断言把 null 当成缺少字段。"
        return FailureNote("指标值为空", field_name, text, group)
    if field_name in mapping and mapping[field_name].kind == "null":
        text = (
            f"断言字段「{field_name}」在返回顶层，但值是 null。\n"
            "接口有结果，这个字段没有数；断言把 null 当成了缺少字段。"
        )
        group = "返回顶层有这个字段，但值是 null。断言把 null 当成缺少字段。"
        return FailureNote("指标值为空", field_name, text, group)

    matched = mapping.get("matched_count")
    hands = mapping.get("hands")
    empty_hands = (matched is not None and matched.text == "0") or (
        hands is not None and hands.kind == "empty_list"
    )
    if empty_hands and (hands is not None or "find_hands" in tool):
        hole = mapping.get("hole_cards_filter")
        hole_txt = ""
        if hole is not None and hole.kind == "str" and hole.text:
            hole_txt = f"筛选 {hole.text}，"
        text = (
            f"断言要字段「{field_name}」，返回里没有这个键。\n"
            f"这次没有匹配到手牌：{hole_txt}matched_count=0，hands 为空。\n"
            f"实际顶层字段：{top_keys}。"
        )
        group = "没有匹配到手牌（matched_count=0，hands 为空），返回里也没有这个字段。"
        return FailureNote("没有匹配到数据", field_name, text, group)

    detail, group = _alias_phrase(field_name, mapping)
    if not detail and row:
        detail, group = _alias_phrase(field_name, row)

    if "opponent_name" in mapping and field_name.endswith("_list"):
        name = _scalar(mapping.get("opponent_name"))
        text = (
            f"断言要字段「{field_name}」，返回里没有这个键。\n"
            f"返回的是单个对手 {name} 的画像，不是列表。\n"
            f"实际顶层字段：{top_keys}。"
        )
        group = "返回的是单个对手画像，没有列表字段。"
        return FailureNote("字段名对不上", field_name, text, group)

    if detail:
        text = (
            f"断言要字段「{field_name}」，返回里没有这个键。{detail}。\n"
            f"实际顶层字段：{top_keys}。"
        )
        return FailureNote("字段名对不上", field_name, text, group)

    if "hand" in mapping and mapping["hand"].kind == "dict":
        hid = _hand_id(mapping)
        hid_txt = f"，hand_id={hid}" if hid else ""
        text = (
            f"断言要字段「{field_name}」，返回里没有这个键。牌谱在 hand 对象里{hid_txt}。\n"
            f"实际顶层字段：{top_keys}。"
        )
        group = "牌谱在 hand 对象里，没有这个字段。"
        return FailureNote("字段名对不上", field_name, text, group)

    if "opponent_name" in mapping:
        name = _scalar(mapping.get("opponent_name"))
        hand_n = _scalar(mapping.get("hands"))
        stats = _pairs(mapping, skip={"opponent_name"}, limit=5)
        who = f"对手 {name}" + (f"（{hand_n} 手）" if hand_n else "")
        stat_txt = f"：{stats}" if stats else ""
        text = (
            f"断言要字段「{field_name}」，返回里没有这个键。\n"
            f"{who}的指标是平铺的{stat_txt}。\n"
            f"实际顶层字段：{top_keys}。"
        )
        group = "对手指标是平铺字段（vpip、pfr 等），没有这个键。"
        return FailureNote("字段名对不上", field_name, text, group)

    if row:
        text = (
            f"断言要字段「{field_name}」，返回顶层和 data 第一行都没有。\n"
            f"data 第一行的字段：{_keys(row)}。\n"
            f"实际顶层字段：{top_keys}。"
        )
        group = "字段不在顶层，也不在 data 第一行。"
        return FailureNote("字段名对不上", field_name, text, group)

    text = (
        f"断言要字段「{field_name}」，返回里没有这个键。\n"
        f"工具 {tool}。实际顶层字段：{top_keys}。"
    )
    group = "返回里没有断言要的字段。"
    return FailureNote("字段名对不上", field_name, text, group)


def _explain_other(message: str, tool: str, field_name: str) -> FailureNote:
    compact = re.sub(r"\s+", " ", message).strip()
    payload_at = compact.find("payload=")
    if payload_at >= 0:
        compact = compact[:payload_at].rstrip("，, ")
    if compact.startswith("HTTP ") or "unauthorized" in compact.lower() or "Session not found" in compact:
        text = f"接口失败：{compact[:180]}"
        return FailureNote("接口失败", field_name, text, text)
    if "无法解析 JSON-RPC" in compact or compact.startswith("jsonrpc error"):
        text = f"响应解析失败：{compact[:180]}"
        return FailureNote("接口失败", field_name, text, text)
    if "未包含" in compact or "!=" in compact or "期望" in compact:
        text = f"结果不符合预期：{compact[:180]}"
        return FailureNote("结果不符合预期", field_name, text, text)
    text = compact[:200] or f"工具 {tool} 失败，原始报错为空。"
    return FailureNote("其它", field_name, text, text)


def _result_bucket(result: str) -> str:
    if result == "passed":
        return "passed"
    if result == "skipped":
        return "skipped"
    if result in {"failed", "error"}:
        return "failed"
    return "other"


def _counts(records: list[ExecRecord]) -> Counter:
    return Counter(_result_bucket(str(r.result or "")) for r in records)


def _rate(passed: int, total: int) -> float | None:
    if not total:
        return None
    return passed / total


def _elapsed(records: list[ExecRecord]) -> tuple[float | None, float | None]:
    vals: list[float] = []
    for rec in records:
        try:
            vals.append(float(rec.elapsed_ms))
        except (TypeError, ValueError):
            continue
    if not vals:
        return None, None
    return round(sum(vals) / len(vals), 1), round(max(vals), 1)


def _banner(ws: Worksheet, row: int, title: str, cols: int = 7) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    fill = _fill(_NAVY)
    font = _font(12, bold=True, color=_WHITE)
    for col in range(1, cols + 1):
        cell = ws.cell(row, col, title if col == 1 else None)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 24


def _paint_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    fill = _fill(_NAVY_2)
    font = _font(10, bold=True, color=_WHITE)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = _CENTER
        cell.border = _THIN
    ws.row_dimensions[row].height = 22


def _paint_data(
    ws: Worksheet,
    row: int,
    values: list[Any],
    *,
    zebra: bool,
    wrap_cols: set[int] | None = None,
    center_cols: set[int] | None = None,
) -> None:
    wrap_cols = wrap_cols or set()
    center_cols = center_cols or set()
    fill = _fill(_ZEBRA if zebra else _WHITE)
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.fill = fill
        cell.font = _font(10)
        cell.border = _THIN
        if col in wrap_cols:
            cell.alignment = _LEFT_TOP
        elif col in center_cols:
            cell.alignment = _CENTER
        else:
            cell.alignment = _LEFT
    ws.row_dimensions[row].height = 32 if wrap_cols else 20


def _write_breakdown_table(
    ws: Worksheet,
    start: int,
    headers: list[str],
    rows: list[list[Any]],
    *,
    wrap_cols: set[int],
    center_cols: set[int],
    percent_cols: set[int] | None = None,
) -> int:
    percent_cols = percent_cols or set()
    _paint_header(ws, start, headers)
    if not rows:
        _paint_data(ws, start + 1, ["—"] + [""] * (len(headers) - 1), zebra=False)
        return start + 2
    for offset, values in enumerate(rows):
        row = start + 1 + offset
        _paint_data(
            ws,
            row,
            values,
            zebra=offset % 2 == 1,
            wrap_cols=wrap_cols,
            center_cols=center_cols,
        )
        for col in percent_cols:
            cell = ws.cell(row, col)
            if isinstance(cell.value, float):
                cell.number_format = "0.0%"
        fail_col = headers.index("失败") + 1 if "失败" in headers else None
        if fail_col and isinstance(values[fail_col - 1], int) and values[fail_col - 1] > 0:
            ws.cell(row, fail_col).font = _font(10, bold=True, color=_RED)
        count_col = headers.index("条数") + 1 if "条数" in headers else None
        if count_col:
            ws.cell(row, count_col).font = _font(10, bold=True, color=_RED)
            ws.cell(row, count_col).alignment = _CENTER
    return start + 1 + len(rows)


def _grouped_rows(records: list[ExecRecord]) -> tuple[list[FailureNote], list[list[Any]], str]:
    notes = [explain_failure(rec) for rec in records]
    buckets: dict[tuple[str, str, str, str], list[str]] = defaultdict(list)
    order: list[tuple[str, str, str, str]] = []
    for rec, note in zip(records, notes):
        if note.category in {"", "跳过"}:
            continue
        example = f"{rec.case_id} step{rec.step_no}".strip()
        key = (
            note.category,
            str(rec.tool_name or "—"),
            note.field or "—",
            note.group_text,
        )
        if key not in buckets:
            order.append(key)
        buckets[key].append(example)
    failed_n = sum(len(v) for v in buckets.values())
    ranked = sorted(order, key=lambda k: (-len(buckets[k]), k[0], k[1], k[2]))
    rows: list[list[Any]] = []
    for key in ranked:
        category, tool, field_name, group_text = key
        examples = buckets[key]
        shown = "、".join(examples[:3])
        if len(examples) > 3:
            shown += f" 等{len(examples)}条"
        share = (len(examples) / failed_n) if failed_n else 0
        rows.append([category, len(examples), share, tool, field_name, group_text, shown])
    by_cat = Counter(note.category for note in notes if note.category not in {"", "跳过"})
    if not by_cat:
        headline = "这次没有失败。"
    else:
        bits = "，".join(f"{n} 条{name}" for name, n in by_cat.most_common())
        headline = f"失败 {failed_n} 条，已按原因归类：{bits}。同一原因在下表只占一行，单条说法在「执行明细」的「结果分析」。"
    return notes, rows, headline


def _stat_line(ws: Worksheet, row: int, label: str, value: Any, *, value_color: str = _INK) -> None:
    label_cell = ws.cell(row, 1, label)
    label_cell.font = _font(10, color=_MUTED)
    label_cell.fill = _fill(_PAPER)
    label_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    label_cell.border = _THIN
    value_cell = ws.cell(row, 2, value)
    value_cell.font = _font(16, bold=True, color=value_color)
    value_cell.fill = _fill(_WHITE)
    value_cell.alignment = _CENTER
    value_cell.border = _THIN
    ws.row_dimensions[row].height = 26


def _build_summary(wb: Workbook, records: list[ExecRecord], fail_rows: list[list[Any]], headline: str) -> None:
    ws = wb.active
    ws.title = "汇总"
    widths = {"A": 18, "B": 14, "C": 12, "D": 30, "E": 26, "F": 56, "G": 36}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    ws.merge_cells("A1:G1")
    title = ws["A1"]
    title.value = "MCP 执行汇总"
    title.font = _font(20, bold=True, color=_WHITE)
    title.fill = _fill(_NAVY)
    title.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 36
    for col in range(1, 8):
        ws.cell(1, col).fill = _fill(_NAVY)

    ws.merge_cells("A2:G2")
    sub = ws["A2"]
    sub.value = datetime.now().strftime("导出时间 %Y-%m-%d %H:%M:%S")
    sub.font = _font(10, color="D6DEE8")
    sub.fill = _fill("243B53")
    sub.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 20
    for col in range(1, 8):
        ws.cell(2, col).fill = _fill("243B53")

    total = len(records)
    counted = _counts(records)
    passed = counted.get("passed", 0)
    failed = counted.get("failed", 0)
    skipped = counted.get("skipped", 0)
    avg_ms, max_ms = _elapsed(records)
    rate = _rate(passed, total)

    _banner(ws, 4, "结果一览")
    stats = [
        ("执行总数", total, _INK),
        ("通过", passed, _GREEN),
        ("失败", failed, _RED),
        ("跳过", skipped, _AMBER),
        ("通过率", rate if rate is not None else "—", _NAVY),
        ("平均耗时", f"{avg_ms} ms" if avg_ms is not None else "—", _INK),
        ("最慢一条", f"{max_ms} ms" if max_ms is not None else "—", _INK),
    ]
    for offset, (label, value, color) in enumerate(stats):
        _stat_line(ws, 5 + offset, label, value, value_color=color)
        if label == "通过率" and isinstance(value, float):
            ws.cell(5 + offset, 2).number_format = "0.0%"

    ws.merge_cells("C5:G11")
    callout = ws["C5"]
    callout.value = headline
    callout.font = _font(12, color=_INK)
    callout.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
    callout.fill = _fill(_PAPER)
    for row in range(5, 12):
        for col in range(3, 8):
            cell = ws.cell(row, col)
            cell.fill = _fill(_PAPER)
            cell.border = _THIN
    callout.border = _THIN

    cursor = 13
    _banner(ws, cursor, "失败原因")
    fail_header = cursor + 1
    cursor = _write_breakdown_table(
        ws,
        cursor + 1,
        ["类型", "条数", "占失败", "工具", "断言字段", "说明", "代表用例"],
        fail_rows,
        wrap_cols={6, 7},
        center_cols={2, 3},
        percent_cols={3},
    )
    for row in range(fail_header + 1, cursor):
        if ws.cell(row, 6).value:
            ws.row_dimensions[row].height = 36

    cursor += 1
    _banner(ws, cursor, "按工具")
    tool_rows = _dimension_rows(records, lambda r: str(r.tool_name or "未知"))
    cursor = _write_breakdown_table(
        ws,
        cursor + 1,
        ["工具", "总数", "通过", "失败", "跳过", "通过率"],
        tool_rows,
        wrap_cols=set(),
        center_cols={2, 3, 4, 5, 6},
        percent_cols={6},
    )
    _tint_rate(ws, cursor, tool_rows, rate_col=6, fail_col=4)

    cursor += 1
    _banner(ws, cursor, "按优先级")
    prio_rows = _dimension_rows(records, lambda r: str(r.priority or "未标"), key_order=_priority_key)
    cursor = _write_breakdown_table(
        ws,
        cursor + 1,
        ["优先级", "总数", "通过", "失败", "跳过", "通过率"],
        prio_rows,
        wrap_cols=set(),
        center_cols={2, 3, 4, 5, 6},
        percent_cols={6},
    )
    _tint_rate(ws, cursor, prio_rows, rate_col=6, fail_col=4)

    cursor += 1
    _banner(ws, cursor, "按分组")
    group_rows = _dimension_rows(records, lambda r: str(r.group or "未分组"))
    cursor = _write_breakdown_table(
        ws,
        cursor + 1,
        ["分组", "总数", "通过", "失败", "跳过", "通过率"],
        group_rows,
        wrap_cols={1},
        center_cols={2, 3, 4, 5, 6},
        percent_cols={6},
    )
    _tint_rate(ws, cursor, group_rows, rate_col=6, fail_col=4)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.tabColor = _NAVY
    ws.print_title_rows = "1:2"
    ws.oddHeader.left.text = "MCP 执行汇总"


def _priority_key(name: str) -> tuple[int, str]:
    order = {"P0": 0, "🔴": 0, "P1": 1, "P2": 2, "P3": 3}
    return (order.get(name, 9), name)


def _dimension_rows(records: list[ExecRecord], label_of, key_order=None) -> list[list[Any]]:
    grouped: dict[str, Counter] = defaultdict(Counter)
    for rec in records:
        grouped[label_of(rec)][_result_bucket(str(rec.result or ""))] += 1

    def sort_key(item: tuple[str, Counter]) -> tuple:
        name, counter = item
        if key_order:
            return key_order(name)
        return (-sum(counter.values()), name)

    rows: list[list[Any]] = []
    for name, counter in sorted(grouped.items(), key=sort_key):
        total = sum(counter.values())
        passed = counter.get("passed", 0)
        failed = counter.get("failed", 0)
        skipped = counter.get("skipped", 0)
        rows.append([name, total, passed, failed, skipped, _rate(passed, total)])
    return rows


def _tint_rate(ws: Worksheet, end_row: int, rows: list[list[Any]], *, rate_col: int, fail_col: int) -> None:
    # end_row 是表后的下一行；数据从 end_row - len(rows) 开始，且前面还有表头
    if not rows:
        return
    first = end_row - len(rows)
    for offset, values in enumerate(rows):
        cell = ws.cell(first + offset, rate_col)
        rate = values[rate_col - 1]
        failed = values[fail_col - 1]
        if not isinstance(rate, float):
            continue
        if failed == 0:
            cell.fill = _fill(_GREEN_BG)
            cell.font = _font(10, bold=True, color=_GREEN)
        elif rate < 0.5:
            cell.fill = _fill(_RED_BG)
            cell.font = _font(10, bold=True, color=_RED)


def _detail_cell_value(header: str, rec: ExecRecord, note: FailureNote) -> Any:
    """原用例列直接取字段；执行列单独映射，避免和 case 字段混名。"""
    if header in CASE_HEADERS:
        raw = getattr(rec, header, "")
        if header == "assert_hints":
            return _trunc(raw, 800)
        return raw if raw is not None else ""

    if header == "请求报文":
        body = (rec.request_body or "").strip()
        # UI 聊天没有 MCP body，执行时发出去的就是问题原文
        if not body and str(rec.role or "").lower() in {"ui", "chat"}:
            body = (rec.user_question or "").strip()
        return _trunc(body, 4000)
    if header == "响应答复":
        reply = (rec.model_reply or "").strip() or (rec.response_raw or "").strip()
        if not reply:
            reply = (rec.response_summary or "").strip()
        return _trunc(reply, 30000)
    if header == "http_status":
        return rec.http_status if rec.http_status is not None else ""
    if header == "执行时间ms":
        return rec.elapsed_ms if rec.elapsed_ms is not None else ""
    if header == "执行结果":
        return rec.result or ""
    if header == "结果分析":
        analysis = (note.text or "").strip()
        err = (rec.error_message or "").strip()
        if analysis and err and err not in analysis:
            return _trunc(f"{analysis}\n---\n原始报错：{err}", 4000)
        return _trunc(analysis or err, 4000)
    return ""


def _build_detail(wb: Workbook, records: list[ExecRecord], notes: list[FailureNote]) -> None:
    ws = wb.create_sheet("执行明细")
    ws.append(DETAIL_HEADERS)
    result_col = DETAIL_HEADERS.index("执行结果") + 1
    note_col = DETAIL_HEADERS.index("结果分析") + 1
    reply_col = DETAIL_HEADERS.index("响应答复") + 1
    req_col = DETAIL_HEADERS.index("请求报文") + 1
    case_n = len(CASE_HEADERS)

    for rec, note in zip(records, notes):
        ws.append([_detail_cell_value(h, rec, note) for h in DETAIL_HEADERS])

    case_fill = _fill(_NAVY)
    exec_fill = _fill(_TEAL)
    header_font = _font(10, bold=True, color=_WHITE)
    for idx, cell in enumerate(ws[1], start=1):
        cell.fill = case_fill if idx <= case_n else exec_fill
        cell.font = header_font
        cell.alignment = _CENTER
    ws.row_dimensions[1].height = 26
    for header, width in DETAIL_WIDTHS.items():
        idx = DETAIL_HEADERS.index(header) + 1
        ws.column_dimensions[get_column_letter(idx)].width = width

    pass_fill = _fill("C8E6C9")
    fail_fill = _fill("FFCDD2")
    skip_fill = _fill("FFF3C4")
    note_fill = _fill("FFF7ED")
    for row in range(2, ws.max_row + 1):
        result_cell = ws.cell(row, result_col)
        status = str(result_cell.value or "").lower()
        if status == "passed":
            result_cell.fill = pass_fill
        elif status in {"failed", "error"}:
            result_cell.fill = fail_fill
        elif status == "skipped":
            result_cell.fill = skip_fill
        result_cell.alignment = _CENTER
        result_cell.font = _font(10, bold=True)
        for col in (req_col, reply_col, note_col):
            cell = ws.cell(row, col)
            cell.alignment = _LEFT_TOP
            cell.font = _font(10)
        note_cell = ws.cell(row, note_col)
        if note_cell.value:
            note_cell.fill = note_fill
            lines = str(note_cell.value).count("\n") + 1
            ws.row_dimensions[row].height = min(96, 16 * lines + 8)
    ws.sheet_properties.tabColor = "C2410C"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_view.zoomScale = 110
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{max(ws.max_row, 1)}"


def export_results_xlsx(records: list[ExecRecord], out: Path) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    notes, fail_rows, headline = _grouped_rows(records)
    wb = Workbook()
    _build_summary(wb, records, fail_rows, headline)
    _build_detail(wb, records, notes)
    wb.save(out)
    return out
